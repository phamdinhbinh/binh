
import sys
import json
import os
import subprocess
import traceback
import importlib
import re
import logging
import threading
from io import BytesIO
from pathlib import Path
from dataclasses import dataclass, asdict
from typing import Dict, List, Tuple, Optional, Any

import numpy as np
import pandas as pd
import pyqtgraph as pg
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment as XlsxAlignment
from openpyxl.utils import get_column_letter
from PySide6.QtCore import Qt, QMimeData, QPoint, Signal, QObject, QThread, QEvent, QTimer, QSettings
from PySide6.QtGui import QDrag, QPainter, QColor, QFont, QPen, QBrush, QIcon, QAction, QKeySequence
from PySide6.QtWidgets import (
    QApplication, QFileDialog, QHBoxLayout, QLabel, QListWidget,
    QListWidgetItem, QMainWindow, QMessageBox, QPushButton,
    QSplitter, QVBoxLayout, QWidget, QLineEdit, QComboBox,
    QCheckBox, QDoubleSpinBox, QGroupBox, QSpinBox, QFrame,
    QTreeWidget, QTreeWidgetItem, QHeaderView, QAbstractItemView,
    QToolButton, QTabWidget, QScrollArea, QSizePolicy, QMenu,
    QDialog, QFormLayout, QDialogButtonBox, QTextEdit, QCompleter, QGridLayout,
    QProgressBar, QTableWidget, QTableWidgetItem, QAbstractScrollArea
)
from asammdf import MDF
from mf4reader.align import estimate_offset, trigger_time
from mf4reader.cache import file_hash


@dataclass
class Alignment:
    offset: float = 0.0
    source: str = "none"
    quality: float = 0.0
    locked: bool = False
    reference: str = ""
    ref_channel: str = ""

    def as_dict(self):
        return asdict(self)


def ensure_python_can():
    try:
        return importlib.import_module("can")
    except ImportError as exc:
        raise ImportError(
            "BLF support requires the bundled python-can package. "
            "Rebuild the application with python-can installed in the build environment."
        ) from exc


def ensure_cantools():
    try:
        return importlib.import_module("cantools")
    except ImportError as exc:
        raise ImportError(
            "DBC decoding requires the bundled cantools package. "
            "Rebuild the application with cantools installed in the build environment."
        ) from exc


def load_dbc_robust(path: str):
    """
    Load a DBC file with fallback across multiple character encodings.
    CANdb++ and OEM DBC files frequently use Windows-1252 or Latin-1.
    """
    cantools = ensure_cantools()
    encodings = ["utf-8", "cp1252", "latin-1", "iso-8859-1", "utf-8-sig"]
    last_err = None
    for enc in encodings:
        try:
            return cantools.database.load_file(path, encoding=enc)
        except UnicodeDecodeError as exc:
            last_err = exc
            continue
        except Exception as exc:
            last_err = exc
            continue

    # Fallback: Read binary and decode with error replacement
    try:
        with open(path, "r", encoding="latin-1", errors="replace") as f:
            content = f.read()
        return cantools.database.load_string(content, database_format="dbc")
    except Exception as exc:
        raise last_err or exc


_DBC_DATABASE_CACHE: Dict[Tuple[str, int], Any] = {}


pg.setConfigOptions(antialias=False, useOpenGL=True)
try:
    import OpenGL  # noqa: F401
except Exception:
    pg.setConfigOptions(useOpenGL=False)


class ThemeManager:
    """Load, apply, and persist the application's external QSS themes."""

    THEMES = {
        "one-dark": ("One Dark", "#282C34", True),
        "one-light": ("One Light", "#FAFAFA", False),
        "solarized-dark": ("Solarized Dark", "#002B36", True),
        "high-contrast": ("High Contrast", "#000000", True),
    }

    def __init__(self, settings: Optional[QSettings] = None):
        self.settings = settings or QSettings("MF4Reader", "MF4VehicleExplorer")
        saved = self.settings.value("theme/name", "one-dark")
        self.name = saved if saved in self.THEMES else "one-dark"

    @property
    def display_name(self) -> str:
        return self.THEMES[self.name][0]

    @property
    def is_dark(self) -> bool:
        return self.THEMES[self.name][2]

    @property
    def background(self) -> str:
        return self.THEMES[self.name][1]

    def stylesheet(self) -> str:
        path = Path(__file__).resolve().parent / "themes" / f"{self.name}.qss"
        try:
            return path.read_text(encoding="utf-8")
        except OSError:
            return STYLE_LIGHT if self.name == "one-light" else STYLE_DARK

    def apply(self, application: QApplication, widget: QWidget) -> None:
        application.setStyleSheet(self.stylesheet())
        widget.setStyleSheet("")
        self.settings.setValue("theme/name", self.name)
        self.settings.sync()

    def set_theme(self, name: str) -> None:
        if name not in self.THEMES:
            raise ValueError(f"Unknown theme: {name}")
        self.name = name


class CommandPalette(QDialog):
    """Small fuzzy command launcher matching the Atom command-palette workflow."""

    def __init__(self, commands: List[Tuple[str, Any]], parent=None):
        super().__init__(parent)
        self.setWindowTitle("Command Palette")
        self.setModal(True)
        self.resize(560, 360)
        self._commands = commands
        layout = QVBoxLayout(self)
        self.search = QLineEdit()
        self.search.setPlaceholderText("Type a command…")
        self.search.textChanged.connect(self._filter)
        layout.addWidget(self.search)
        self.results = QListWidget()
        self.results.itemActivated.connect(self._run_selected)
        layout.addWidget(self.results)
        self._filter("")

    def _filter(self, query: str) -> None:
        terms = query.casefold().split()
        self.results.clear()
        for label, callback in self._commands:
            haystack = label.casefold()
            if all(term in haystack for term in terms):
                item = QListWidgetItem(label)
                item.setData(Qt.UserRole, callback)
                self.results.addItem(item)
        if self.results.count():
            self.results.setCurrentRow(0)

    def _run_selected(self, item: QListWidgetItem) -> None:
        callback = item.data(Qt.UserRole)
        self.accept()
        callback()

    def keyPressEvent(self, event) -> None:
        if event.key() in (Qt.Key_Return, Qt.Key_Enter) and self.results.currentItem():
            self._run_selected(self.results.currentItem())
            return
        super().keyPressEvent(event)

DARK_BG = "#161618"
DARK_SURFACE = "#1E1E22"
DARK_SURFACE_ALT = "#28282D"
DARK_BORDER = "#333338"
DARK_TEXT_PRIMARY = "#F5F5F7"
DARK_TEXT_SECONDARY = "#86868B"
DARK_ACCENT = "#0A84FF"

LIGHT_BG = "#F5F5F7"
LIGHT_SURFACE = "#FFFFFF"
LIGHT_SURFACE_ALT = "#EBEBF0"
LIGHT_BORDER = "#D1D1D6"
LIGHT_TEXT_PRIMARY = "#1D1D1F"
LIGHT_TEXT_SECONDARY = "#86868B"
LIGHT_ACCENT = "#0071E3"

PALETTE_DARK = [
    "#0A84FF", "#30D158", "#FF9F0A", "#FF375F", "#BF5AF2",
    "#64D2FF", "#FFD60A", "#FF453A", "#40C8E0", "#AC8E68",
]

PALETTE_LIGHT = [
    "#0071E3", "#34C759", "#FF9500", "#FF2D55", "#AF52DE",
    "#00C7BE", "#5856D6", "#FF3B30", "#A2845E", "#5AC8FA",
]

STYLE_DARK = f"""
QMainWindow, QWidget {{
    background-color: {DARK_BG};
    color: {DARK_TEXT_PRIMARY};
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
    font-size: 12px;
}}
QFrame#sidebar, QFrame#inspector {{
    background-color: {DARK_SURFACE};
    border-radius: 12px;
    border: 1px solid {DARK_BORDER};
}}
QFrame#card {{
    background-color: {DARK_SURFACE_ALT};
    border-radius: 8px;
    border: 1px solid {DARK_BORDER};
}}
QLineEdit, QComboBox, QDoubleSpinBox, QSpinBox {{
    background-color: {DARK_SURFACE_ALT};
    color: {DARK_TEXT_PRIMARY};
    border: 1px solid {DARK_BORDER};
    border-radius: 6px;
    padding: 5px 8px;
    selection-background-color: {DARK_ACCENT};
}}
QLineEdit:focus, QComboBox:focus, QDoubleSpinBox:focus, QSpinBox:focus {{
    border: 1px solid {DARK_ACCENT};
}}
QPushButton {{
    background-color: {DARK_SURFACE_ALT};
    color: {DARK_TEXT_PRIMARY};
    border: 1px solid {DARK_BORDER};
    border-radius: 7px;
    padding: 6px 14px;
    font-weight: 500;
}}
QPushButton:hover {{
    background-color: #383840;
    border-color: #4A4A52;
}}
QPushButton:pressed {{
    background-color: #2D2D34;
}}
QPushButton#primaryBtn {{
    background-color: {DARK_ACCENT};
    color: #FFFFFF;
    border: none;
    font-weight: 600;
}}
QPushButton#primaryBtn:hover {{
    background-color: #2997FF;
}}
QPushButton#accentBtn {{
    background-color: #FF9F0A;
    color: #000000;
    border: none;
    font-weight: 600;
}}
QPushButton#accentBtn:hover {{
    background-color: #FFB340;
}}
QTreeWidget, QListWidget {{
    background-color: {DARK_SURFACE};
    border: 1px solid {DARK_BORDER};
    border-radius: 8px;
    padding: 4px;
    outline: none;
}}
QTreeWidget::item, QListWidget::item {{
    height: 28px;
    border-radius: 5px;
    padding: 2px 6px;
}}
QTreeWidget::item:hover, QListWidget::item:hover {{
    background-color: rgba(255, 255, 255, 0.05);
}}
QTreeWidget::item:selected, QListWidget::item:selected {{
    background-color: {DARK_ACCENT};
    color: #FFFFFF;
}}
QTreeWidget::branch {{
    background: transparent;
}}
QHeaderView::section {{
    background-color: {DARK_SURFACE};
    color: {DARK_TEXT_SECONDARY};
    border: none;
    border-bottom: 1px solid {DARK_BORDER};
    padding: 4px;
    font-weight: 600;
    text-transform: uppercase;
    font-size: 10px;
}}
QCheckBox {{
    color: {DARK_TEXT_PRIMARY};
    spacing: 6px;
}}
QCheckBox::indicator {{
    width: 16px;
    height: 16px;
    border-radius: 4px;
    border: 1px solid {DARK_BORDER};
    background-color: {DARK_SURFACE_ALT};
}}
QCheckBox::indicator:checked {{
    background-color: {DARK_ACCENT};
    border-color: {DARK_ACCENT};
}}
QSplitter::handle {{
    background-color: transparent;
    width: 6px;
    height: 6px;
}}
QScrollBar:vertical {{
    background: transparent;
    width: 8px;
    margin: 0px;
}}
QScrollBar::handle:vertical {{
    background: #44444C;
    min-height: 20px;
    border-radius: 4px;
}}
QScrollBar::handle:vertical:hover {{
    background: #555560;
}}
QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical,
QScrollBar::add-page:vertical, QScrollBar::sub-page:vertical {{
    background: none;
}}
QStatusBar {{
    background-color: {DARK_BG};
    color: {DARK_TEXT_SECONDARY};
    border-top: 1px solid {DARK_BORDER};
    font-size: 11px;
}}
QProgressBar {{
    background-color: {DARK_SURFACE_ALT};
    color: {DARK_TEXT_PRIMARY};
    border: 1px solid {DARK_BORDER};
    border-radius: 6px;
    text-align: center;
    font-size: 10px;
    font-weight: 600;
}}
QProgressBar::chunk {{
    background-color: {DARK_ACCENT};
    border-radius: 5px;
}}
"""

STYLE_LIGHT = f"""
QMainWindow, QWidget {{
    background-color: {LIGHT_BG};
    color: {LIGHT_TEXT_PRIMARY};
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
    font-size: 12px;
}}
QFrame#sidebar, QFrame#inspector {{
    background-color: {LIGHT_SURFACE};
    border-radius: 12px;
    border: 1px solid {LIGHT_BORDER};
}}
QFrame#card {{
    background-color: {LIGHT_SURFACE_ALT};
    border-radius: 8px;
    border: 1px solid {LIGHT_BORDER};
}}
QLineEdit, QComboBox, QDoubleSpinBox, QSpinBox {{
    background-color: {LIGHT_SURFACE};
    color: {LIGHT_TEXT_PRIMARY};
    border: 1px solid {LIGHT_BORDER};
    border-radius: 6px;
    padding: 5px 8px;
    selection-background-color: {LIGHT_ACCENT};
}}
QLineEdit:focus, QComboBox:focus, QDoubleSpinBox:focus, QSpinBox:focus {{
    border: 1px solid {LIGHT_ACCENT};
}}
QPushButton {{
    background-color: {LIGHT_SURFACE};
    color: {LIGHT_TEXT_PRIMARY};
    border: 1px solid {LIGHT_BORDER};
    border-radius: 7px;
    padding: 6px 14px;
    font-weight: 500;
}}
QPushButton:hover {{
    background-color: #F0F0F5;
    border-color: #B8B8C0;
}}
QPushButton:pressed {{
    background-color: #E2E2E8;
}}
QPushButton#primaryBtn {{
    background-color: {LIGHT_ACCENT};
    color: #FFFFFF;
    border: none;
    font-weight: 600;
}}
QPushButton#primaryBtn:hover {{
    background-color: #0077ED;
}}
QPushButton#accentBtn {{
    background-color: #FF9500;
    color: #FFFFFF;
    border: none;
    font-weight: 600;
}}
QPushButton#accentBtn:hover {{
    background-color: #E08500;
}}
QTreeWidget, QListWidget {{
    background-color: {LIGHT_SURFACE};
    border: 1px solid {LIGHT_BORDER};
    border-radius: 8px;
    padding: 4px;
    outline: none;
}}
QTreeWidget::item, QListWidget::item {{
    height: 28px;
    border-radius: 5px;
    padding: 2px 6px;
}}
QTreeWidget::item:hover, QListWidget::item:hover {{
    background-color: rgba(0, 0, 0, 0.04);
}}
QTreeWidget::item:selected, QListWidget::item:selected {{
    background-color: {LIGHT_ACCENT};
    color: #FFFFFF;
}}
QHeaderView::section {{
    background-color: {LIGHT_SURFACE};
    color: {LIGHT_TEXT_SECONDARY};
    border: none;
    border-bottom: 1px solid {LIGHT_BORDER};
    padding: 4px;
    font-weight: 600;
    text-transform: uppercase;
    font-size: 10px;
}}
QCheckBox {{
    color: {LIGHT_TEXT_PRIMARY};
    spacing: 6px;
}}
QCheckBox::indicator {{
    width: 16px;
    height: 16px;
    border-radius: 4px;
    border: 1px solid {LIGHT_BORDER};
    background-color: {LIGHT_SURFACE};
}}
QCheckBox::indicator:checked {{
    background-color: {LIGHT_ACCENT};
    border-color: {LIGHT_ACCENT};
}}
QSplitter::handle {{
    background-color: transparent;
    width: 6px;
    height: 6px;
}}
QScrollBar:vertical {{
    background: transparent;
    width: 8px;
    margin: 0px;
}}
QScrollBar::handle:vertical {{
    background: #C7C7CC;
    min-height: 20px;
    border-radius: 4px;
}}
QScrollBar::handle:vertical:hover {{
    background: #AEAEB2;
}}
QStatusBar {{
    background-color: {LIGHT_BG};
    color: {LIGHT_TEXT_SECONDARY};
    border-top: 1px solid {LIGHT_BORDER};
    font-size: 11px;
}}
QProgressBar {{
    background-color: {LIGHT_SURFACE_ALT};
    color: {LIGHT_TEXT_PRIMARY};
    border: 1px solid {LIGHT_BORDER};
    border-radius: 6px;
    text-align: center;
    font-size: 10px;
    font-weight: 600;
}}
QProgressBar::chunk {{
    background-color: {LIGHT_ACCENT};
    border-radius: 5px;
}}
"""


def extract_signal_only_name(name: str) -> str:
    """
    Extracts strictly the leaf signal name itself, removing all group paths,
    device prefixes, CAN bus IDs, and message names.
    Examples:
      'BMS_Status.PackVoltage' -> 'PackVoltage'
      'CAN1/Sensors/Cell_01_Temp' -> 'Cell_01_Temp'
      'Group_01\\Channel_02\\Battery_Current' -> 'Battery_Current'
      'Engine_0x120.RPM' -> 'RPM'
      '0x100::Inverter_Temp' -> 'Inverter_Temp'
    """
    if not name:
        return ""
    text = name.replace("\\", "/").strip()
    if "/" in text:
        text = text.split("/")[-1].strip()
    if "::" in text:
        text = text.split("::")[-1].strip()
    if "." in text:
        dot_parts = [p.strip() for p in text.split(".") if p.strip()]
        if dot_parts:
            text = dot_parts[-1]
    if ":" in text:
        text = text.split(":")[-1].strip()
    return text or name


def simplify_signal_name(name: str, max_length: int = 48) -> str:
    leaf = extract_signal_only_name(name)
    if len(leaf) <= max_length:
        return leaf
    return f"{leaf[:max_length - 3].rstrip()}..."


def _tokenize_signal_name(name: str) -> set:
    leaf = extract_signal_only_name(name).lower()
    return {token for token in re.split(r"[^a-z0-9]+", leaf) if token}


def match_signals(records: Dict[str, "MDFFileRecord"], min_overlap: float = 0.6
                 ) -> Dict[str, Dict[str, str]]:
    """
    Build a canonical signal map across N tests.
    Returns {canonical_key: {file_id: actual_channel_name}}.
    """
    file_ids = list(records.keys())
    if not file_ids:
        return {}

    base_id = file_ids[0]
    base_signals = records[base_id].signal_names
    matches: Dict[str, Dict[str, str]] = {}

    for base_name in base_signals:
        base_tokens = _tokenize_signal_name(base_name)
        if not base_tokens:
            continue
        canonical_key = extract_signal_only_name(base_name)
        row = {base_id: base_name}

        for other_id in file_ids[1:]:
            other_record = records[other_id]
            best_name, best_score = None, 0.0
            for other_name in other_record.signal_names:
                other_tokens = _tokenize_signal_name(other_name)
                if not other_tokens:
                    continue
                overlap = len(base_tokens & other_tokens) / len(base_tokens | other_tokens)
                if overlap > best_score:
                    best_score, best_name = overlap, other_name
            if best_name is not None and best_score >= min_overlap:
                row[other_id] = best_name

        if len(row) == len(file_ids):
            matches[canonical_key] = row

    return matches


def _signal_stats(x: np.ndarray, y: np.ndarray) -> Dict[str, float]:
    if len(y) == 0:
        return {}
    rms = float(np.sqrt(np.mean(y ** 2)))
    peak = float(np.max(np.abs(y)))
    integrate = np.trapezoid if hasattr(np, "trapezoid") else np.trapz
    energy_ws = float(integrate(y, x)) if len(x) > 1 else 0.0
    return {
        "peak": peak,
        "min": float(np.min(y)),
        "mean": float(np.mean(y)),
        "rms": rms,
        "crest_factor": peak / rms if rms else np.nan,
        "samples": len(y),
        "duration_s": float(x[-1] - x[0]) if len(x) > 1 else 0.0,
        "integral": energy_ws,
    }


def compute_metrics(records: Dict[str, "MDFFileRecord"],
                    signal_map: Dict[str, Dict[str, str]],
                    align_to_zero: bool = True) -> pd.DataFrame:
    """Long-format metrics table: one row per (signal, test)."""
    rows = []
    for canonical_key, per_test_names in signal_map.items():
        for file_id, actual_name in per_test_names.items():
            record = records[file_id]
            data = record.get_signal_data(actual_name)
            if data is None:
                continue
            x, y = data
            if align_to_zero and len(x):
                x = x - x[0]
            valid = np.isfinite(x) & np.isfinite(y)
            x, y = x[valid], y[valid]
            if len(y) == 0:
                continue
            stats = _signal_stats(x, y)
            stats.update({
                "signal": canonical_key,
                "test": file_id,
                "unit": record.signal_units.get(actual_name, ""),
                "channel": actual_name,
            })
            rows.append(stats)
    return pd.DataFrame(rows)


def compute_deltas(metrics: pd.DataFrame, reference_id: str,
                    metric_cols: Tuple[str, ...] = ("peak", "mean", "rms", "integral")
                    ) -> pd.DataFrame:
    """For every non-reference test, compute % delta vs. the reference test."""
    if metrics.empty:
        return pd.DataFrame()

    reference = metrics[metrics["test"] == reference_id].set_index("signal")
    others = metrics[metrics["test"] != reference_id]

    rows = []
    for _, row in others.iterrows():
        signal = row["signal"]
        if signal not in reference.index:
            continue
        ref_row = reference.loc[signal]
        record = {"signal": signal, "test": row["test"], "reference": reference_id,
                  "unit": row["unit"]}
        max_abs_pct = 0.0
        for col in metric_cols:
            ref_val = ref_row.get(col, np.nan)
            test_val = row.get(col, np.nan)
            if pd.isna(ref_val) or pd.isna(test_val) or ref_val == 0:
                pct = np.nan
            else:
                pct = 100.0 * (test_val - ref_val) / abs(ref_val)
                max_abs_pct = max(max_abs_pct, abs(pct))
            record[f"{col}_ref"] = ref_val
            record[f"{col}_test"] = test_val
            record[f"{col}_delta_pct"] = pct
        record["max_abs_delta_pct"] = max_abs_pct
        rows.append(record)

    delta_df = pd.DataFrame(rows)
    if not delta_df.empty:
        delta_df = delta_df.sort_values("max_abs_delta_pct", ascending=False).reset_index(drop=True)
    return delta_df


FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name=FONT, size=10)
FLAG_FILL = PatternFill("solid", fgColor="FFCDD2")
WARN_FILL = PatternFill("solid", fgColor="FFF3CD")


def _write_table(ws, df: pd.DataFrame, start_row: int = 1):
    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = XlsxAlignment(horizontal="center")
    for row_offset, (_, row) in enumerate(df.iterrows(), start=1):
        for col_idx, col_name in enumerate(df.columns, start=1):
            value = row[col_name]
            if isinstance(value, float) and np.isfinite(value):
                value = round(value, 6)
            elif isinstance(value, float):
                value = None
            ws.cell(row=start_row + row_offset, column=col_idx, value=value).font = BODY_FONT
    for col_idx, col_name in enumerate(df.columns, start=1):
        width = max(12, min(28, len(str(col_name)) + 4))
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    return start_row + len(df)


def _flag_delta_rows(ws, header_row: int, n_rows: int, delta_col_idx: int,
                      warn_pct: float = 5.0, flag_pct: float = 15.0):
    for r in range(header_row + 1, header_row + 1 + n_rows):
        cell = ws.cell(row=r, column=delta_col_idx)
        if cell.value is None:
            continue
        abs_val = abs(cell.value)
        if abs_val >= flag_pct:
            for c in range(1, ws.max_column + 1):
                ws.cell(row=r, column=c).fill = FLAG_FILL
        elif abs_val >= warn_pct:
            for c in range(1, ws.max_column + 1):
                ws.cell(row=r, column=c).fill = WARN_FILL


def generate_comparison_report(
    records: Dict[str, "MDFFileRecord"],
    output_path: str,
    reference_id: Optional[str] = None,
    top_n_charts: int = 8,
    warn_pct: float = 5.0,
    flag_pct: float = 15.0,
    align_to_zero: bool = True,
) -> Dict[str, Any]:
    """
    records: {file_id: MDFFileRecord} — 2 to 3 loaded tests to compare.
    output_path: destination .xlsx path.
    reference_id: baseline test id; defaults to the first key in records.
    Returns a summary dict (counts, flagged signal names) for a status message.
    """
    if len(records) < 2:
        raise ValueError("Need at least two tests to compare.")
    reference_id = reference_id or next(iter(records))
    if reference_id not in records:
        raise ValueError(f"reference_id '{reference_id}' not in records")

    signal_map = match_signals(records)
    if not signal_map:
        raise ValueError("No signals matched across all tests — check channel naming.")

    metrics = compute_metrics(records, signal_map, align_to_zero=align_to_zero)
    deltas = compute_deltas(metrics, reference_id=reference_id)

    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.cell(row=1, column=1, value=f"Test comparison — reference: {reference_id}").font = \
        Font(name=FONT, bold=True, size=12)
    ws_summary.cell(row=2, column=1,
                     value=f"Tests: {', '.join(records.keys())}  |  "
                           f"Matched signals: {len(signal_map)}").font = Font(name=FONT, italic=True, size=9)
    summary_cols = ["signal", "test", "unit", "peak", "min", "mean", "rms",
                    "crest_factor", "integral", "duration_s", "samples"]
    summary_view = metrics[summary_cols].sort_values(["signal", "test"])
    _write_table(ws_summary, summary_view, start_row=4)

    ws_delta = wb.create_sheet("Deltas")
    delta_cols = ["signal", "test", "reference", "unit",
                  "peak_ref", "peak_test", "peak_delta_pct",
                  "mean_ref", "mean_test", "mean_delta_pct",
                  "rms_ref", "rms_test", "rms_delta_pct",
                  "integral_ref", "integral_test", "integral_delta_pct",
                  "max_abs_delta_pct"]
    delta_view = deltas[[c for c in delta_cols if c in deltas.columns]]
    header_row = 2
    ws_delta.cell(row=1, column=1,
                  value="Sorted by worst |delta| first — red ≥ flag threshold, amber ≥ warn threshold").font = \
        Font(name=FONT, italic=True, size=9)
    _write_table(ws_delta, delta_view, start_row=header_row)
    max_delta_col = delta_view.columns.get_loc("max_abs_delta_pct") + 1
    _flag_delta_rows(ws_delta, header_row, len(delta_view), max_delta_col, warn_pct, flag_pct)

    ws_charts = wb.create_sheet("Top Deltas - Charts")
    flagged_signals = deltas.drop_duplicates("signal").nlargest(top_n_charts, "max_abs_delta_pct")
    chart_row = 1
    hidden_col_start = 40

    for _, flagged in flagged_signals.iterrows():
        signal_key = flagged["signal"]
        per_test_names = signal_map[signal_key]

        ref_name = per_test_names[reference_id]
        ref_data = records[reference_id].get_signal_data(ref_name)
        if ref_data is None:
            continue
        ref_x, _ = ref_data
        if align_to_zero and len(ref_x):
            ref_x = ref_x - ref_x[0]
        t_grid = np.linspace(ref_x[0], ref_x[-1], min(500, len(ref_x))) if len(ref_x) else np.array([])
        if len(t_grid) < 2:
            continue

        data_block = {"time_s": t_grid}
        for file_id, actual_name in per_test_names.items():
            data = records[file_id].get_signal_data(actual_name)
            if data is None:
                continue
            x, y = data
            if align_to_zero and len(x):
                x = x - x[0]
            data_block[file_id] = np.interp(t_grid, x, y)

        block_df = pd.DataFrame(data_block)
        start_col = hidden_col_start
        for col_idx, col_name in enumerate(block_df.columns, start=start_col):
            ws_charts.cell(row=chart_row, column=col_idx, value=col_name)
            for row_idx, val in enumerate(block_df[col_name], start=chart_row + 1):
                ws_charts.cell(row=row_idx, column=col_idx, value=round(float(val), 6))

        chart = LineChart()
        chart.title = f"{signal_key}  (Δmax {flagged['max_abs_delta_pct']:.1f}%)"
        chart.y_axis.title = flagged.get("unit", "")
        chart.x_axis.title = "Time (s)"
        chart.height, chart.width = 8, 16
        n_rows = len(block_df)
        cats = Reference(ws_charts, min_col=start_col, min_row=chart_row + 1,
                          max_row=chart_row + n_rows)
        for offset, test_id in enumerate(per_test_names.keys(), start=1):
            data_ref = Reference(ws_charts, min_col=start_col + offset,
                                  min_row=chart_row, max_row=chart_row + n_rows)
            chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats)

        chart_index = list(flagged_signals.index).index(flagged.name)
        anchor = f"A{1 + chart_index * 17}"
        ws_charts.add_chart(chart, anchor)
        chart_row += n_rows + 3

    for col_idx in range(hidden_col_start, hidden_col_start + 6):
        ws_charts.column_dimensions[get_column_letter(col_idx)].hidden = True

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    return {
        "output_path": output_path,
        "tests_compared": list(records.keys()),
        "reference": reference_id,
        "signals_matched": len(signal_map),
        "signals_flagged": int((deltas.drop_duplicates("signal")["max_abs_delta_pct"] >= warn_pct).sum())
                            if not deltas.empty else 0,
        "top_flagged_signals": flagged_signals["signal"].tolist() if not flagged_signals.empty else [],
    }


CHN_CHANNEL_LINE = re.compile(
    r"^Name of channel\s+(?P<number>\d+)\s+:\s*"
    r"(?P<code>\S+)\s*/\s*(?P<description>.*?)\s*$",
    re.IGNORECASE,
)
CHN_EXCEL_MAPPING = Path(
    r"C:\Users\phupt6\OneDrive - TAP DOAN VINGROUP\Python\Mf4_reader\CHANNEL\Engine Mount VF9 Analysis.xlsx"
)


def read_chn_metadata(path: Path) -> Tuple[str, Dict[int, Tuple[str, str]]]:
    standard = ""
    channels: Dict[int, Tuple[str, str]] = {}
    block: Dict[str, str] = {}
    block_number = 0

    def flush_block():
        nonlocal block_number
        if not block:
            return
        code = (block.get("channel_code") or block.get("laboratory_channel_code") or
                block.get("customer_channel_code") or "").strip()
        name = (block.get("name_of_the_channel") or block.get("channel_name") or "").strip()
        if code or name:
            block_number += 1
            channels[block_number] = (code or name, name or code)
        block.clear()

    with path.open("r", encoding="ascii", errors="replace") as handle:
        for raw_line in handle:
            line = raw_line.strip()
            if not line:
                flush_block()
                continue
            if line.lower().startswith("instrumentation standard"):
                standard = line.split(":", 1)[-1].strip()
            match = CHN_CHANNEL_LINE.match(line)
            if match:
                channels[int(match.group("number"))] = (
                    match.group("code"),
                    match.group("description").strip(),
                )
                continue
            if ":" in line:
                label, value = line.split(":", 1)
                normalized = _chn_mapping_key(label)
                if normalized in {
                    "name_of_the_channel", "channel_name", "channel_code",
                    "laboratory_channel_code", "customer_channel_code",
                }:
                    block[normalized] = value.strip()
    flush_block()
    if not channels:
        raise ValueError(f"No channel definitions found in {path}")
    return standard, channels


def _compact_chn_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value).lower())


def _chn_direction(*values: Any) -> str:
    text = " ".join(str(value) for value in values if value).upper()
    match = re.search(r"(?:^|[^A-Z])([XYZ])(?:$|[^A-Z])", text)
    if match:
        return match.group(1)
    match = re.search(r"(?:AC|DC)?([XYZ])(?:P|N)?$", text)
    return match.group(1) if match else ""


def _chn_location(*values: Any) -> str:
    text = " ".join(str(value).strip() for value in values if value).strip()
    text = re.sub(r"\s+(?:X|Y|Z)\s*$", "", text, flags=re.IGNORECASE)
    return text


def read_chn_channel_metadata(path: Path) -> Dict[str, Dict[str, str]]:
    """Read CHN channel definitions and expose aliases for tabular mapping."""
    standard, definitions = read_chn_metadata(path)
    metadata: Dict[str, Dict[str, str]] = {}
    column_metadata: Dict[Any, Dict[str, str]] = {}
    for number, (code, description) in definitions.items():
        record = {
            "channel_code": code,
            "channel_name": description or code,
            "laboratory_channel_code": code,
            "customer_channel_code": code,
            "unit": "",
            "direction": _chn_direction(code, description),
            "location": _chn_location(description),
            "instrumentation_standard": standard,
        }
        metadata[_compact_chn_key(code)] = record

        sidecar = path.with_suffix(f".{number:03d}")
        if not sidecar.is_file():
            sidecar = path.with_suffix(f".{number}")
        if sidecar.is_file():
            file_metadata = read_chn_file_metadata(sidecar)
            for key, value in file_metadata.items():
                record[key] = value
            record["channel_name"] = (
                file_metadata.get("name_of_the_channel")
                or record["channel_name"]
            )
            record["channel_code"] = (
                file_metadata.get("channel_code")
                or record["channel_code"]
            )
            record["laboratory_channel_code"] = (
                file_metadata.get("laboratory_channel_code")
                or record["laboratory_channel_code"]
            )
            record["customer_channel_code"] = (
                file_metadata.get("customer_channel_code")
                or record["customer_channel_code"]
            )
            record["unit"] = file_metadata.get("unit", record["unit"])
            record["direction"] = _chn_direction(
                record["direction"], record["channel_name"], record["channel_code"]
            )
            record["location"] = _chn_location(
                record["location"], record["channel_name"]
            )
        for alias in (
            record.get("channel_code"),
            record.get("laboratory_channel_code"),
            record.get("customer_channel_code"),
        ):
            if alias:
                metadata[_compact_chn_key(alias)] = record
    return metadata


def read_chn_file_metadata(path: Path) -> Dict[str, str]:
    metadata: Dict[str, str] = {}
    with path.open("r", encoding="ascii", errors="strict") as handle:
        for raw_line in handle:
            line = raw_line.strip()
            if not line:
                continue
            if ":" not in line:
                break
            key, value = line.split(":", 1)
            metadata[key.strip().lower().replace(" ", "_")] = value.strip()
    return metadata


def read_chn_values(path: Path) -> np.ndarray:
    values: List[float] = []
    started = False
    with path.open("r", encoding="ascii", errors="strict") as handle:
        for line_number, raw_line in enumerate(handle, start=1):
            value = raw_line.strip()
            if not value:
                continue
            if not started and ":" in value:
                continue
            try:
                values.append(float(value))
            except ValueError as exc:
                raise ValueError(
                    f"Invalid numeric value in {path.name}:{line_number}: {value!r}"
                ) from exc
            started = True
    return np.asarray(values, dtype=float)


def _chn_mapping_key(value: Any) -> str:
    return str(value).strip().lower().replace(" ", "_").rstrip(":")


def read_chn_excel_mapping(path: Path) -> Dict[str, Dict[str, str]]:
    """Read channel descriptions from the Excel workbook, keyed by channel code."""
    if not path.is_file():
        return {}
    try:
        sheets = pd.read_excel(path, sheet_name=None, header=None, engine="openpyxl")
    except Exception as exc:
        logging.warning("Pandas cannot read CHN Excel mapping %s: %s", path, exc)
        try:
            import win32com.client

            excel = win32com.client.DispatchEx("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False
            workbook = excel.Workbooks.Open(str(path.resolve()), None, True)
            sheets = {}
            for worksheet in workbook.Worksheets:
                used_range = worksheet.UsedRange
                raw_values = used_range.Value
                if raw_values is None:
                    continue
                if not isinstance(raw_values, tuple):
                    raw_values = ((raw_values,),)
                elif raw_values and not isinstance(raw_values[0], tuple):
                    raw_values = (raw_values,)
                sheets[worksheet.Name] = pd.DataFrame(raw_values)
            workbook.Close(False)
            excel.Quit()
        except Exception as com_exc:
            logging.warning("Excel COM fallback unavailable for %s: %s", path, com_exc)
            return {}

    mapping: Dict[str, Dict[str, str]] = {}
    for sheet_name, frame in sheets.items():
        for column_index in range(frame.shape[1]):
            values: Dict[str, str] = {}
            for raw_cell in frame.iloc[:, column_index].tolist():
                if pd.isna(raw_cell):
                    continue
                cell = str(raw_cell).strip()
                if ":" not in cell:
                    continue
                label, value = cell.split(":", 1)
                values[_chn_mapping_key(label)] = value.strip()
            code = values.get("channel_code") or values.get("customer_channel_code")
            if code:
                mapping[code.lower()] = {
                    "name": values.get("name_of_the_channel", ""),
                    "group": str(sheet_name).strip(),
                    "unit": values.get("unit", ""),
                    "sampling_interval": values.get("sampling_interval", ""),
                    "time_of_first_sample": values.get("time_of_first_sample", ""),
                }
    logging.info("Loaded %d CHN channel mappings from %s", len(mapping), path.name)
    return mapping


def read_channel_metadata_csv(path: Path) -> Dict[str, Dict[str, str]]:
    """Read a metadata CSV whose columns are sensors and rows are key/value fields."""
    frame = pd.read_csv(path, header=None, dtype=str, keep_default_na=False)
    if frame.empty or not any(":" in str(value) for value in frame.iloc[0].tolist()):
        return {}
    metadata: Dict[str, Dict[str, str]] = {}
    column_metadata: Dict[Any, Dict[str, str]] = {}
    for column in frame.columns:
        values: Dict[str, str] = {}
        for cell in frame[column].tolist():
            text = str(cell).strip()
            if ":" not in text:
                continue
            label, value = text.split(":", 1)
            values[_chn_mapping_key(label)] = value.strip()
        code = (values.get("channel_code") or values.get("laboratory_channel_code") or
                values.get("customer_channel_code"))
        name = values.get("name_of_the_channel") or values.get("channel_name")
        if not code and not name:
            continue
        record = {
            **values,
            "channel_code": code or name,
            "channel_name": name or code,
            "laboratory_channel_code": values.get("laboratory_channel_code", code or name),
            "customer_channel_code": values.get("customer_channel_code", code or name),
            "direction": values.get("direction", "") or _chn_direction(code, name),
            "location": _chn_location(values.get(".group"), name),
        }
        for alias in (record["channel_code"], record["laboratory_channel_code"], record["customer_channel_code"]):
            metadata[_compact_chn_key(alias)] = record
        column_metadata[column] = record
    return metadata


def read_measurement_table(path: Path, metadata_rows: int = 43, sheet_name: Any = 0, raw_data: Optional[pd.DataFrame] = None) -> Tuple[pd.DataFrame, Dict[str, Dict[str, str]]]:
    """Read NVH tables with channel metadata in rows 1-43 and samples from row 44."""
    suffix = path.suffix.lower()
    if raw_data is not None:
        raw = raw_data
    elif suffix in {".xlsx", ".xls"}:
        sheets = _read_excel_sheets_resilient(path)
        if sheet_name == "auto":
            for candidate in sheets:
                candidate_frame, candidate_metadata = read_measurement_table(path, metadata_rows, candidate)
                if candidate_metadata and not candidate_frame.empty:
                    return candidate_frame, candidate_metadata
            sheet_name = next(iter(sheets), 0)
        raw = sheets[sheet_name] if sheet_name in sheets else next(iter(sheets.values()), pd.DataFrame())
    elif suffix == ".txt":
        raw = pd.read_csv(path, header=None, dtype=str, sep=None, engine="python")
    else:
        raw = pd.read_csv(path, header=None, dtype=str, keep_default_na=False)
    if raw.empty or raw.shape[0] <= metadata_rows:
        return pd.DataFrame(), {}

    metadata: Dict[str, Dict[str, str]] = {}
    column_metadata: Dict[Any, Dict[str, str]] = {}
    for column in raw.columns:
        values: Dict[str, str] = {}
        for cell in raw.iloc[:metadata_rows, column].tolist():
            text = str(cell).strip()
            if ":" not in text:
                continue
            label, value = text.split(":", 1)
            values[_chn_mapping_key(label)] = value.strip()
        code = (values.get("channel_code") or values.get("laboratory_channel_code") or
                values.get("customer_channel_code"))
        name = values.get("name_of_the_channel") or values.get("channel_name")
        if not code and not name:
            continue
        record = {
            **values,
            "channel_code": code or name,
            "channel_name": name or code,
            "laboratory_channel_code": values.get("laboratory_channel_code", code or name),
            "customer_channel_code": values.get("customer_channel_code", code or name),
            "direction": values.get("direction", "") or _chn_direction(code, name),
            "location": _chn_location(values.get(".group"), name),
        }
        for alias in (record["channel_code"], record["laboratory_channel_code"], record["customer_channel_code"]):
            metadata[_compact_chn_key(alias)] = record
        column_metadata[column] = record

    if not metadata:
        return pd.DataFrame(), {}

    sample_frame = raw.iloc[metadata_rows:].copy()
    selected_columns = []
    column_names = []
    for column in sample_frame.columns:
        metadata_for_column = column_metadata.get(column)
        column_values = pd.to_numeric(sample_frame[column], errors="coerce")
        if not column_values.notna().any():
            continue
        selected_columns.append(column_values.reset_index(drop=True))
        column_names.append(metadata_for_column.get("channel_code", str(column)) if metadata_for_column else str(column))
    if not selected_columns:
        return pd.DataFrame(), {}
    measurements = pd.concat(selected_columns, axis=1)
    measurements.columns = column_names
    return measurements, metadata


def read_excel_all_sheets_fast(path: Path) -> Dict[str, pd.DataFrame]:
    """Read every Excel worksheet once, preferring the Rust calamine engine."""
    try:
        return pd.read_excel(path, sheet_name=None, header=None, dtype=str, engine="calamine")
    except Exception:
        return _read_excel_sheets_resilient(path)


def read_two_row_excel_table(raw: pd.DataFrame) -> Optional[Tuple[pd.DataFrame, np.ndarray, Dict[str, str]]]:
    """Parse a worksheet whose first row is names and second row is units.

    The first column must be a time column. This also prevents workbook
    information/event sheets from being treated as measurements.
    """
    if raw.empty or raw.shape[0] < 3 or raw.shape[1] < 2:
        return None

    names = ["" if pd.isna(value) else str(value).strip() for value in raw.iloc[0]]
    units = ["" if pd.isna(value) else str(value).strip() for value in raw.iloc[1]]
    if names[0].lower() not in {"time", "t", "timestamp", "timestamps"}:
        return None

    time_values = pd.to_numeric(raw.iloc[2:, 0], errors="coerce")
    valid_time = time_values.notna()
    if valid_time.sum() < 2:
        return None

    columns: List[pd.Series] = []
    column_names: List[str] = []
    signal_units: Dict[str, str] = {}
    used_names: Dict[str, int] = {}
    for position in range(1, raw.shape[1]):
        base_name = names[position]
        if not base_name:
            continue
        values = pd.to_numeric(raw.iloc[2:, position], errors="coerce")
        if not values[valid_time].notna().any():
            continue
        occurrence = used_names.get(base_name, 0) + 1
        used_names[base_name] = occurrence
        name = base_name if occurrence == 1 else f"{base_name} ({occurrence})"
        columns.append(values[valid_time].reset_index(drop=True))
        column_names.append(name)
        signal_units[name] = units[position]

    if not columns:
        return None
    measurements = pd.concat(columns, axis=1)
    measurements.columns = column_names
    timestamps = time_values[valid_time].to_numpy(dtype=float)
    return measurements, timestamps, signal_units


def _read_excel_sheets_resilient(path: Path) -> Dict[Any, pd.DataFrame]:
    """Read Excel sheets even when OneDrive/Excel has the source file locked."""
    try:
        return pd.read_excel(path, sheet_name=None, header=None, dtype=str)
    except (PermissionError, OSError) as first_error:
        try:
            with path.open("rb") as handle:
                workbook_bytes = BytesIO(handle.read())
            return pd.read_excel(workbook_bytes, sheet_name=None, header=None, dtype=str)
        except Exception:
            pass

        try:
            import win32com.client
            excel = win32com.client.DispatchEx("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False
            workbook = excel.Workbooks.Open(str(path.resolve()), ReadOnly=True, UpdateLinks=0)
            sheets = {}
            for worksheet in workbook.Worksheets:
                used_range = worksheet.UsedRange
                values = used_range.Value
                if values is None:
                    continue
                if not isinstance(values, tuple):
                    values = ((values,),)
                elif values and not isinstance(values[0], tuple):
                    values = (values,)
                sheets[worksheet.Name] = pd.DataFrame(values)
            workbook.Close(False)
            excel.Quit()
            if sheets:
                return sheets
        except Exception as com_error:
            try:
                workbook.Close(False)
            except Exception:
                pass
            try:
                excel.Quit()
            except Exception:
                pass
            raise PermissionError(
                f"Cannot read Excel file '{path.name}'. Close it in Excel and make sure the OneDrive file is available offline."
            ) from com_error
        raise PermissionError(
            f"Cannot read Excel file '{path.name}'. Close it in Excel and make sure the OneDrive file is available offline."
        ) from first_error


class MDFFileRecord:
    def __init__(self, file_id: str, path: str, color_tag: str, dbc_paths: Optional[List[str]] = None, progress_callback=None, sheet_name: Any = "auto", raw_data: Optional[pd.DataFrame] = None):
        self.file_id = file_id
        self.path = path
        self.filename = Path(path).name
        self.dbc_paths = dbc_paths or []
        self.progress_callback = progress_callback
        self.sheet_name = sheet_name
        self.raw_data = raw_data
        self.color_tag = color_tag
        self.mdf = None
        self.csv_data: Optional[pd.DataFrame] = None
        self.blf_data: Optional[Dict[str, Tuple[np.ndarray, np.ndarray]]] = None
        self.blf_dbc_names: List[str] = []
        self.signal_names: List[str] = []
        self.signal_units: Dict[str, str] = {}
        self.signal_descriptions: Dict[str, str] = {}
        self.signal_display_names: Dict[str, str] = {}
        self.signal_metadata: Dict[str, Dict[str, str]] = {}
        self.chn_metadata: Dict[str, Dict[str, str]] = {}
        self.chn_data: Optional[Dict[str, Tuple[np.ndarray, np.ndarray]]] = None
        self.signal_cache: Dict[str, Tuple[np.ndarray, np.ndarray]] = {}
        self.t_start: float = 0.0
        self.t_end: float = 0.0
        suffix = Path(path).suffix.lower()
        self.source_type = "BLF" if suffix == ".blf" else (
            "MF4" if suffix in {".mf4", ".mdf", ".dat"} else suffix.lstrip(".").upper()
        )
        if suffix in {".csv", ".xlsx", ".xls", ".txt"}:
            self._report_progress(f"Reading tabular data in {self.filename}...", 0.15)
            self._load_tabular_file(path, sheet_name=sheet_name, raw_data=raw_data)
            self._report_progress(f"Finished loading {self.filename}", 1.0)
        elif suffix == ".chn":
            self._report_progress(f"Reading CHN channels in {self.filename}...", 0.05)
            self._load_chn(path)
            self._report_progress(f"Finished loading {self.filename}", 1.0)
        elif suffix == ".blf":
            self._report_progress(f"Scanning and decoding {self.filename}...", 0.05)
            self._load_blf(path, self.dbc_paths)
            self._report_progress(f"Finished decoding {self.filename}", 1.0)
        else:
            self._report_progress(f"Opening MDF file {self.filename}...", 0.2)
            self.mdf = MDF(path)
            self._report_progress(f"Inspecting channels in {self.filename}...", 0.75)
            self._inspect_channels()
            self._report_progress(f"Finished loading {self.filename}", 1.0)

    def _report_progress(self, msg: str, frac: float = 0.5):
        if self.progress_callback is not None:
            try:
                self.progress_callback(msg, frac)
            except TypeError:
                self.progress_callback(msg)

    def _load_tabular_file(self, path: str, sheet_name: Any = "auto", raw_data: Optional[pd.DataFrame] = None):
        suffix = Path(path).suffix.lower()
        if suffix in {".xlsx", ".xls"}:
            excel_raw = raw_data
            if excel_raw is None:
                sheets = _read_excel_sheets_resilient(Path(path))
                excel_raw = sheets.get(sheet_name, next(iter(sheets.values()), pd.DataFrame()))
            parsed = read_two_row_excel_table(excel_raw)
            if parsed is not None:
                self.csv_data, self._csv_timestamps, self.signal_units = parsed
                self.signal_names = list(self.csv_data.columns)
                self.signal_metadata = {
                    name: {
                        "channel_code": name,
                        "channel_name": name,
                        "unit": self.signal_units.get(name, ""),
                        "sheet_name": str(sheet_name),
                    }
                    for name in self.signal_names
                }
                self.signal_display_names = {name: name for name in self.signal_names}
                self._apply_chn_metadata_to_tabular(path)
                self._update_time_bounds()
                return
            if raw_data is not None:
                raise ValueError(
                    "not a measurement sheet (expected row 1 names, row 2 units, and a Time column)"
                )

        metadata_measurements, metadata = read_measurement_table(
            Path(path), metadata_rows=43, sheet_name=sheet_name, raw_data=raw_data
        )
        if metadata and not metadata_measurements.empty:
            self.csv_data = metadata_measurements
            self.signal_names = [str(column) for column in metadata_measurements.columns]
            self.signal_units = {name: "" for name in self.signal_names}
            self._csv_timestamps = np.arange(len(metadata_measurements), dtype=float)
            self.apply_chn_metadata(metadata)
            interval_values = [
                float(info["sampling_interval"])
                for info in metadata.values()
                if info.get("sampling_interval")
                and _compact_chn_key(info.get("channel_code", "")) in metadata
            ]
            first_sample_values = [
                float(info["time_of_first_sample"])
                for info in metadata.values()
                if info.get("time_of_first_sample")
                and _compact_chn_key(info.get("channel_code", "")) in metadata
            ]
            if interval_values:
                interval = float(np.median(interval_values))
                first_sample = float(np.median(first_sample_values)) if first_sample_values else 0.0
                self._csv_timestamps = first_sample + np.arange(len(metadata_measurements), dtype=float) * interval
            return

        if raw_data is not None:
            data = raw_data
        elif Path(path).suffix.lower() in {".xlsx", ".xls"}:
            sheets = _read_excel_sheets_resilient(Path(path))
            data = sheets.get(sheet_name, next(iter(sheets.values()), pd.DataFrame()))
        elif Path(path).suffix.lower() == ".txt":
            data = pd.read_csv(path, sep=None, engine="python")
        else:
            data = pd.read_csv(path)
        if data.empty or len(data.columns) == 0:
            raise ValueError("Tabular file does not contain any data columns")

        data.columns = [str(column).strip() for column in data.columns]
        if any(not column for column in data.columns):
            raise ValueError("Tabular file contains an empty column name")

        time_column = data.columns[0]
        time_values = pd.to_numeric(data[time_column], errors="coerce")
        time_name = time_column.lower().strip()
        use_first_as_time = (
            len(data.columns) > 1
            and (
                time_name in {"time", "t", "timestamp", "timestamps"}
                or (time_values.notna().all() and time_values.is_monotonic_increasing)
            )
        )
        if use_first_as_time:
            self.csv_data = data.iloc[:, 1:].copy()
            self._csv_timestamps = time_values.to_numpy(dtype=float)
        else:
            self.csv_data = data.copy()
            self._csv_timestamps = np.arange(len(data), dtype=float)

        numeric_columns = []
        for name in self.csv_data.columns:
            values = pd.to_numeric(self.csv_data[name], errors="coerce")
            if values.notna().any():
                self.csv_data[name] = values
                numeric_columns.append(name)
        self.signal_names = numeric_columns
        self.signal_units = {name: "" for name in self.signal_names}
        self._apply_chn_metadata_to_tabular(path)
        if not self.signal_names:
            raise ValueError("Tabular file does not contain numeric signal columns")
        self._update_time_bounds()

    def apply_chn_metadata(self, metadata: Dict[str, Dict[str, str]]):
        self.chn_metadata = metadata
        for column in self.signal_names:
            normalized_column = _compact_chn_key(column)
            channel_metadata = metadata.get(normalized_column)
            if channel_metadata is None:
                channel_metadata = next(
                    (candidate for key, candidate in metadata.items()
                     if key and (key in normalized_column or normalized_column in key)),
                    None,
                )
            if channel_metadata is None:
                continue
            self.signal_metadata[column] = dict(channel_metadata)
            display_name = channel_metadata.get("channel_name") or column
            direction = channel_metadata.get("direction", "")
            self.signal_display_names[column] = f"{display_name} [{direction}]" if direction else display_name
            self.signal_descriptions[column] = channel_metadata.get("location", "") or display_name
            self.signal_units[column] = channel_metadata.get("unit", "") or self.signal_units.get(column, "")
            if channel_metadata.get("sampling_interval") and not channel_metadata.get("sampling_rate"):
                try:
                    channel_metadata["sampling_rate"] = str(1.0 / float(channel_metadata["sampling_interval"]))
                except (TypeError, ValueError, ZeroDivisionError):
                    pass

    def _apply_chn_metadata_to_tabular(self, path: str):
        source = Path(path)
        candidates = [source.with_suffix(".chn"), source.with_suffix(".CHN")]
        if not any(candidate.is_file() for candidate in candidates):
            candidates = sorted(source.parent.glob("*.chn")) + sorted(source.parent.glob("*.CHN"))
        chn_path = next((candidate for candidate in candidates if candidate.is_file()), None)
        if chn_path is None:
            return
        try:
            self.chn_metadata = read_chn_channel_metadata(chn_path)
        except Exception as exc:
            logging.warning("Unable to map CHN metadata for %s: %s", source.name, exc)
            return
        self.apply_chn_metadata(self.chn_metadata)

    def _load_chn(self, path: str):
        metadata_path = Path(path)
        standard, definitions = read_chn_metadata(metadata_path)
        self.chn_metadata = read_chn_channel_metadata(metadata_path)
        values = read_chn_values(metadata_path)
        code, description = next(iter(definitions.values()))
        name = description or code
        interval = 1.0
        metadata = self.chn_metadata.get(_compact_chn_key(code), {})
        try:
            interval = float(metadata.get("sampling_interval", interval))
        except (TypeError, ValueError):
            pass
        timestamps = np.arange(len(values), dtype=float) * interval
        self.chn_data = {name: (timestamps, values)}
        self.signal_names = [name]
        self.signal_units[name] = str(metadata.get("unit", ""))
        self.chn_standard = standard
        self._update_time_bounds()

    def _update_time_bounds(self):
        bounds = []
        for name in self.signal_names:
            data = self.get_signal_data(name)
            if data is not None and len(data[0]):
                bounds.append((float(data[0][0]), float(data[0][-1])))
        if bounds:
            self.t_start = min(start for start, _ in bounds)
            self.t_end = max(end for _, end in bounds)

    def _load_blf(self, path: str, dbc_paths: List[str]):
        """
        Reads Vector BLF files and decodes CAN frames into time-series signals
        using uploaded DBC files. Handles 11-bit and 29-bit CAN IDs, multi-channel logs,
        and unit extraction.
        """
        can = ensure_python_can()
        dbc_files = [Path(dbc_path) for dbc_path in dbc_paths if Path(dbc_path).is_file()]
        if not dbc_files:
            raise FileNotFoundError(
                "A DBC file is required to decode signals from BLF.\n"
                "Please click '+ Add DBC' or drag & drop a .dbc file first."
            )

        # Load all DBC files robustly
        loaded_dbs = []
        database_by_id = {}
        database_by_masked_id = {}
        for dbc_file in dbc_files:
            try:
                cache_key = (os.path.normcase(str(dbc_file.resolve())), dbc_file.stat().st_mtime_ns)
                database = _DBC_DATABASE_CACHE.get(cache_key)
                if database is None:
                    database = load_dbc_robust(str(dbc_file))
                    _DBC_DATABASE_CACHE[cache_key] = database
                loaded_dbs.append((database, dbc_file.name))
                for message in database.messages:
                    definition = (message, dbc_file.name)
                    database_by_id.setdefault(int(message.frame_id), definition)
                    database_by_masked_id.setdefault(
                        int(message.frame_id) & 0x1FFFFFFF, definition
                    )
                    database_by_masked_id.setdefault(
                        int(message.frame_id) & 0x7FF, definition
                    )
            except Exception as exc:
                print(f"[Warning] Failed loading DBC {dbc_file.name}: {exc}")
                continue

        if not loaded_dbs:
            raise ValueError("None of the specified DBC files could be loaded successfully.")

        signals: Dict[str, Tuple[List[float], List[float]]] = {}
        decoded_messages = 0
        frame_count = 0
        # Cache signal units from messages
        unit_by_key: Dict[str, str] = {}
        for db, _ in loaded_dbs:
            for msg in db.messages:
                for sig in msg.signals:
                    unit_str = str(sig.unit or "").strip()
                    unit_by_key[f"{msg.name}.{sig.name}"] = unit_str
                    if sig.name not in unit_by_key:
                        unit_by_key[sig.name] = unit_str

        file_size = os.path.getsize(path) if os.path.exists(path) else 1
        with can.BLFReader(path) as reader:
            for message in reader:
                frame_count += 1
                if frame_count % 20_000 == 0:
                    pos = 0
                    try:
                        f_obj = getattr(reader, "_file", getattr(reader, "file", None))
                        if f_obj and hasattr(f_obj, "tell"):
                            pos = f_obj.tell()
                    except Exception:
                        pass
                    if pos > 0 and file_size > 0:
                        fraction = min(0.98, max(0.05, pos / file_size))
                    else:
                        fraction = min(0.95, max(0.05, frame_count / (frame_count + 80_000)))
                    self._report_progress(
                        f"Decoding {self.filename}: {frame_count:,} frames, {decoded_messages:,} signals",
                        fraction
                    )

                if not hasattr(message, "arbitration_id") or not hasattr(message, "data"):
                    continue

                arb_id = int(message.arbitration_id)
                timestamp = float(message.timestamp)

                # Match CAN Message Definition
                definition = database_by_id.get(arb_id)
                if definition is None:
                    definition = database_by_masked_id.get(arb_id & 0x1FFFFFFF)
                if definition is None:
                    definition = database_by_masked_id.get(arb_id & 0x7FF)
                if definition is None and arb_id not in database_by_id:
                    for db, db_name in loaded_dbs:
                        try:
                            msg_def = db.get_message_by_frame_id(arb_id & 0x1FFFFFFF)
                            if msg_def is not None:
                                definition = (msg_def, db_name)
                                break
                        except Exception:
                            continue
                    database_by_id[arb_id] = definition

                if definition is not None:
                    can_message, dbc_name = definition
                    try:
                        # Decode CAN payload (handles truncated / variable DLC gracefully)
                        try:
                            decoded = can_message.decode(message.data, decode_choices=False, allow_truncated=True)
                        except TypeError:
                            decoded = can_message.decode(message.data, decode_choices=False)

                        for signal_name, signal_value in decoded.items():
                            try:
                                val = float(signal_value)
                            except (ValueError, TypeError):
                                continue

                            # Standardized signal name: Message.Signal (e.g. BMS_Status.PackVoltage)
                            full_name = f"{can_message.name}.{signal_name}" if can_message.name else signal_name
                            if full_name not in signals:
                                signals[full_name] = ([], [])
                                self.signal_units[full_name] = unit_by_key.get(full_name, unit_by_key.get(signal_name, ""))

                            signals[full_name][0].append(timestamp)
                            signals[full_name][1].append(val)

                        decoded_messages += 1
                        if dbc_name not in self.blf_dbc_names:
                            self.blf_dbc_names.append(dbc_name)
                    except Exception:
                        pass

        # Convert to numpy arrays
        self.blf_data = {
            name: (np.asarray(timestamps, dtype=float), np.asarray(values, dtype=float))
            for name, (timestamps, values) in signals.items()
            if len(timestamps) > 0
        }
        self.signal_names = sorted(self.blf_data, key=str.lower)
        for name in self.signal_names:
            if name not in self.signal_units:
                self.signal_units[name] = ""
        self._update_time_bounds()

        if not self.signal_names:
            raise ValueError(
                f"BLF file '{self.filename}' was scanned ({frame_count:,} frames), "
                f"but no CAN messages matched the loaded DBC message IDs.\n"
                f"Please verify that the correct DBC database for this vehicle bus is uploaded."
            )

    def rescan_blf(self, dbc_paths: List[str]) -> int:
        """
        Re-scans this Vector BLF file with updated DBC database files,
        decoding all newly matched CAN messages into signals.
        """
        if Path(self.path).suffix.lower() == ".blf":
            self.dbc_paths = list(dict.fromkeys(dbc_paths + self.dbc_paths))
            self.signal_cache.clear()
            self._load_blf(self.path, self.dbc_paths)
            return len(self.signal_names)
        return 0

    def _inspect_channels(self):
        if self.mdf is None:
            return
        names = set()
        for name in self.mdf.channels_db.keys():
            if not name or str(name).startswith("$"):
                continue
            if str(name).strip().lower() in ("time", "t"):
                continue
            names.add(str(name))
        self.signal_names = sorted(names, key=str.lower)
        for name in self.signal_names:
            unit = self._lookup_unit(name)
            self.signal_units[name] = unit

    def _lookup_unit(self, name: str) -> str:
        if self.mdf is None:
            return ""
        try:
            unit = self.mdf.get_channel_unit(name)
            if unit:
                return str(unit)
        except Exception:
            pass
        try:
            group, index = self.mdf.channels_db[name][0]
            ch = self.mdf.groups[group].channels[index]
            return str(getattr(ch, "unit", "") or "")
        except Exception:
            return ""

    def get_signal_data(self, name: str) -> Optional[Tuple[np.ndarray, np.ndarray]]:
        if name in self.signal_cache:
            return self.signal_cache[name]
        if self.blf_data is not None:
            data = self.blf_data.get(name)
            if data is not None:
                self.signal_cache[name] = data
                return data
        if self.chn_data is not None:
            data = self.chn_data.get(name)
            if data is not None:
                self.signal_cache[name] = data
                return data
        if self.csv_data is not None:
            try:
                if name not in self.csv_data:
                    return None
                values = self.csv_data[name].to_numpy(dtype=float)
                timestamps = self._csv_timestamps
                valid = np.isfinite(values) & np.isfinite(timestamps)
                if not valid.any():
                    return None
                res = (timestamps[valid], values[valid])
                self.signal_cache[name] = res
                return res
            except Exception:
                return None
        try:
            group, index = self.mdf.channels_db[name][0]
            sig = self.mdf.get(group=group, index=index)
            y = np.asarray(sig.samples)
            if y.ndim != 1 or not np.issubdtype(y.dtype, np.number):
                return None
            x = np.asarray(sig.timestamps, dtype=float)
            if len(x) == 0:
                return None
            res = (x, y.astype(float))
            if len(self.signal_cache) > 100:
                self.signal_cache.pop(next(iter(self.signal_cache)))
            self.signal_cache[name] = res
            return res
        except Exception:
            return None

    def close(self):
        if self.mdf is not None:
            try:
                self.mdf.close()
            except Exception:
                pass


class DraggableSignalTree(QTreeWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setHeaderLabels(["Signal", "Unit", "Source"])
        self.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self.setDragEnabled(True)
        self.setAcceptDrops(False)
        self.setRootIsDecorated(False)
        self.setUniformRowHeights(True)
        self.setSortingEnabled(True)
        self.header().setSectionResizeMode(0, QHeaderView.Stretch)
        self.header().setSectionResizeMode(1, QHeaderView.Fixed)
        self.header().setSectionResizeMode(2, QHeaderView.Fixed)
        self.header().resizeSection(1, 62)
        self.header().resizeSection(2, 170)
        self.header().setMinimumSectionSize(42)
        self.setColumnWidth(0, 240)

    def startDrag(self, supportedActions):
        selected_items = self.selectedItems()
        if not selected_items:
            return
        drag_data = []
        for item in selected_items:
            payload = item.data(0, Qt.UserRole)
            if payload:
                drag_data.append(payload)

        if not drag_data:
            return

        mime_data = QMimeData()
        mime_data.setData("application/x-mf4-signals", json.dumps(drag_data).encode("utf-8"))

        drag = QDrag(self)
        drag.setMimeData(mime_data)
        drag.exec_(Qt.CopyAction)


class DropGraphicsLayoutWidget(pg.GraphicsLayoutWidget):
    signals_dropped = Signal(list)
    files_dropped = Signal(list)

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setAcceptDrops(True)
        self.is_drag_hovering = False

    def dragEnterEvent(self, event):
        if event.mimeData().hasFormat("application/x-mf4-signals") or event.mimeData().hasUrls():
            event.acceptProposedAction()
            self.is_drag_hovering = True
            self.update()
        else:
            event.ignore()

    def dragMoveEvent(self, event):
        if event.mimeData().hasFormat("application/x-mf4-signals") or event.mimeData().hasUrls():
            event.acceptProposedAction()
        else:
            event.ignore()

    def dragLeaveEvent(self, event):
        self.is_drag_hovering = False
        self.update()
        event.accept()

    def dropEvent(self, event):
        self.is_drag_hovering = False
        self.update()
        if event.mimeData().hasFormat("application/x-mf4-signals"):
            raw = event.mimeData().data("application/x-mf4-signals").data().decode("utf-8")
            try:
                sig_list = json.loads(raw)
                self.signals_dropped.emit(sig_list)
                event.acceptProposedAction()
            except Exception:
                event.ignore()
        elif event.mimeData().hasUrls():
            paths = []
            for url in event.mimeData().urls():
                p = url.toLocalFile()
                if Path(p).suffix.lower() in [".mf4", ".mdf", ".dat", ".csv", ".xlsx", ".xls", ".txt", ".blf", ".chn", ".dbc"]:
                    paths.append(p)
            if paths:
                self.files_dropped.emit(paths)
                event.acceptProposedAction()
            else:
                event.ignore()
        else:
            event.ignore()

    def paintEvent(self, event):
        super().paintEvent(event)
        if self.is_drag_hovering:
            painter = QPainter(self.viewport())
            painter.setRenderHint(QPainter.Antialiasing)
            pen = QPen(QColor("#0A84FF"), 2, Qt.DashLine)
            painter.setPen(pen)
            painter.setBrush(QColor(10, 132, 255, 30))
            painter.drawRoundedRect(self.rect().adjusted(4, 4, -4, -4), 10, 10)
            painter.end()


class FileLoadWorker(QObject):
    finished = Signal(object, object)
    progress = Signal(str, int)
    cancelled = Signal()

    def __init__(self, jobs):
        super().__init__()
        self.jobs = jobs
        self._cancel_requested = threading.Event()

    def request_cancel(self):
        self._cancel_requested.set()

    def is_cancelled(self):
        return self._cancel_requested.is_set()

    def run(self):
        records = {}
        errors = []
        total = max(1, len(self.jobs))
        for idx, (file_id, path, color_tag, dbc_paths) in enumerate(self.jobs):
            if self.is_cancelled():
                self.cancelled.emit()
                self.finished.emit(records, errors)
                return
            fname = Path(path).name
            base_pct = int((idx / total) * 100)
            self.progress.emit(f"Loading {fname}...", max(1, base_pct))

            def make_sub_cb(file_idx=idx, filename=fname):
                def _sub_cb(msg: str, frac: float = 0.5):
                    calc_pct = int(min(99, ((file_idx + max(0.0, min(1.0, frac))) / total) * 100))
                    self.progress.emit(msg, max(1, calc_pct))
                return _sub_cb

            try:
                cached_sheets = None
                sheet_names = [None]
                if Path(path).suffix.lower() in {".xlsx", ".xls"}:
                    self.progress.emit(f"Fast reading {fname}...", min(99, base_pct + 5))
                    cached_sheets = read_excel_all_sheets_fast(Path(path))
                    sheet_names = list(cached_sheets.keys())
                loaded_sheet_count = 0
                for sheet_name in sheet_names:
                    if self.is_cancelled():
                        self.cancelled.emit()
                        self.finished.emit(records, errors)
                        return
                    record_id = file_id if sheet_name is None else f"{file_id} [{sheet_name}]"
                    try:
                        record = MDFFileRecord(
                            file_id=record_id,
                            path=path,
                            color_tag=color_tag,
                            dbc_paths=dbc_paths,
                            progress_callback=make_sub_cb(idx, fname),
                            sheet_name=sheet_name if sheet_name is not None else "auto",
                            raw_data=cached_sheets.get(sheet_name) if cached_sheets is not None else None,
                        )
                    except (ValueError, KeyError) as sheet_error:
                        if sheet_name is not None:
                            self.progress.emit(f"Skipped sheet {sheet_name}: {sheet_error}", max(1, base_pct))
                            continue
                        raise
                    if sheet_name is not None and not record.signal_names:
                        continue
                    records[record_id] = record
                    loaded_sheet_count += 1
                if loaded_sheet_count == 0:
                    raise ValueError(f"No measurement sheets found in {fname}")
                end_pct = int(((idx + 1) / total) * 100)
                self.progress.emit(f"Loaded {fname}", min(100 if idx + 1 == total else 99, end_pct))
            except Exception as exc:
                errors.append((fname, str(exc)))

        self.progress.emit("Finished loading all files", 100)
        self.finished.emit(records, errors)


class PowerEnergyCalculatorDialog(QDialog):
    """
    Dialog to calculate:
    Power = Voltage * Current [W]
    Sample Energy = Voltage * Current * dt [W.s] (default dt = 0.01s)
    Cumulative Energy = sum(V * I * dt) / 1000 / 3600 [kW.h]

    Features dedicated real-time search and filtering options for Voltage and Current signals,
    auto-completion, quick filter chips, and smart signal pair detection.
    """
    def __init__(self, main_window: "MF4VehicleExplorer", parent=None):
        super().__init__(parent)
        self.main_window = main_window
        self.setWindowTitle("⚡ Power & Energy Calculator (P = V * I)")
        self.setMinimumWidth(620)
        self.setModal(True)
        self.all_signals: List[Tuple[str, str, str, str]] = []  # (display_text, file_id, signal_name, unit)
        self._collect_available_signals()
        self._init_ui()

    def _collect_available_signals(self):
        self.all_signals = []
        for sname, info in self.main_window.derived_signals.items():
            unit = info.get("unit", "")
            leaf = extract_signal_only_name(sname)
            disp = f"⚡ Derived :: {leaf} [{unit}]" if unit else f"⚡ Derived :: {leaf}"
            self.all_signals.append((disp, "derived", sname, unit))

        for fid, rec in self.main_window.loaded_files.items():
            for sname in rec.signal_names:
                unit = rec.signal_units.get(sname, "")
                leaf = extract_signal_only_name(sname)
                disp = f"{fid} :: {leaf} [{unit}]" if unit else f"{fid} :: {leaf}"
                self.all_signals.append((disp, fid, sname, unit))

    def _init_ui(self):
        layout = QVBoxLayout(self)
        layout.setSpacing(12)

        header_card = QFrame()
        header_card.setObjectName("card")
        header_l = QVBoxLayout(header_card)
        title = QLabel("Calculate Electrical Power & Energy from Signals")
        title.setStyleSheet("font-weight: bold; font-size: 13px; color: #0A84FF;")
        formula_lbl = QLabel(
            "• Instantaneous Power: Power [W] = Voltage [V] × Current [A]\n"
            "• Integration Step: Energy [W.s] = Voltage × Current × Δt (0.01s)\n"
            "• Cumulative Energy: Energy [kW.h] = Sum(Energy [W.s]) / 1000 / 3600"
        )
        formula_lbl.setStyleSheet("color: #86868B; font-size: 11px;")
        header_l.addWidget(title)
        header_l.addWidget(formula_lbl)
        layout.addWidget(header_card)

        sig_group = QGroupBox("Signal Selection (Search Voltage & Current Channels):")
        sig_layout = QVBoxLayout(sig_group)
        sig_layout.setSpacing(10)

        auto_bar = QHBoxLayout()
        auto_bar.setSpacing(8)
        lbl_auto_hint = QLabel("Search or select channels, or auto-match:")
        lbl_auto_hint.setStyleSheet("color: #86868B; font-size: 11px;")
        auto_bar.addWidget(lbl_auto_hint)
        auto_bar.addStretch()

        self.btn_auto_pair = QPushButton("⚡ Auto-Detect & Pair V & I")
        self.btn_auto_pair.setToolTip("Automatically searches and selects best matching voltage and current signals")
        self.btn_auto_pair.clicked.connect(self._auto_detect_signals)
        auto_bar.addWidget(self.btn_auto_pair)
        sig_layout.addLayout(auto_bar)

        # 1. VOLTAGE CARD
        volt_card = QFrame()
        volt_card.setObjectName("card")
        volt_box = QVBoxLayout(volt_card)
        volt_box.setContentsMargins(10, 8, 10, 10)
        volt_box.setSpacing(6)

        volt_hdr = QHBoxLayout()
        lbl_v_title = QLabel("Voltage Signal (V):")
        lbl_v_title.setStyleSheet("font-weight: 600; color: #0A84FF;")
        volt_hdr.addWidget(lbl_v_title)
        volt_hdr.addStretch()
        self.lbl_v_count = QLabel(f"{len(self.all_signals)} available")
        self.lbl_v_count.setStyleSheet("color: #86868B; font-size: 11px;")
        volt_hdr.addWidget(self.lbl_v_count)
        volt_box.addLayout(volt_hdr)

        v_search_row = QHBoxLayout()
        v_search_row.setSpacing(6)
        self.search_voltage = QLineEdit()
        self.search_voltage.setPlaceholderText("🔍 Search voltage (e.g. 'volt', 'vbatt', 'hv_v', '[v]', 'pack')...")
        self.search_voltage.setClearButtonEnabled(True)
        self.search_voltage.textChanged.connect(self._filter_voltage_signals)
        v_search_row.addWidget(self.search_voltage)
        volt_box.addLayout(v_search_row)

        v_chips_row = QHBoxLayout()
        v_chips_row.setSpacing(4)
        v_chips_lbl = QLabel("Quick filters:")
        v_chips_lbl.setStyleSheet("color: #86868B; font-size: 10px;")
        v_chips_row.addWidget(v_chips_lbl)
        for chip_text in ["[V]", "volt", "vbatt", "pack_v", "cell"]:
            btn_chip = QPushButton(chip_text)
            btn_chip.setFixedHeight(22)
            btn_chip.setStyleSheet("padding: 2px 7px; font-size: 10px; border-radius: 4px;")
            btn_chip.clicked.connect(lambda _, t=chip_text: self.search_voltage.setText(t))
            v_chips_row.addWidget(btn_chip)
        v_chips_row.addStretch()
        volt_box.addLayout(v_chips_row)

        self.combo_voltage = QComboBox()
        self.combo_voltage.setEditable(True)
        self.combo_voltage.setInsertPolicy(QComboBox.NoInsert)
        volt_box.addWidget(self.combo_voltage)

        sig_layout.addWidget(volt_card)

        # 2. CURRENT CARD
        curr_card = QFrame()
        curr_card.setObjectName("card")
        curr_box = QVBoxLayout(curr_card)
        curr_box.setContentsMargins(10, 8, 10, 10)
        curr_box.setSpacing(6)

        curr_hdr = QHBoxLayout()
        lbl_i_title = QLabel("Current Signal (I):")
        lbl_i_title.setStyleSheet("font-weight: 600; color: #30D158;")
        curr_hdr.addWidget(lbl_i_title)
        curr_hdr.addStretch()
        self.lbl_i_count = QLabel(f"{len(self.all_signals)} available")
        self.lbl_i_count.setStyleSheet("color: #86868B; font-size: 11px;")
        curr_hdr.addWidget(self.lbl_i_count)
        curr_box.addLayout(curr_hdr)

        i_search_row = QHBoxLayout()
        i_search_row.setSpacing(6)
        self.search_current = QLineEdit()
        self.search_current.setPlaceholderText("🔍 Search current (e.g. 'curr', 'ibatt', 'amps', '[a]', 'phase')...")
        self.search_current.setClearButtonEnabled(True)
        self.search_current.textChanged.connect(self._filter_current_signals)
        i_search_row.addWidget(self.search_current)
        curr_box.addLayout(i_search_row)

        i_chips_row = QHBoxLayout()
        i_chips_row.setSpacing(4)
        i_chips_lbl = QLabel("Quick filters:")
        i_chips_lbl.setStyleSheet("color: #86868B; font-size: 10px;")
        i_chips_row.addWidget(i_chips_lbl)
        for chip_text in ["[A]", "curr", "ibatt", "pack_i", "phase"]:
            btn_chip = QPushButton(chip_text)
            btn_chip.setFixedHeight(22)
            btn_chip.setStyleSheet("padding: 2px 7px; font-size: 10px; border-radius: 4px;")
            btn_chip.clicked.connect(lambda _, t=chip_text: self.search_current.setText(t))
            i_chips_row.addWidget(btn_chip)
        i_chips_row.addStretch()
        curr_box.addLayout(i_chips_row)

        self.combo_current = QComboBox()
        self.combo_current.setEditable(True)
        self.combo_current.setInsertPolicy(QComboBox.NoInsert)
        curr_box.addWidget(self.combo_current)

        sig_layout.addWidget(curr_card)
        layout.addWidget(sig_group)

        self._filter_voltage_signals("")
        self._filter_current_signals("")
        self._guess_defaults()

        form_layout = QFormLayout()
        form_layout.setSpacing(10)

        dt_layout = QHBoxLayout()
        self.spin_dt = QDoubleSpinBox()
        self.spin_dt.setRange(0.0001, 10.0)
        self.spin_dt.setDecimals(4)
        self.spin_dt.setSingleStep(0.005)
        self.spin_dt.setValue(0.01)
        self.spin_dt.setSuffix(" s (10 ms)")
        dt_layout.addWidget(self.spin_dt)

        self.chk_auto_dt = QCheckBox("Auto-detect dt from time vector")
        self.chk_auto_dt.setChecked(False)
        dt_layout.addWidget(self.chk_auto_dt)
        form_layout.addRow("Sampling Interval (Δt):", dt_layout)

        time_group = QGroupBox("Calculation Time & Sample Start Window:")
        time_layout = QVBoxLayout(time_group)
        time_layout.setSpacing(6)

        mode_row = QHBoxLayout()
        self.combo_time_mode = QComboBox()
        self.combo_time_mode.addItems(["Custom Time (s)", "Sample Index (#)", "From Origin (0s)"])
        mode_row.addWidget(QLabel("Start Mode:"))
        mode_row.addWidget(self.combo_time_mode)
        time_layout.addLayout(mode_row)

        start_inputs_layout = QHBoxLayout()
        start_inputs_layout.addWidget(QLabel("Start Time (t_start):"))
        self.spin_start_time = QDoubleSpinBox()
        self.spin_start_time.setRange(-1e6, 1e6)
        self.spin_start_time.setDecimals(3)
        self.spin_start_time.setSingleStep(0.5)
        self.spin_start_time.setValue(0.0)
        self.spin_start_time.setSuffix(" s")
        start_inputs_layout.addWidget(self.spin_start_time)

        start_inputs_layout.addWidget(QLabel("Sample Index:"))
        self.spin_start_sample = QSpinBox()
        self.spin_start_sample.setRange(0, 10000000)
        self.spin_start_sample.setValue(0)
        self.spin_start_sample.setEnabled(False)
        start_inputs_layout.addWidget(self.spin_start_sample)
        time_layout.addLayout(start_inputs_layout)

        def _on_time_mode_changed(idx):
            self.spin_start_time.setEnabled(idx == 0)
            self.spin_start_sample.setEnabled(idx == 1)
        self.combo_time_mode.currentIndexChanged.connect(_on_time_mode_changed)

        end_row = QHBoxLayout()
        self.chk_custom_end = QCheckBox("Limit End Time (t_end):")
        self.chk_custom_end.setChecked(False)
        end_row.addWidget(self.chk_custom_end)
        self.spin_end_time = QDoubleSpinBox()
        self.spin_end_time.setRange(-1e6, 1e6)
        self.spin_end_time.setDecimals(3)
        self.spin_end_time.setValue(10.0)
        self.spin_end_time.setSuffix(" s")
        self.spin_end_time.setEnabled(False)
        self.chk_custom_end.toggled.connect(self.spin_end_time.setEnabled)
        end_row.addWidget(self.spin_end_time)
        time_layout.addLayout(end_row)

        self.chk_align_zero = QCheckBox("Zero-align output timestamps (start at relative t = 0.0s)")
        self.chk_align_zero.setChecked(False)
        time_layout.addWidget(self.chk_align_zero)

        layout.addWidget(time_group)

        self.txt_prefix = QLineEdit("Calc")
        form_layout.addRow("Signal Name Tag:", self.txt_prefix)
        layout.addLayout(form_layout)

        gen_group = QGroupBox("Signals to Generate and Add to Tree:")
        gen_layout = QVBoxLayout(gen_group)

        self.chk_gen_power_w = QCheckBox("Power = V * I  [W]")
        self.chk_gen_power_w.setChecked(True)
        gen_layout.addWidget(self.chk_gen_power_w)

        self.chk_gen_power_kw = QCheckBox("Power in kW = (V * I) / 1000  [kW]")
        self.chk_gen_power_kw.setChecked(True)
        gen_layout.addWidget(self.chk_gen_power_kw)

        self.chk_gen_energy_ws = QCheckBox("Instantaneous Step Energy = V * I * dt  [W.s]")
        self.chk_gen_energy_ws.setChecked(False)
        gen_layout.addWidget(self.chk_gen_energy_ws)

        self.chk_gen_energy_kwh = QCheckBox("Cumulative Energy = sum(V * I * dt) / 1000 / 3600  [kW.h]")
        self.chk_gen_energy_kwh.setChecked(True)
        gen_layout.addWidget(self.chk_gen_energy_kwh)

        self.chk_plot_immediately = QCheckBox("Plot generated signals immediately")
        self.chk_plot_immediately.setChecked(True)
        gen_layout.addWidget(self.chk_plot_immediately)

        layout.addWidget(gen_group)

        btn_box = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        btn_ok = btn_box.button(QDialogButtonBox.Ok)
        btn_ok.setText("Compute & Add Signals")
        btn_ok.setObjectName("primaryBtn")
        btn_box.accepted.connect(self.calculate)
        btn_box.rejected.connect(self.reject)
        layout.addWidget(btn_box)

    def _filter_voltage_signals(self, query: str = ""):
        prev_data = self.combo_voltage.currentData()
        q_clean = query.lower().strip()
        tokens = q_clean.split() if q_clean else []

        matching = []
        for text, fid, sname, unit in self.all_signals:
            if not tokens:
                matching.append((text, fid, sname, unit))
                continue
            haystack = f"{fid} {sname} {unit} [{unit}]".lower()
            if all(tok in haystack for tok in tokens):
                matching.append((text, fid, sname, unit))

        self.combo_voltage.blockSignals(True)
        self.combo_voltage.clear()

        selected_idx = -1
        for idx, (text, fid, sname, unit) in enumerate(matching):
            self.combo_voltage.addItem(text, (fid, sname, unit))
            if prev_data and (fid, sname) == (prev_data[0], prev_data[1]):
                selected_idx = idx

        if selected_idx >= 0:
            self.combo_voltage.setCurrentIndex(selected_idx)
        elif self.combo_voltage.count() > 0:
            self.combo_voltage.setCurrentIndex(0)

        self.combo_voltage.blockSignals(False)

        completer = QCompleter([item[0] for item in matching], self.combo_voltage)
        completer.setCaseSensitivity(Qt.CaseInsensitive)
        completer.setFilterMode(Qt.MatchContains)
        self.combo_voltage.setCompleter(completer)

        total = len(self.all_signals)
        match_count = len(matching)
        if q_clean:
            self.lbl_v_count.setText(f"{match_count} / {total} matches")
            self.lbl_v_count.setStyleSheet("color: #0A84FF; font-weight: 600; font-size: 11px;")
        else:
            self.lbl_v_count.setText(f"{total} available")
            self.lbl_v_count.setStyleSheet("color: #86868B; font-size: 11px;")

    def _filter_current_signals(self, query: str = ""):
        prev_data = self.combo_current.currentData()
        q_clean = query.lower().strip()
        tokens = q_clean.split() if q_clean else []

        matching = []
        for text, fid, sname, unit in self.all_signals:
            if not tokens:
                matching.append((text, fid, sname, unit))
                continue
            haystack = f"{fid} {sname} {unit} [{unit}]".lower()
            if all(tok in haystack for tok in tokens):
                matching.append((text, fid, sname, unit))

        self.combo_current.blockSignals(True)
        self.combo_current.clear()

        selected_idx = -1
        for idx, (text, fid, sname, unit) in enumerate(matching):
            self.combo_current.addItem(text, (fid, sname, unit))
            if prev_data and (fid, sname) == (prev_data[0], prev_data[1]):
                selected_idx = idx

        if selected_idx >= 0:
            self.combo_current.setCurrentIndex(selected_idx)
        elif self.combo_current.count() > 0:
            self.combo_current.setCurrentIndex(0)

        self.combo_current.blockSignals(False)

        completer = QCompleter([item[0] for item in matching], self.combo_current)
        completer.setCaseSensitivity(Qt.CaseInsensitive)
        completer.setFilterMode(Qt.MatchContains)
        self.combo_current.setCompleter(completer)

        total = len(self.all_signals)
        match_count = len(matching)
        if q_clean:
            self.lbl_i_count.setText(f"{match_count} / {total} matches")
            self.lbl_i_count.setStyleSheet("color: #30D158; font-weight: 600; font-size: 11px;")
        else:
            self.lbl_i_count.setText(f"{total} available")
            self.lbl_i_count.setStyleSheet("color: #86868B; font-size: 11px;")

    def _guess_defaults(self):
        for i in range(self.combo_voltage.count()):
            txt = self.combo_voltage.itemText(i).lower()
            if "volt" in txt or "u_" in txt or "[v]" in txt or "_v" in txt or "vbatt" in txt:
                self.combo_voltage.setCurrentIndex(i)
                break

        for i in range(self.combo_current.count()):
            txt = self.combo_current.itemText(i).lower()
            if "curr" in txt or "i_" in txt or "[a]" in txt or "_i" in txt or "ibatt" in txt:
                self.combo_current.setCurrentIndex(i)
                break

    def _auto_detect_signals(self):
        v_keywords = ["volt", "u_", "_v", "vbatt", "v_pack", "v_bus", "pack_v", "vdc"]
        i_keywords = ["curr", "i_", "_i", "ibatt", "i_pack", "i_bus", "pack_i", "idc", "amps"]

        best_v_idx = -1
        for i in range(self.combo_voltage.count()):
            txt = self.combo_voltage.itemText(i).lower()
            if any(kw in txt for kw in v_keywords) or "[v]" in txt:
                best_v_idx = i
                break

        best_i_idx = -1
        for i in range(self.combo_current.count()):
            txt = self.combo_current.itemText(i).lower()
            if any(kw in txt for kw in i_keywords) or "[a]" in txt:
                best_i_idx = i
                break

        if best_v_idx >= 0:
            self.combo_voltage.setCurrentIndex(best_v_idx)
        if best_i_idx >= 0:
            self.combo_current.setCurrentIndex(best_i_idx)

        v_name = self.combo_voltage.currentText()
        i_name = self.combo_current.currentText()
        if best_v_idx >= 0 or best_i_idx >= 0:
            QMessageBox.information(
                self, "Signals Auto-Detected",
                f"Auto-selected signals for power calculation:\n\n"
                f"• Voltage: {v_name}\n"
                f"• Current: {i_name}"
            )
        else:
            QMessageBox.warning(self, "Auto-Detection", "Could not find standard voltage or current keywords.")

    @staticmethod
    def _prepare_numeric_signal(
        timestamps: np.ndarray, samples: np.ndarray
    ) -> Tuple[np.ndarray, np.ndarray]:
        """Return finite, time-sorted, interpolation-safe numeric samples."""
        x = np.asarray(timestamps, dtype=float).reshape(-1)
        y = np.asarray(samples, dtype=float).reshape(-1)
        if x.size != y.size:
            raise ValueError("timestamp and sample arrays have different lengths")
        valid = np.isfinite(x) & np.isfinite(y)
        x, y = x[valid], y[valid]
        if x.size == 0:
            return x, y
        order = np.argsort(x, kind="stable")
        x, y = x[order], y[order]
        # np.interp requires increasing timestamps. Keep the last value when
        # a logger supplied multiple samples at an identical timestamp.
        reverse_unique = np.unique(x[::-1], return_index=True)[1]
        keep = np.sort(x.size - 1 - reverse_unique)
        return x[keep], y[keep]

    @staticmethod
    def _electrical_unit_scale(unit: str, quantity: str) -> Optional[float]:
        """Scale common electrical units to base SI units (V or A)."""
        normalized = str(unit or "").strip().replace("µ", "u").replace("μ", "u")
        scales = {
            "voltage": {
                "V": 1.0, "volt": 1.0, "volts": 1.0,
                "mV": 1e-3, "uV": 1e-6, "kV": 1e3,
            },
            "current": {
                "A": 1.0, "amp": 1.0, "amps": 1.0,
                "ampere": 1.0, "amperes": 1.0,
                "mA": 1e-3, "uA": 1e-6, "kA": 1e3,
            },
        }
        if not normalized:
            return 1.0
        for expected, scale in scales[quantity].items():
            if normalized.casefold() == expected.casefold():
                return scale
        return None

    def calculate(self):
        v_data = self.combo_voltage.currentData()
        i_data = self.combo_current.currentData()

        if not v_data or not i_data:
            QMessageBox.warning(self, "Missing Selection", "Please select both a voltage and a current signal.")
            return

        if not any((
            self.chk_gen_power_w.isChecked(),
            self.chk_gen_power_kw.isChecked(),
            self.chk_gen_energy_ws.isChecked(),
            self.chk_gen_energy_kwh.isChecked(),
        )):
            QMessageBox.warning(self, "Nothing Selected", "Select at least one power or energy signal to generate.")
            return

        fid_v, name_v, _ = v_data
        fid_i, name_i, _ = i_data

        sig_v = self.main_window._fetch_raw_signal("file" if fid_v != "derived" else "derived", fid_v, name_v)
        sig_i = self.main_window._fetch_raw_signal("file" if fid_i != "derived" else "derived", fid_i, name_i)

        if sig_v is None or sig_i is None:
            QMessageBox.critical(self, "Data Error", "Could not fetch sample arrays for the selected signals.")
            return

        xv, yv, unit_v = sig_v
        xi, yi, unit_i = sig_i

        voltage_scale = self._electrical_unit_scale(unit_v, "voltage")
        current_scale = self._electrical_unit_scale(unit_i, "current")
        if voltage_scale is None or current_scale is None:
            QMessageBox.warning(
                self,
                "Incompatible Units",
                "Power requires a voltage signal and a current signal.\n\n"
                f"Selected voltage unit: {unit_v or '(missing)'}\n"
                f"Selected current unit: {unit_i or '(missing)'}\n\n"
                "Supported units: V, mV, uV, kV and A, mA, uA, kA.",
            )
            return

        try:
            xv, yv = self._prepare_numeric_signal(xv, yv)
            xi, yi = self._prepare_numeric_signal(xi, yi)
        except (TypeError, ValueError) as exc:
            QMessageBox.critical(self, "Data Error", f"Voltage/current samples are not numeric: {exc}")
            return

        yv = yv * voltage_scale
        yi = yi * current_scale

        if len(xv) == 0 or len(xi) == 0:
            QMessageBox.warning(self, "Empty Signal", "One of the signals contains 0 samples.")
            return

        t_start = max(xv[0], xi[0])
        t_end = min(xv[-1], xi[-1])

        if t_end <= t_start:
            QMessageBox.warning(self, "No Overlap", "The selected signals do not have overlapping timestamps.")
            return

        if self.chk_auto_dt.isChecked():
            intervals = []
            for time_values in (xv, xi):
                positive_steps = np.diff(time_values)
                positive_steps = positive_steps[positive_steps > 0]
                if positive_steps.size:
                    intervals.append(float(np.median(positive_steps)))
            # Use the finer source interval so the faster input is not lost.
            dt = min(intervals) if intervals else 0.01
        else:
            dt = float(self.spin_dt.value())

        mode_idx = self.combo_time_mode.currentIndex()
        if mode_idx == 0:
            t_start = max(t_start, float(self.spin_start_time.value()))
        elif mode_idx == 1:
            sample_offset = int(self.spin_start_sample.value()) * dt
            t_start = max(t_start, t_start + sample_offset)

        if self.chk_custom_end.isChecked():
            t_end = min(t_end, float(self.spin_end_time.value()))

        if t_end <= t_start:
            QMessageBox.warning(self, "Invalid Window", f"End time ({t_end:.3f}s) must be greater than Start time ({t_start:.3f}s).")
            return

        common_t = np.arange(t_start, t_end + dt * 0.5, dt)
        common_t = common_t[common_t <= t_end]
        if len(common_t) < 2:
            QMessageBox.warning(self, "Range Too Small", "Not enough samples generated in the overlapping interval.")
            return

        v_interp = np.interp(common_t, xv, yv)
        i_interp = np.interp(common_t, xi, yi)

        power_w = v_interp * i_interp
        power_kw = power_w / 1000.0
        energy_step_ws = np.gradient(common_t) * power_w
        cumulative_energy_ws = np.concatenate((
            [0.0],
            np.cumsum((power_w[:-1] + power_w[1:]) * np.diff(common_t) / 2.0),
        ))
        cumulative_energy_kwh = cumulative_energy_ws / (1000.0 * 3600.0)

        plot_t = (common_t - common_t[0]) if self.chk_align_zero.isChecked() else common_t

        tag = self.txt_prefix.text().strip()
        prefix = f"{tag}_" if tag else ""

        created_signals = []

        if self.chk_gen_power_w.isChecked():
            sig_name = f"{prefix}Power"
            self.main_window.derived_signals[sig_name] = {"x": plot_t, "y": power_w, "unit": "W"}
            created_signals.append(sig_name)

        if self.chk_gen_power_kw.isChecked():
            sig_name = f"{prefix}Power_kW"
            self.main_window.derived_signals[sig_name] = {"x": plot_t, "y": power_kw, "unit": "kW"}
            created_signals.append(sig_name)

        if self.chk_gen_energy_ws.isChecked():
            sig_name = f"{prefix}Energy_Step_Ws"
            self.main_window.derived_signals[sig_name] = {"x": plot_t, "y": energy_step_ws, "unit": "W.s"}
            created_signals.append(sig_name)

        if self.chk_gen_energy_kwh.isChecked():
            sig_name = f"{prefix}Cumulative_Energy_kWh"
            self.main_window.derived_signals[sig_name] = {"x": plot_t, "y": cumulative_energy_kwh, "unit": "kW.h"}
            created_signals.append(sig_name)

        self.main_window.populate_signal_tree()

        integrate = getattr(np, "trapezoid", None)
        if integrate is None:
            integrate = np.trapz
        total_ws = float(integrate(power_w, common_t)) if len(common_t) > 1 else 0.0
        consumed_ws = float(integrate(np.clip(power_w, 0.0, None), common_t)) if len(common_t) > 1 else 0.0
        regenerated_ws = float(integrate(np.clip(power_w, None, 0.0), common_t)) if len(common_t) > 1 else 0.0
        total_kwh = total_ws / (1000.0 * 3600.0)
        consumed_kwh = consumed_ws / (1000.0 * 3600.0)
        regenerated_kwh = regenerated_ws / (1000.0 * 3600.0)
        peak_power_kw = np.max(power_kw) if len(power_kw) > 0 else 0.0
        min_power_kw = np.min(power_kw) if len(power_kw) > 0 else 0.0
        duration_s = common_t[-1] - common_t[0]

        summary_msg = (
            f"Successfully computed {len(created_signals)} derived signal(s):\n"
            f"• {', '.join(created_signals)}\n\n"
            f"Summary Metrics (Δt = {dt:.4f} s, {len(common_t):,} samples):\n"
            f"• Net Energy: {total_kwh:.4f} kW.h ({total_kwh * 3.6:.2f} MJ)\n"
            f"• Consumed Energy: {consumed_kwh:.4f} kW.h\n"
            f"• Regenerated Energy: {abs(regenerated_kwh):.4f} kW.h\n"
            f"• Peak Electrical Power: {peak_power_kw:.2f} kW\n"
            f"• Min Electrical Power (Regen): {min_power_kw:.2f} kW\n"
            f"• Time Duration: {duration_s:.2f} s"
        )

        if self.chk_plot_immediately.isChecked():
            for name in created_signals:
                self.main_window._add_signal_to_plotted({
                    "type": "derived",
                    "file_id": "derived",
                    "name": name
                })
            # Rebuild once so replaced derived signals and their axes cannot
            # leave stale or duplicate curves on the canvas.
            self.main_window.plot_current_curves(rebuild_layout=True)

        QMessageBox.information(self, "Calculation Complete", summary_msg)
        self.accept()


class OffsetComparisonPlot(pg.PlotWidget):
    """Comparison plot with right-button horizontal dragging for one trace."""

    offsetDragged = Signal(str, float)

    def __init__(self, parent=None):
        super().__init__(parent)
        self._dragging_offset = False
        self._drag_start_pixel = 0.0
        self._drag_units_per_pixel = 1.0
        self._drag_start_offset = 0.0
        self._drag_target = None
        self._trace_data = []
        self._offset_history = []

    def _scene_position(self, event):
        position = event.position().toPoint() if hasattr(event, "position") else event.pos()
        return self.mapToScene(position)

    def mousePressEvent(self, event):
        if event.button() == Qt.RightButton:
            scene_position = self._scene_position(event)
            view_position = self.plotItem.vb.mapSceneToView(scene_position)
            target = self._nearest_trace(float(view_position.x()), float(view_position.y()))
            if target is None:
                event.ignore()
                return
            self._drag_target = target
            self._dragging_offset = True
            self._drag_start_pixel = float(scene_position.x())
            view_range = self.plotItem.vb.viewRange()[0]
            scene_width = max(self.plotItem.vb.sceneBoundingRect().width(), 1.0)
            self._drag_units_per_pixel = (view_range[1] - view_range[0]) / scene_width
            self._drag_start_offset = float(target["offset"])
            self._offset_history.append((target["label"], self._drag_start_offset))
            self._offset_history = self._offset_history[-50:]
            self.setFocus()
            event.accept()
            return
        super().mousePressEvent(event)

    def mouseMoveEvent(self, event):
        if self._dragging_offset and (event.buttons() & Qt.RightButton):
            current_pixel = float(self._scene_position(event).x())
            self.offsetDragged.emit(
                self._drag_target["label"],
                self._drag_start_offset
                + (current_pixel - self._drag_start_pixel) * self._drag_units_per_pixel
            )
            event.accept()
            return
        super().mouseMoveEvent(event)

    def mouseReleaseEvent(self, event):
        if self._dragging_offset and event.button() == Qt.RightButton:
            self._dragging_offset = False
            self._drag_target = None
            event.accept()
            return
        super().mouseReleaseEvent(event)

    def keyPressEvent(self, event):
        target = self._drag_target or (self._trace_data[0] if self._trace_data else None)
        if target is None:
            super().keyPressEvent(event)
            return
        if event.key() in (Qt.Key_Left, Qt.Key_Right):
            direction = -1.0 if event.key() == Qt.Key_Left else 1.0
            step = float(np.median(np.diff(target["base_x"]))) if len(target["base_x"]) > 1 else 0.001
            if not np.isfinite(step) or step <= 0:
                step = 0.001
            if event.modifiers() & Qt.ShiftModifier:
                step *= 10.0
            self.offsetDragged.emit(target["label"], float(target["offset"]) + direction * step)
            event.accept()
            return
        if event.key() == Qt.Key_Z and event.modifiers() & Qt.ControlModifier:
            if self._offset_history:
                label, offset = self._offset_history.pop()
                self.offsetDragged.emit(label, offset)
            event.accept()
            return
        super().keyPressEvent(event)

    def set_trace_data(self, trace_data):
        self._trace_data = trace_data

    def _nearest_trace(self, x_value: float, y_value: float):
        if not self._trace_data:
            return None
        scene_point = self.plotItem.vb.mapViewToScene(pg.Point(x_value, y_value))
        best_trace = None
        best_distance = float("inf")
        for trace in self._trace_data:
            x = trace["x"]
            y = trace["y"]
            if len(x) < 2 or x_value < x[0] or x_value > x[-1]:
                continue
            trace_y = float(np.interp(x_value, x, y))
            trace_point = self.plotItem.vb.mapViewToScene(pg.Point(x_value, trace_y))
            distance = abs(float(trace_point.y()) - float(scene_point.y()))
            if distance < best_distance:
                best_distance = distance
                best_trace = trace
        return best_trace if best_distance <= 45.0 else None


class SpeedComparisonDialog(QDialog):
    def __init__(self, main_window: "MF4VehicleExplorer", parent=None):
        super().__init__(parent)
        self.main_window = main_window
        self.setWindowTitle("Signal Comparison - Manual Time Alignment")
        self.setMinimumSize(760, 760)
        self.setModal(False)
        self.options: List[Tuple[str, str, str, str]] = []
        self._build_ui()
        self._populate_signals()

    def _build_ui(self):
        layout = QVBoxLayout(self)
        layout.setSpacing(12)

        header = QFrame()
        header.setObjectName("card")
        header_layout = QVBoxLayout(header)
        title = QLabel("Align Two Speed Tests")
        title.setStyleSheet("font-weight: bold; font-size: 13px; color: #0A84FF;")
        description = QLabel(
            "Select one signal from each test. Both traces are shown without automatic "
            "synchronization; right-drag either trace to align it manually."
        )
        description.setWordWrap(True)
        description.setStyleSheet("color: #86868B; font-size: 11px;")
        header_layout.addWidget(title)
        header_layout.addWidget(description)
        layout.addWidget(header)

        self.search_test1 = QLineEdit()
        self.search_test1.setPlaceholderText("Search Test 1 signals...")
        self.search_test1.setClearButtonEnabled(True)
        self.search_test1.textChanged.connect(lambda text: self._filter_combo(self.combo_test1, text))
        layout.addWidget(self.search_test1)

        self.combo_test1 = QComboBox()
        self.combo_test1.setMinimumHeight(32)
        layout.addWidget(self.combo_test1)

        self.search_test2 = QLineEdit()
        self.search_test2.setPlaceholderText("Search Test 2 signals...")
        self.search_test2.setClearButtonEnabled(True)
        self.search_test2.textChanged.connect(lambda text: self._filter_combo(self.combo_test2, text))
        layout.addWidget(self.search_test2)

        self.combo_test2 = QComboBox()
        self.combo_test2.setMinimumHeight(32)
        layout.addWidget(self.combo_test2)

        offset_layout = QHBoxLayout()
        offset_layout.addWidget(QLabel("Signal 1 offset:"))
        self.spin_test1_offset = self._offset_spin()
        self.spin_test1_offset.valueChanged.connect(
            lambda value: self._numeric_offset_changed(
                self._combo_labels.get(
                    self.combo_test1,
                    self._selected_signal_label(self.combo_test1),
                ),
                value,
            )
        )
        offset_layout.addWidget(self.spin_test1_offset)
        offset_layout.addWidget(QLabel("Signal 2 offset:"))
        self.spin_test2_offset = self._offset_spin()
        self.spin_test2_offset.valueChanged.connect(
            lambda value: self._numeric_offset_changed(
                self._combo_labels.get(
                    self.combo_test2,
                    self._selected_signal_label(self.combo_test2),
                ),
                value,
            )
        )
        offset_layout.addWidget(self.spin_test2_offset)
        layout.addLayout(offset_layout)

        self.extra_datasets_layout = QVBoxLayout()
        layout.addLayout(self.extra_datasets_layout)
        self.extra_datasets = []
        self._combo_labels = {}
        add_dataset_button = QPushButton("+ Add dataset")
        add_dataset_button.clicked.connect(self._add_dataset)
        layout.addWidget(add_dataset_button)

        self.lbl_result = QLabel("The comparison plot will appear after alignment.")
        self.lbl_result.setStyleSheet("color: #86868B; font-size: 11px;")
        layout.addWidget(self.lbl_result)
        drag_hint = QLabel(
            "After alignment, hold the right mouse button on the comparison plot "
            "and drag horizontally to move the selected signal."
        )
        drag_hint.setWordWrap(True)
        drag_hint.setStyleSheet("color: #86868B; font-size: 11px;")
        layout.addWidget(drag_hint)

        self.chk_difference_view = QCheckBox("Difference vs reference")
        self.chk_difference_view.setToolTip("Plot the relative difference from the first selected signal instead of the raw values")
        layout.addWidget(self.chk_difference_view)

        self.reset_offset_button = QPushButton("Reset Offsets")
        self.reset_offset_button.clicked.connect(self._reset_offsets)
        layout.addWidget(self.reset_offset_button)
        self.auto_align_button = QPushButton("Auto-align all")
        self.auto_align_button.clicked.connect(self._auto_align_all)
        layout.addWidget(self.auto_align_button)

        self.btn_generate_report = QPushButton("Generate Difference Report")
        self.btn_generate_report.setObjectName("primaryBtn")
        self.btn_generate_report.clicked.connect(self._generate_difference_report)

        report_group = QGroupBox("Difference report parameters")
        report_layout = QGridLayout(report_group)
        report_layout.addWidget(QLabel("Reference test:"), 0, 0)
        self.report_reference = QComboBox()
        self.report_reference.setToolTip("Baseline test used to calculate percentage differences.")
        self.report_reference.addItems(list(self.main_window.loaded_files))
        report_layout.addWidget(self.report_reference, 0, 1)

        report_layout.addWidget(QLabel("Warn at:"), 1, 0)
        self.report_warn_pct = QDoubleSpinBox()
        self.report_warn_pct.setRange(0.0, 100000.0)
        self.report_warn_pct.setDecimals(1)
        self.report_warn_pct.setValue(5.0)
        self.report_warn_pct.setSuffix(" %")
        self.report_warn_pct.setToolTip("Amber rows in the Deltas sheet at or above this absolute change.")
        report_layout.addWidget(self.report_warn_pct, 1, 1)

        report_layout.addWidget(QLabel("Flag at:"), 1, 2)
        self.report_flag_pct = QDoubleSpinBox()
        self.report_flag_pct.setRange(0.0, 100000.0)
        self.report_flag_pct.setDecimals(1)
        self.report_flag_pct.setValue(15.0)
        self.report_flag_pct.setSuffix(" %")
        self.report_flag_pct.setToolTip("Red rows in the Deltas sheet at or above this absolute change.")
        report_layout.addWidget(self.report_flag_pct, 1, 3)

        report_layout.addWidget(QLabel("Charts:"), 2, 0)
        self.report_chart_count = QSpinBox()
        self.report_chart_count.setRange(0, 100)
        self.report_chart_count.setValue(8)
        self.report_chart_count.setToolTip("Number of highest-difference signals to include in the chart sheet.")
        report_layout.addWidget(self.report_chart_count, 2, 1)

        self.report_align_zero = QCheckBox("Align each test to T = 0 before comparing")
        self.report_align_zero.setChecked(True)
        report_layout.addWidget(self.report_align_zero, 2, 2, 1, 2)
        report_layout.addWidget(
            QLabel("The report contains Summary, Deltas, and Top Deltas - Charts sheets."),
            3, 0, 1, 4,
        )
        report_layout.addWidget(self.btn_generate_report, 4, 0, 1, 4)
        layout.addWidget(report_group)

        event_layout = QHBoxLayout()
        event_layout.addWidget(QLabel("Event trigger:"))
        self.event_rule = QComboBox()
        self.event_rule.addItems(["rising", "falling"])
        event_layout.addWidget(self.event_rule)
        self.event_level = QDoubleSpinBox()
        self.event_level.setRange(-1e9, 1e9)
        self.event_level.setValue(5.0)
        self.event_level.setSuffix(" level")
        event_layout.addWidget(self.event_level)
        self.event_align_button = QPushButton("Align to event")
        self.event_align_button.clicked.connect(self._align_to_event_all)
        event_layout.addWidget(self.event_align_button)
        layout.addLayout(event_layout)

        buttons = QDialogButtonBox(QDialogButtonBox.Close)
        buttons.rejected.connect(self.reject)
        self.btn_compare = QPushButton("Show Comparison")
        self.btn_compare.setObjectName("primaryBtn")
        self.btn_compare.clicked.connect(self._compare)
        buttons.addButton(self.btn_compare, QDialogButtonBox.ActionRole)
        layout.addWidget(buttons)

    @staticmethod
    def _offset_spin():
        spin = QDoubleSpinBox()
        spin.setRange(-1e6, 1e6)
        spin.setDecimals(3)
        spin.setSingleStep(0.1)
        spin.setSuffix(" s")
        return spin

    def _add_dataset(self):
        index = len(self.extra_datasets) + 3
        row = QHBoxLayout()
        label_widget = QLabel(f"Signal {index}:")
        search = QLineEdit()
        search.setPlaceholderText("Search signal name, file, or unit...")
        search.setClearButtonEnabled(True)
        search.setMinimumWidth(230)
        combo = QComboBox()
        combo.setMinimumWidth(420)
        self._set_combo_options(combo, self.options)
        offset = self._offset_spin()
        remove = QPushButton("Remove")
        search.textChanged.connect(
            lambda text, target=combo: self._filter_combo(target, text)
        )
        row.addWidget(label_widget)
        row.addWidget(search)
        row.addWidget(combo)
        row.addWidget(QLabel("offset:"))
        row.addWidget(offset)
        row.addWidget(remove)
        self.extra_datasets_layout.addLayout(row)
        entry = {
            "label": f"Signal {index}",
            "label_widget": label_widget,
            "search": search,
            "combo": combo,
            "offset": offset,
            "layout": row,
        }
        self.extra_datasets.append(entry)
        offset.valueChanged.connect(
            lambda value, item=entry: self._numeric_offset_changed(
                item.get("comparison_label", item["label"]), value
            )
        )
        remove.clicked.connect(lambda: self._remove_dataset(entry))

    def _remove_dataset(self, entry):
        if entry not in self.extra_datasets:
            return
        self.extra_datasets.remove(entry)
        while entry["layout"].count():
            item = entry["layout"].takeAt(0)
            widget = item.widget()
            if widget is not None:
                widget.deleteLater()
        entry["layout"].deleteLater()
        for index, current in enumerate(self.extra_datasets, start=3):
            current["label"] = f"Signal {index}"
            current["label_widget"].setText(f"Signal {index}:")

    def _numeric_offset_changed(self, label, value):
        if any(current_label == label for current_label, _ in self.main_window._speed_comparison_selection):
            self.main_window._speed_comparison_offsets[label] = float(value)
            alignment = self.main_window._speed_alignments.setdefault(label, Alignment())
            alignment.offset = float(value)
            alignment.source = "manual"
            alignment.quality = alignment.quality
            self.main_window._update_speed_comparison(auto_range=False)

    def _auto_align_all(self):
        selections = self.main_window._speed_comparison_selection
        if len(selections) < 2:
            self.lbl_result.setText("Select at least two signals first.")
            return
        ref_label, ref_selection = selections[0]
        ref_data = self.main_window._fetch_raw_signal("file", ref_selection[0], ref_selection[1])
        if ref_data is None:
            self.lbl_result.setText("The reference signal could not be loaded.")
            return
        ref_x, ref_y, _ = ref_data
        results = []
        for label, selection in selections:
            if label == ref_label:
                self.main_window._speed_alignments[label] = Alignment(
                    reference=ref_label, ref_channel=ref_selection[1]
                )
                self.main_window._speed_comparison_offsets[label] = 0.0
                continue
            existing = self.main_window._speed_alignments.get(label)
            if existing is not None and existing.locked:
                results.append(f"{label}: locked")
                continue
            data = self.main_window._fetch_raw_signal("file", selection[0], selection[1])
            if data is None:
                continue
            x, y, _ = data
            offset, quality = estimate_offset(ref_x, ref_y, x, y)
            self.main_window._speed_comparison_offsets[label] = offset
            self.main_window._speed_alignments[label] = Alignment(
                offset=offset, source="auto", quality=quality,
                reference=ref_label, ref_channel=ref_selection[1],
            )
            results.append(f"{label}: {offset:+.3f}s ({quality:.0%})")
        self.main_window._update_speed_comparison(auto_range=False)
        self.lbl_result.setText("Auto alignment: " + " | ".join(results))

    def _align_to_event_all(self):
        selections = self.main_window._speed_comparison_selection
        if len(selections) < 2:
            self.lbl_result.setText("Select at least two signals first.")
            return
        reference_label, reference_selection = selections[0]
        reference_data = self.main_window._fetch_raw_signal(
            "file", reference_selection[0], reference_selection[1]
        )
        if reference_data is None:
            return
        reference_time = trigger_time(
            reference_data[0], reference_data[1],
            self.event_rule.currentText(), self.event_level.value(),
        )
        if reference_time is None:
            self.lbl_result.setText("No event crossing found in the reference signal.")
            return
        results = []
        for label, selection in selections:
            data = self.main_window._fetch_raw_signal("file", selection[0], selection[1])
            event_time = trigger_time(
                data[0], data[1], self.event_rule.currentText(), self.event_level.value()
            ) if data is not None else None
            if event_time is None:
                results.append(f"{label}: no event")
                continue
            offset = 0.0 if label == reference_label else reference_time - event_time
            self.main_window._speed_comparison_offsets[label] = offset
            self.main_window._speed_alignments[label] = Alignment(
                offset=offset, source="event", quality=1.0,
                reference=reference_label, ref_channel=reference_selection[1],
            )
            results.append(f"{label}: {offset:+.3f}s")
        self.main_window._update_speed_comparison(auto_range=False)
        self.lbl_result.setText("Event alignment: " + " | ".join(results))

    @staticmethod
    def _selected_signal_label(combo):
        signal = combo.currentData()
        if not signal:
            return ""
        return extract_signal_only_name(signal[1])

    def _unique_signal_label(self, signal, used_labels):
        base_label = extract_signal_only_name(signal[1])
        label = base_label
        if label in used_labels:
            label = f"{base_label} [{signal[0]}]"
        suffix = 2
        while label in used_labels:
            label = f"{base_label} [{signal[0]}] #{suffix}"
            suffix += 1
        used_labels.add(label)
        return label

    def _populate_signals(self):
        self.options = []
        for file_id, record in self.main_window.loaded_files.items():
            for name in record.signal_names:
                unit = record.signal_units.get(name, "")
                label = f"{file_id}: {extract_signal_only_name(name)}"
                self.options.append((label, file_id, name, unit))
        self._set_combo_options(self.combo_test1, self.options)
        self._set_combo_options(self.combo_test2, self.options)

        speed_indices = [
            index for index, (label, _fid, _name, unit) in enumerate(self.options)
            if any(token in f"{label} {unit}".lower() for token in ("speed", "velocity", "vehicle spd"))
        ]
        if speed_indices:
            self.combo_test1.setCurrentIndex(speed_indices[0])
            self.combo_test2.setCurrentIndex(speed_indices[min(1, len(speed_indices) - 1)])

    def _set_combo_options(self, combo: QComboBox, options):
        combo.clear()
        for label, file_id, name, unit in options:
            combo.addItem(label, (file_id, name, unit))
        completer = QCompleter([option[0] for option in options], combo)
        completer.setCaseSensitivity(Qt.CaseInsensitive)
        completer.setFilterMode(Qt.MatchContains)
        combo.setCompleter(completer)

    def _filter_combo(self, combo: QComboBox, query: str):
        query = query.lower().strip()
        combo.blockSignals(True)
        combo.clear()
        filtered = [
            option for option in self.options
            if not query or query in f"{option[0]} {option[3]}".lower()
        ]
        self._set_combo_options(combo, filtered)
        if filtered:
            combo.setCurrentIndex(0)
        combo.blockSignals(False)

    def _compare(self):
        used_labels = set()
        signal1 = self.combo_test1.currentData()
        signal2 = self.combo_test2.currentData()
        selections = [
            (
                self._unique_signal_label(signal1, used_labels) if signal1 else "",
                signal1,
                self.spin_test1_offset.value(),
            ),
            (
                self._unique_signal_label(signal2, used_labels) if signal2 else "",
                signal2,
                self.spin_test2_offset.value(),
            ),
        ]
        selections.extend(
            (
                self._unique_signal_label(entry["combo"].currentData(), used_labels)
                if entry["combo"].currentData()
                else "",
                entry["combo"].currentData(),
                entry["offset"].value(),
            )
            for entry in self.extra_datasets
        )
        selections = [item for item in selections if item[1]]
        if len(selections) < 2:
            self.lbl_result.setText("Select at least two valid signals.")
            return
        self.main_window._speed_comparison_selection = [
            (label, signal) for label, signal, _offset in selections
        ]
        self._combo_labels = {
            self.combo_test1: selections[0][0],
            self.combo_test2: selections[1][0],
        }
        for entry, item in zip(self.extra_datasets, selections[2:]):
            entry["comparison_label"] = item[0]
        self.main_window._speed_comparison_offsets = {
            label: float(offset) for label, _signal, offset in selections
        }
        reference_label = selections[0][0]
        self.main_window._speed_alignments = {
            label: Alignment(offset=float(offset), source="manual" if offset else "none",
                             reference=reference_label, ref_channel=selections[0][1][1])
            for label, _signal, offset in selections
        }
        self.main_window._speed_comparison_offset = 0.0
        self.main_window._speed_alignment_method = "manual"
        self.main_window._comparison_difference_mode = bool(self.chk_difference_view.isChecked())
        self.main_window._update_speed_comparison(auto_range=True)
        self.main_window.speed_comparison_plot.show()
        self.lbl_result.setText(
            "Manual mode active. Right-drag a signal trace to move it horizontally."
        )

    def _generate_difference_report(self):
        if not self.main_window._speed_comparison_selection:
            self.lbl_result.setText("Pick at least two signals first, then generate the report.")
            return
        records = {}
        seen = set()
        for label, selection in self.main_window._speed_comparison_selection:
            file_id, name = selection[:2]
            if file_id not in records:
                records[file_id] = self.main_window.loaded_files[file_id]
                seen.add(file_id)
        if len(records) < 2:
            self.lbl_result.setText("Need signals from at least two files to compare differences.")
            return

        default_name = "difference_report.xlsx"
        save_path, _ = QFileDialog.getSaveFileName(
            self,
            "Save difference report",
            default_name,
            "Excel Files (*.xlsx)",
        )
        if not save_path:
            return

        reference_id = next(iter(records))
        selected_reference = self.report_reference.currentText()
        if selected_reference in records:
            reference_id = selected_reference
        else:
            self.report_reference.setCurrentText(reference_id)
        try:
            result = generate_comparison_report(
                records,
                save_path,
                reference_id=self.report_reference.currentText() or reference_id,
                top_n_charts=self.report_chart_count.value(),
                warn_pct=self.report_warn_pct.value(),
                flag_pct=self.report_flag_pct.value(),
                align_to_zero=self.report_align_zero.isChecked(),
            )
            self.lbl_result.setText(
                f"Difference report saved: {result['output_path']} | "
                f"Matched signals: {result['signals_matched']}"
            )
        except Exception as exc:
            self.lbl_result.setText(f"Report generation failed: {exc}")
            QMessageBox.warning(self, "Difference report", str(exc))

    def _reset_offsets(self):
        for spin in [self.spin_test1_offset, self.spin_test2_offset]:
            spin.setValue(0.0)
        for entry in self.extra_datasets:
            entry["offset"].setValue(0.0)
        self.main_window._speed_comparison_offsets = {
            label: 0.0 for label, _signal in self.main_window._speed_comparison_selection
        }
        self.main_window._speed_alignments = {
            label: Alignment(reference=self.main_window._speed_comparison_selection[0][0])
            for label, _signal in self.main_window._speed_comparison_selection
        }
        self.main_window._speed_comparison_offset = 0.0
        self.main_window._update_speed_comparison(auto_range=True)
        self.lbl_result.setText("Offsets reset to 0.000 s for all selected signals.")


class MF4VehicleExplorer(QMainWindow):
    def __init__(self):
        super().__init__()
        self.theme_manager = ThemeManager()
        self._theme_actions: Dict[str, QAction] = {}
        self.setWindowTitle("MF4 / BLF Vehicle Signal Explorer & Power Calculator")
        self.resize(1440, 900)
        self.setAcceptDrops(True)

        self.loaded_files: Dict[str, MDFFileRecord] = {}
        self.dbc_paths: List[str] = []
        self._load_thread: Optional[QThread] = None
        self._load_worker: Optional[FileLoadWorker] = None
        self._queued_load_paths: List[str] = []
        self._load_cancelled = False
        self._blf_prompt_shown = False
        self.derived_signals: Dict[str, Dict[str, Any]] = {}
        self.max_plot_points: int = 25_000
        self.export_raster: float = 0.01
        self.dark_mode: bool = self.theme_manager.is_dark
        self.align_to_zero: bool = True

        self._overlay_plot: Optional[pg.PlotItem] = None
        self._overlay_legend: Optional[pg.LegendItem] = None
        self._overlay_unit_main: Optional[str] = None
        self._overlay_extra_viewboxes: Dict[str, Tuple[pg.ViewBox, pg.AxisItem]] = {}
        self._plotted_curves: List[Dict[str, Any]] = []
        self._separate_plots: List[pg.PlotItem] = []
        self._last_overlay_mode: Optional[bool] = None
        self._cursor_lines: Dict[pg.PlotItem, pg.InfiniteLine] = {}
        self._cursor_value_labels: Dict[pg.PlotItem, pg.TextItem] = {}
        self._cursor_delta_lines: Dict[pg.PlotItem, pg.InfiniteLine] = {}
        self._cursor_delta_labels: Dict[pg.PlotItem, pg.TextItem] = {}
        self._cursor_horizontal_lines: Dict[pg.PlotItem, pg.InfiniteLine] = {}
        self._cursor_delta_horizontal_lines: Dict[pg.PlotItem, pg.InfiniteLine] = {}
        self._plot_signal_map: Dict[pg.PlotItem, Dict[str, Any]] = {}
        self._plot_legend_map: Dict[pg.PlotItem, pg.LegendItem] = {}
        self._legend_signal_map: Dict[QObject, Dict[str, Any]] = {}
        self._comparison_plot: Optional[pg.PlotItem] = None
        self._comparison_curves: List[Dict[str, Any]] = []
        self._comparison_difference_mode: bool = False
        self._speed_comparison_offset: float = 0.0
        self._speed_comparison_offsets: Dict[str, float] = {
            "Test 1": 0.0,
            "Test 2": 0.0,
        }
        self._speed_alignments: Dict[str, Alignment] = {}
        self._speed_alignment_method: str = "trend"
        self._speed_comparison_selection = []
        self.cursor_visible: bool = False
        self.cursor_locked: bool = False
        self.delta_cursor_visible: bool = False
        self.cursor1_x: Optional[float] = None
        self.cursor2_x: Optional[float] = None
        self.cursor1_y: Optional[float] = None
        self.cursor2_y: Optional[float] = None
        self._is_dragging_line: bool = False
        self.x_zoom_only: bool = True

        self._build_ui()
        self._create_menus()
        self.plot_canvas.scene().sigMouseMoved.connect(self.update_cursor)
        self.plot_canvas.scene().sigMouseClicked.connect(self.toggle_cursor_lock)
        self.apply_theme()

    def _build_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        main_layout = QVBoxLayout(central)
        main_layout.setContentsMargins(14, 14, 14, 14)
        main_layout.setSpacing(12)

        top_bar = self._create_top_bar()
        main_layout.addLayout(top_bar)

        self.main_splitter = QSplitter(Qt.Horizontal)

        left_pane = self._create_sidebar()
        self.main_splitter.addWidget(left_pane)

        right_pane = self._create_plot_stage()
        self.main_splitter.addWidget(right_pane)

        left_pane.setMinimumWidth(300)
        right_pane.setMinimumWidth(520)
        self.main_splitter.setStretchFactor(0, 0)
        self.main_splitter.setStretchFactor(1, 1)
        self.main_splitter.setSizes([420, 1020])
        main_layout.addWidget(self.main_splitter)

        self.progress_bar = QProgressBar()
        self.progress_bar.setRange(0, 100)
        self.progress_bar.setValue(0)
        self.progress_bar.setFixedHeight(16)
        self.progress_bar.setFixedWidth(220)
        self.progress_bar.setTextVisible(True)
        self.progress_bar.hide()
        self.statusBar().addPermanentWidget(self.progress_bar)
        self.cancel_load_button = QToolButton()
        self.cancel_load_button.setText("×")
        self.cancel_load_button.setToolTip("Cancel the current file load")
        self.cancel_load_button.setObjectName("cancelLoadButton")
        self.cancel_load_button.clicked.connect(self.cancel_loading)
        self.cancel_load_button.hide()
        self.statusBar().addPermanentWidget(self.cancel_load_button)

        self.statusBar().showMessage("Ready — Open or Drag & Drop MF4, CSV, or BLF + DBC files")

    def _create_menus(self) -> None:
        menu_bar = self.menuBar()
        file_menu = menu_bar.addMenu("&File")
        open_action = file_menu.addAction("Open Files…")
        open_action.triggered.connect(self.open_files_dialog)
        file_menu.addSeparator()
        export_action = file_menu.addAction("Export Plotted CSV")
        export_action.triggered.connect(self.export_plotted_csv)
        file_menu.addSeparator()
        file_menu.addAction("Quit", self.close, QKeySequence.Quit)

        tools_menu = menu_bar.addMenu("&Tools")
        tools_menu.addAction("Power Calculator", self.open_power_calculator)
        tools_menu.addAction("Compare Tests", self.open_speed_comparison)
        tools_menu.addAction("Clear Plot", self.clear_plots)

        settings_menu = menu_bar.addMenu("&Settings")
        theme_menu = settings_menu.addMenu("Theme")
        for theme_name, (label, _background, _is_dark) in ThemeManager.THEMES.items():
            action = theme_menu.addAction(label)
            action.setCheckable(True)
            action.setChecked(theme_name == self.theme_manager.name)
            action.triggered.connect(
                lambda checked=False, name=theme_name: self.set_theme(name)
            )
            self._theme_actions[theme_name] = action

        command_action = QAction("Command Palette", self)
        command_action.setShortcut(QKeySequence("Ctrl+Shift+P"))
        command_action.triggered.connect(self.open_command_palette)
        self.addAction(command_action)

    def open_command_palette(self) -> None:
        commands = [
            ("Open Files", self.open_files_dialog),
            ("Open Power Calculator", self.open_power_calculator),
            ("Compare Tests", self.open_speed_comparison),
            ("Generate Difference Report", self.open_speed_comparison),
            ("Export CSV", self.export_plotted_csv),
            ("Clear Plot", self.clear_plots),
        ]
        CommandPalette(commands, self).exec()

    def _create_top_bar(self) -> QHBoxLayout:
        top = QHBoxLayout()
        top.setSpacing(8)

        self.btn_open = QPushButton("+ Open Files")
        self.btn_open.setObjectName("primaryBtn")
        self.btn_open.setToolTip("Add MF4 / MDF / CSV / XLSX / BLF / DBC files")
        self.btn_open.clicked.connect(self.open_files_dialog)
        top.addWidget(self.btn_open)

        self.btn_add_dbc = QPushButton("+ Add DBC")
        self.btn_add_dbc.setToolTip("Upload DBC CAN database files used to decode BLF logs")
        self.btn_add_dbc.clicked.connect(self.add_dbc_dialog)
        top.addWidget(self.btn_add_dbc)

    

        self.btn_compare_speed = QPushButton("Time Offset")
        self.btn_compare_speed.setObjectName("accentBtn")
        self.btn_compare_speed.setToolTip("Open the speed comparison tool and find the Test 2 time offset")
        self.btn_compare_speed.clicked.connect(self.open_speed_comparison)
        top.addWidget(self.btn_compare_speed)


        self.btn_clear_plot = QPushButton("Clear Plot")
        self.btn_clear_plot.setToolTip("Reset canvas and clear plotted curves")
        self.btn_clear_plot.clicked.connect(self.clear_plots)
        top.addWidget(self.btn_clear_plot)

        top.addSpacing(12)

        lbl_mode = QLabel("Layout:")
        lbl_mode.setStyleSheet("color: #86868B; font-weight: 500;")
        top.addWidget(lbl_mode)

        self.combo_plot_mode = QComboBox()
        self.combo_plot_mode.addItems(["Overlay (Multi-Y)", "Stacked Subplots"])
        self.combo_plot_mode.currentIndexChanged.connect(lambda: self.plot_current_curves(rebuild_layout=True))
        top.addWidget(self.combo_plot_mode)

        self.chk_align_zero = QCheckBox("Align T = 0s")
        self.chk_align_zero.setChecked(True)
        self.chk_align_zero.setToolTip("Shift each file's start time to 0s for direct run comparison")
        self.chk_align_zero.toggled.connect(self.on_align_zero_toggled)
        top.addWidget(self.chk_align_zero)


        self.chk_cursor = QCheckBox("Cursor")
        self.chk_cursor.setToolTip("Show the vertical cursor and nearest signal values")
        self.chk_cursor.toggled.connect(self.set_cursor_visible)
        top.addWidget(self.chk_cursor)

        self.chk_horizontal_cursor = QCheckBox("Horizontal cursor")
        self.chk_horizontal_cursor.setToolTip("Show the horizontal cursor and Y position")
        self.chk_horizontal_cursor.toggled.connect(self.set_horizontal_cursor_visible)
        top.addWidget(self.chk_horizontal_cursor)

        self.chk_delta_cursor = QCheckBox("2nd Cursor")
        self.chk_delta_cursor.setToolTip("Show a second cursor and calculate Delta X and Delta Y")
        self.chk_delta_cursor.toggled.connect(self.set_delta_cursor_visible)
        top.addWidget(self.chk_delta_cursor)

    
        top.addStretch()

        self.btn_export_csv = QPushButton("Export CSV")
        self.btn_export_csv.setToolTip("Resample and export plotted signals to a clean, fast CSV")
        self.btn_export_csv.clicked.connect(self.export_plotted_csv)
        top.addWidget(self.btn_export_csv)

        self.btn_theme = QPushButton("Light" if self.dark_mode else "Dark")
        self.btn_theme.clicked.connect(self.toggle_theme)
        top.addWidget(self.btn_theme)

        return top

    def _create_sidebar(self) -> QWidget:
        sidebar = QFrame()
        sidebar.setObjectName("sidebar")
        layout = QVBoxLayout(sidebar)
        layout.setContentsMargins(12, 12, 12, 12)
        layout.setSpacing(10)

        # Loaded DBC Databases Status Box
        dbc_header = QHBoxLayout()
        lbl_dbc = QLabel("Active DBC CAN Databases")
        lbl_dbc.setStyleSheet("font-weight: 600; font-size: 11px; text-transform: uppercase; color: #30D158;")
        dbc_header.addWidget(lbl_dbc)
        dbc_header.addStretch()
        self.lbl_dbc_count = QLabel("0 DBCs")
        self.lbl_dbc_count.setStyleSheet("color: #86868B; font-size: 11px;")
        dbc_header.addWidget(self.lbl_dbc_count)
        layout.addLayout(dbc_header)

        self.dbc_list = QListWidget()
        self.dbc_list.setMaximumHeight(65)
        layout.addWidget(self.dbc_list)

        dbc_actions = QHBoxLayout()
        self.btn_sidebar_add_dbc = QPushButton("+ Add DBC")
        self.btn_sidebar_add_dbc.clicked.connect(self.add_dbc_dialog)
        dbc_actions.addWidget(self.btn_sidebar_add_dbc)

    
        self.btn_clear_dbc = QPushButton("Clear DBCs")
        self.btn_clear_dbc.clicked.connect(self.clear_dbc_paths)
        dbc_actions.addWidget(self.btn_clear_dbc)
        layout.addLayout(dbc_actions)

        # Loaded Files Section
        files_header = QHBoxLayout()
        lbl_files = QLabel("Loaded Measurement Files")
        lbl_files.setStyleSheet("font-weight: 600; font-size: 11px; text-transform: uppercase; color: #86868B;")
        files_header.addWidget(lbl_files)
        files_header.addStretch()
        self.lbl_file_count = QLabel("0 files")
        self.lbl_file_count.setStyleSheet("color: #86868B; font-size: 11px;")
        files_header.addWidget(self.lbl_file_count)
        layout.addLayout(files_header)

        self.file_list = QListWidget()
        self.file_list.setMaximumHeight(85)
        self.file_list.itemSelectionChanged.connect(self.on_file_selected)
        layout.addWidget(self.file_list)

        file_actions = QHBoxLayout()
        self.btn_load_new = QPushButton("Load New")
        self.btn_load_new.setToolTip("Close all loaded files and open a new selection")
        self.btn_load_new.clicked.connect(self.load_new_files_dialog)
        file_actions.addWidget(self.btn_load_new)

        self.btn_remove_file = QPushButton("Remove Selected")
        self.btn_remove_file.setToolTip("Remove the selected file and its signals")
        self.btn_remove_file.clicked.connect(self.remove_selected_file)
        file_actions.addWidget(self.btn_remove_file)
        layout.addLayout(file_actions)

        # Power Calculator quick button

        self.search_box = QLineEdit()
        self.search_box.setPlaceholderText("Filter signals by name (e.g. BMS, volt, curr, speed)...")
        self.search_box.setClearButtonEnabled(True)
        self.search_box.textChanged.connect(self.filter_signals)
        layout.addWidget(self.search_box)

        self.lbl_search_results = QLabel("All signals")
        self.lbl_search_results.setStyleSheet("color: #86868B; font-size: 10px;")
        layout.addWidget(self.lbl_search_results)

        drag_hint = QLabel("Drag a signal onto the plot or select to plot:")
        drag_hint.setStyleSheet("color: #86868B; font-size: 11px;")
        layout.addWidget(drag_hint)
        self.signal_tree = DraggableSignalTree()
        layout.addWidget(self.signal_tree)

        return sidebar

    def _create_plot_stage(self) -> QWidget:
        container = QWidget()
        layout = QVBoxLayout(container)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(6)

        self.plot_canvas = DropGraphicsLayoutWidget()
        self.plot_canvas.signals_dropped.connect(self.handle_signals_dropped)
        self.plot_canvas.files_dropped.connect(self.handle_files_dropped)
        layout.addWidget(self.plot_canvas)

        self.speed_comparison_plot = OffsetComparisonPlot()
        self.speed_comparison_plot.offsetDragged.connect(
            self._on_comparison_offset_dragged
        )
        self.speed_comparison_plot.setMinimumHeight(260)
        self._comparison_plot = self.speed_comparison_plot.getPlotItem()
        self._style_plot(self._comparison_plot, title="Signal Comparison | Manual offsets")
        self._comparison_plot.addLegend(offset=(10, 10), labelTextColor=self.current_fg())
        self._comparison_plot.setLabel("left", "Speed")
        self._comparison_plot.setLabel("bottom", "Time", units="s")
        self._install_cursor(self._comparison_plot)
        layout.addWidget(self.speed_comparison_plot)
        self.speed_comparison_plot.hide()

        return container

    def _on_comparison_offset_dragged(self, label: str, offset: float):
        self._speed_comparison_offsets[label] = float(offset)
        alignment = self._speed_alignments.setdefault(label, Alignment())
        alignment.offset = float(offset)
        alignment.source = "manual"
        self._speed_comparison_offset = self._speed_comparison_offsets.get(
            "Test 2", 0.0
        )
        self._speed_alignment_method = "manual"
        trace = next(
            (item for item in self._comparison_curves if item["label"] == label),
            None,
        )
        if trace is not None:
            # Updating curve data must not let pyqtgraph auto-range away from
            # the second currently being used as the visual alignment point.
            view_box = self._comparison_plot.getViewBox()
            x_range, y_range = view_box.viewRange()
            view_box.disableAutoRange()
            trace["x_offset"] = float(offset)
            trace["curve"].setData(
                trace["base_x"] + float(offset),
                trace["y"],
            )
            view_box.setRange(xRange=x_range, yRange=y_range, padding=0.0)
            for target in self.speed_comparison_plot._trace_data:
                if target["label"] == label:
                    target["offset"] = float(offset)
                    target["x"] = trace["base_x"] + float(offset)
                    break
            self._comparison_plot.setTitle(
                self._comparison_title(),
                color=self.current_fg(),
                size="11pt",
            )
        self.statusBar().showMessage(
            f"{label} offset: {offset:+.3f} s — right-drag to continue aligning"
        )

    def _comparison_title(self):
        offset_summary = " | ".join(
            f"{label}: {value:+.3f} s"
            + (f" ({self._speed_alignments[label].quality:.0%})"
               if label in self._speed_alignments and self._speed_alignments[label].quality
               else "")
            for label, value in self._speed_comparison_offsets.items()
        )
        return f"Signal Comparison | {offset_summary}"

    def alignment_snapshot(self):
        """Return JSON-serializable alignment provenance for session/report export."""
        return {
            label: alignment.as_dict()
            for label, alignment in self._speed_alignments.items()
        }

    def save_alignment_snapshot(self, path: str):
        payload = {
            "files": {
                file_id: {
                    "path": record.path,
                    "sha256": file_hash(record.path),
                }
                for file_id, record in self.loaded_files.items()
                if Path(record.path).is_file()
            },
            "reference": (
                self._speed_comparison_selection[0][0]
                if self._speed_comparison_selection else None
            ),
            "alignments": self.alignment_snapshot(),
        }
        Path(path).write_text(json.dumps(payload, indent=2, sort_keys=True), encoding="utf-8")

    def open_power_calculator(self):
        if not self.loaded_files and not self.derived_signals:
            QMessageBox.information(
                self, "No Signals Available",
                "Please open an MF4, CSV, or BLF file with voltage and current signals first."
            )
            return
        dialog = PowerEnergyCalculatorDialog(self, self)
        dialog.exec()

    def open_nvh_analysis(self):
        records = [record for record in self.loaded_files.values()
                   if record.csv_data is not None or record.chn_data is not None]
        if not records:
            QMessageBox.information(
                self, "No NVH Measurements",
                "Load CSV/XLSX measurements or a CHN file before opening NVH analysis."
            )
            return
        dialog = EngineMountNvhDialog(self, records, self.dark_mode)
        dialog.exec()

    def open_speed_comparison(self):
        if len(self.loaded_files) < 1:
            QMessageBox.information(
                self,
                "No Signals Available",
                "Please open at least one measurement file before comparing speed signals.",
            )
            return
        dialog = SpeedComparisonDialog(self, self)
        self._speed_comparison_dialog = dialog
        dialog.show()

    def toggle_theme(self):
        names = list(ThemeManager.THEMES)
        next_name = names[(names.index(self.theme_manager.name) + 1) % len(names)]
        self.set_theme(next_name)

    def set_theme(self, name: str):
        self.theme_manager.set_theme(name)
        self.dark_mode = self.theme_manager.is_dark
        self.btn_theme.setText(self.theme_manager.display_name)
        for theme_name, action in self._theme_actions.items():
            action.setChecked(theme_name == name)
        self.apply_theme()

    def apply_theme(self):
        self.theme_manager.apply(QApplication.instance(), self)
        self.plot_canvas.setBackground(self.theme_manager.background)
        self.plot_current_curves(rebuild_layout=True)

    def current_palette(self) -> List[str]:
        if self.theme_manager.name == "high-contrast":
            return ["#FFFF00", "#00FFFF", "#00FF00", "#FF00FF", "#FFFFFF"]
        if self.theme_manager.name == "solarized-dark":
            return ["#268BD2", "#2AA198", "#B58900", "#DC322F", "#6C71C4"]
        return PALETTE_DARK if self.dark_mode else PALETTE_LIGHT

    def current_fg(self) -> str:
        return "#FFFFFF" if self.theme_manager.name == "high-contrast" else (
            DARK_TEXT_PRIMARY if self.dark_mode else LIGHT_TEXT_PRIMARY
        )

    def current_border(self) -> str:
        if self.theme_manager.name == "high-contrast":
            return "#FFFFFF"
        if self.theme_manager.name == "solarized-dark":
            return "#586E75"
        return DARK_BORDER if self.dark_mode else LIGHT_BORDER

    def open_files_dialog(self):
        paths, _ = QFileDialog.getOpenFileNames(
            self, "Open Measurement and DBC Files", "",
            "Supported Files (*.mf4 *.mdf *.dat *.csv *.xlsx *.xls *.txt *.blf *.chn *.dbc);;All Files (*.*)"
        )
        if paths:
            self.load_files(paths)

    def load_chn_database_dialog(self):
        paths, _ = QFileDialog.getOpenFileNames(
            self, "Load CHN Database", "",
            "CHN files (*.chn *.CHN);;Metadata CSV (*.csv);;All Files (*.*)"
        )
        if paths:
            self.load_files(paths)

    def open_excel_dialog(self):
        paths, _ = QFileDialog.getOpenFileNames(
            self, "Open Excel Measurement File", "",
            "Microsoft Excel (*.xlsx *.xls);;All Files (*.*)"
        )
        if paths:
            self.load_files(paths)

    def add_dbc_dialog(self):
        paths, _ = QFileDialog.getOpenFileNames(
            self, "Add DBC CAN Database Files", "", "CAN Database Files (*.dbc);;All Files (*.*)"
        )
        if not paths:
            return
        self.add_dbc_paths(paths)

    def clear_dbc_paths(self):
        self.dbc_paths.clear()
        self.refresh_dbc_list()
        self.statusBar().showMessage("Cleared active DBC databases.")

    def refresh_dbc_list(self):
        self.dbc_list.clear()
        for p in self.dbc_paths:
            item = QListWidgetItem(Path(p).name)
            item.setToolTip(p)
            self.dbc_list.addItem(item)
        self.lbl_dbc_count.setText(f"{len(self.dbc_paths)} DBC{'s' if len(self.dbc_paths) != 1 else ''}")

    def load_new_files_dialog(self):
        paths, _ = QFileDialog.getOpenFileNames(
            self, "Load New Measurement and DBC Files", "",
            "Supported Files (*.mf4 *.mdf *.dat *.csv *.xlsx *.xls *.txt *.blf *.chn *.dbc);;All Files (*.*)"
        )
        if not paths:
            return
        self.load_files(paths)

    def clear_loaded_files(self):
        for rec in self.loaded_files.values():
            rec.close()
        self.loaded_files.clear()
        self.refresh_file_list()
        self.populate_signal_tree()
        self.lbl_file_count.setText("0 loaded")

    def remove_selected_file(self):
        item = self.file_list.currentItem()
        if item is None:
            QMessageBox.information(self, "Remove File", "Select a loaded file first.")
            return

        file_id = item.data(Qt.UserRole)
        if not file_id:
            QMessageBox.information(self, "Remove File", "Select a specific loaded file to remove.")
            return
        record = self.loaded_files.pop(file_id, None)
        if record is None:
            return
        record.close()
        self._plotted_curves = [
            sig for sig in self._plotted_curves if sig.get("file_id") != file_id
        ]
        self._reset_plot_layout()
        self.refresh_file_list()
        self.populate_signal_tree()
        self.lbl_file_count.setText(f"{len(self.loaded_files)} loaded")
        self.statusBar().showMessage(f"Removed {file_id}. All remaining signals are available.")

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()
        else:
            event.ignore()

    def dropEvent(self, event):
        if event.mimeData().hasUrls():
            paths = [u.toLocalFile() for u in event.mimeData().urls() if Path(u.toLocalFile()).suffix.lower() in [".mf4", ".mdf", ".dat", ".csv", ".xlsx", ".xls", ".txt", ".blf", ".chn", ".dbc"]]
            if paths:
                self.load_files(paths)
                event.acceptProposedAction()

    def handle_files_dropped(self, paths: List[str]):
        self.load_files(paths)

    def load_files(self, paths: List[str]):
        paths = list(dict.fromkeys(str(Path(path)) for path in paths))
        if self._load_thread is not None and self._load_thread.isRunning():
            self._queued_load_paths.extend(
                path for path in paths if path not in self._queued_load_paths
            )
            self.statusBar().showMessage(
                f"Queued {len(paths)} file(s); current upload will finish first."
            )
            return
        # Extract DBC files first if provided
        dbc_paths = [path for path in paths if Path(path).suffix.lower() == ".dbc"]
        if dbc_paths:
            self.add_dbc_paths(dbc_paths, auto_rescan=False)
            paths = [path for path in paths if Path(path).suffix.lower() != ".dbc"]
        if not paths:
            return

        # Check if BLF files are included without a DBC uploaded
        blf_present = any(Path(path).suffix.lower() == ".blf" for path in paths)
        if blf_present and not self.dbc_paths:
            reply = QMessageBox.question(
                self,
                "DBC Required for BLF",
                "Opening BLF files requires a DBC database to decode CAN signals into names and units.\n\n"
                "Would you like to select a DBC file now?",
                QMessageBox.Yes | QMessageBox.No,
                QMessageBox.Yes
            )
            if reply == QMessageBox.Yes:
                dbc_select_paths, _ = QFileDialog.getOpenFileNames(
                    self, "Select CAN Database (DBC)", "", "CAN Database (*.dbc);;All Files (*.*)"
                )
                if dbc_select_paths:
                    self.add_dbc_paths(dbc_select_paths, auto_rescan=False)
                else:
                    return
            else:
                return

        palette = self.current_palette()
        jobs = []
        reserved_ids = set(self.loaded_files)
        existing_by_path = {
            os.path.normcase(os.path.abspath(record.path)): file_id
            for file_id, record in self.loaded_files.items()
        }
        for p in paths:
            if not os.path.exists(p):
                continue
            normalized_path = os.path.normcase(os.path.abspath(p))
            file_id = existing_by_path.get(normalized_path, Path(p).name)
            count = 1
            original_id = file_id
            while file_id in reserved_ids and existing_by_path.get(normalized_path) != file_id:
                file_id = f"{original_id} ({count})"
                count += 1

            reserved_ids.add(file_id)
            color_tag = palette[(len(self.loaded_files) + len(jobs)) % len(palette)]
            jobs.append((file_id, p, color_tag, list(self.dbc_paths)))

        if not jobs:
            return
        self.btn_open.setEnabled(False)
        self.btn_add_dbc.setEnabled(False)
        self.btn_load_new.setEnabled(False)
        self.progress_bar.show()
        self._load_cancelled = False
        self.cancel_load_button.setEnabled(True)
        self.cancel_load_button.show()
        self.progress_bar.setValue(0)
        self.progress_bar.setFormat("Loading files... 0%")
        self.statusBar().showMessage(f"Reading {len(jobs)} file(s) in the background...")
        self._load_thread = QThread(self)
        self._load_worker = FileLoadWorker(jobs)
        self._load_worker.moveToThread(self._load_thread)
        self._load_thread.started.connect(self._load_worker.run)
        self._load_worker.progress.connect(self._on_load_progress)
        self._load_worker.finished.connect(self._finish_loading_files)
        self._load_worker.cancelled.connect(
            lambda: self.statusBar().showMessage("Load cancelled. Existing files were kept.")
        )
        self._load_worker.finished.connect(self._load_thread.quit)
        self._load_worker.finished.connect(self._load_worker.deleteLater)
        self._load_thread.finished.connect(self._load_thread.deleteLater)
        self._load_thread.finished.connect(self._clear_load_state)
        self._load_thread.start()

    def _on_load_progress(self, msg: str, percent: int):
        self.progress_bar.show()
        self.progress_bar.setValue(percent)
        self.progress_bar.setFormat(f"%p% - {msg}" if percent < 100 else "100% - Finished loading")
        self.statusBar().showMessage(f"[{percent}%] {msg}")

    def _finish_loading_files(self, records, errors):
        retained_curves = list(self._plotted_curves)
        loaded_measurement = bool(records)
        for file_id, record in records.items():
            old_record = self.loaded_files.get(file_id)
            if old_record is not None and old_record is not record:
                old_record.close()
        self.loaded_files.update(records)
        chn_metadata = {}
        for record in self.loaded_files.values():
            if record.chn_metadata:
                chn_metadata.update(record.chn_metadata)
        if chn_metadata:
            for record in self.loaded_files.values():
                if record.csv_data is not None and record.signal_names:
                    record.apply_chn_metadata(chn_metadata)
        for filename, error in errors:
            QMessageBox.critical(self, "Load Error", f"Failed to load {filename}:\n{error}")
        self.refresh_file_list()
        # Keep the user's active query while loading. The tree is rebuilt in
        # full below, then the same query is applied to old and new signals.
        self.populate_signal_tree()
        self._plotted_curves = [
            signal for signal in retained_curves
            if signal.get("type", "file") == "derived"
            or (
                signal.get("file_id") in self.loaded_files
                and signal.get("name") in self.loaded_files[signal["file_id"]].signal_names
            )
        ]
        if self._plotted_curves:
            self.plot_current_curves(rebuild_layout=True)
        self.lbl_file_count.setText(f"{len(self.loaded_files)} loaded")
        decoded = [
            f"{record.filename}: {len(record.signal_names)} signals"
            for record in records.values()
            if record.blf_data is not None
        ]
        summary = f" ({'; '.join(decoded)})" if decoded else ""
        self.progress_bar.setValue(100)
        self.progress_bar.setFormat("100% - Finished loading")
        if self._load_cancelled:
            self.statusBar().showMessage(
                f"Load cancelled. Kept {len(records)} file(s) loaded before cancellation."
            )
        else:
            self.statusBar().showMessage(
                f"Successfully loaded {len(records)} file(s). Ready to plot.{summary}"
            )
        QTimer.singleShot(4000, self._hide_progress_bar_if_idle)

    def _hide_progress_bar_if_idle(self):
        if self._load_thread is None:
            self.progress_bar.hide()

    def _clear_load_state(self):
        self.btn_open.setEnabled(True)
        self.btn_add_dbc.setEnabled(True)
        self.btn_load_new.setEnabled(True)
        self._load_thread = None
        self._load_worker = None
        self.cancel_load_button.hide()
        queued = list(dict.fromkeys(self._queued_load_paths))
        self._queued_load_paths.clear()
        if queued:
            QTimer.singleShot(0, lambda: self.load_files(queued))

    def cancel_loading(self):
        if self._load_worker is None:
            return
        self._load_worker.request_cancel()
        self._load_cancelled = True
        self.cancel_load_button.setEnabled(False)
        self.progress_bar.setFormat("Cancelling... current file will finish safely")
        self.statusBar().showMessage("Cancelling load...")

    def add_dbc_paths(self, paths: List[str], auto_rescan: bool = True):
        valid_paths = []
        for path in paths:
            try:
                # Test load with robust multi-encoding
                db = load_dbc_robust(path)
                valid_paths.append(path)
            except Exception as exc:
                QMessageBox.critical(
                    self,
                    "Invalid DBC",
                    f"Failed to read {Path(path).name}:\n{exc}",
                )
        if not valid_paths:
            return

        self.dbc_paths = list(dict.fromkeys(valid_paths + self.dbc_paths))
        self.refresh_dbc_list()

        # Check if any BLF files are already loaded in workspace
        blf_records = [rec for rec in self.loaded_files.values() if Path(rec.path).suffix.lower() == ".blf"]
        if blf_records and auto_rescan:
            total_rescanned = 0
            for blf_rec in blf_records:
                try:
                    count = blf_rec.rescan_blf(self.dbc_paths)
                    total_rescanned += count
                except Exception as exc:
                    print(f"Error rescanning BLF {blf_rec.filename}: {exc}")
            self.refresh_file_list()
            self.populate_signal_tree()
            QMessageBox.information(
                self,
                "BLF Rescanned with DBC",
                f"Added {len(valid_paths)} DBC database(s):\n"
                + "\n".join(f"• {Path(p).name}" for p in valid_paths)
                + f"\n\nRescanned {len(blf_records)} loaded BLF file(s). "
                f"Identified and decoded {total_rescanned} CAN signals across BLF logs!"
            )
            self.statusBar().showMessage(
                f"DBC ready: Rescanned BLF with {total_rescanned} active signals."
            )
        elif not blf_records and auto_rescan and not self._blf_prompt_shown:
            # If no BLF files are loaded yet, offer to select a BLF file to scan right now
            self._blf_prompt_shown = True
            reply = QMessageBox.question(
                self,
                "Scan BLF with DBC",
                f"Loaded DBC database: {Path(valid_paths[0]).name}\n\n"
                "Would you like to select a Vector BLF file now to scan and read its signals with this DBC?",
                QMessageBox.Yes | QMessageBox.No,
                QMessageBox.Yes,
            )
            if reply == QMessageBox.Yes:
                blf_paths, _ = QFileDialog.getOpenFileNames(
                    self, "Select Vector BLF File to Scan", "", "Vector BLF (*.blf);;All Files (*.*)"
                )
                if blf_paths:
                    self.load_files(blf_paths)
            else:
                self.statusBar().showMessage(
                    f"DBC database ready ({len(self.dbc_paths)} active). Open a BLF file or click 'Scan BLF with DBC'."
                )

    def scan_blf_with_dbc_dialog(self):
        """Dedicated workflow: Select DBC and scan through BLF file to read and extract signals."""
        if not self.dbc_paths:
            reply = QMessageBox.question(
                self,
                "Select DBC First",
                "No CAN DBC database is currently loaded.\n\n"
                "To scan and decode signals from a BLF file, a DBC database is required.\n"
                "Would you like to select a DBC file first?",
                QMessageBox.Yes | QMessageBox.No,
                QMessageBox.Yes
            )
            if reply != QMessageBox.Yes:
                return
            dbc_paths, _ = QFileDialog.getOpenFileNames(
                self, "Select CAN Database (DBC)", "", "CAN Database (*.dbc);;All Files (*.*)"
            )
            if not dbc_paths:
                return
            self.add_dbc_paths(dbc_paths, auto_rescan=False)

        blf_paths, _ = QFileDialog.getOpenFileNames(
            self, "Select Vector BLF File to Scan", "", "Vector BLF (*.blf);;All Files (*.*)"
        )
        if not blf_paths:
            return
        self.load_files(blf_paths)

    def refresh_file_list(self):
        self.file_list.clear()
        all_item = QListWidgetItem("All loaded files")
        all_item.setData(Qt.UserRole, None)
        all_item.setToolTip("Show signals from every loaded file")
        self.file_list.addItem(all_item)
        for fid, rec in self.loaded_files.items():
            source_type = getattr(rec, "source_type", Path(rec.path).suffix.upper().lstrip("."))
            item = QListWidgetItem(
                f"{fid}  ·  {source_type}  ·  {len(rec.signal_names)} readable signals"
            )
            item.setData(Qt.UserRole, fid)
            item.setForeground(QColor(rec.color_tag))
            item.setToolTip(
                f"File: {rec.filename}\nFormat: {source_type}\n"
                f"Readable signals: {len(rec.signal_names)}"
            )
            self.file_list.addItem(item)
        if self.file_list.count() > 0:
            self.file_list.setCurrentRow(0)
        self._refresh_speed_comparison_signals()

    def _refresh_speed_comparison_signals(self):
        invalid = any(
            signal and signal[0] not in self.loaded_files
            for _label, signal in self._speed_comparison_selection
        )
        if invalid:
            self._speed_comparison_selection = []
            self.speed_comparison_plot.hide()

    def _find_speed_offset(
        self,
        test1=None,
        test2=None,
        method: str = "trend",
        max_shift: Optional[float] = None,
    ):
        stored = self._speed_comparison_selection
        test1 = test1 or (stored[0][1] if len(stored) > 0 else None)
        test2 = test2 or (stored[1][1] if len(stored) > 1 else None)
        if not test1 or not test2:
            return False

        data1 = self._fetch_raw_signal("file", test1[0], test1[1])
        data2 = self._fetch_raw_signal("file", test2[0], test2[1])
        if data1 is None or data2 is None:
            return False

        x1, y1, _ = data1
        x2, y2, _ = data2
        if len(x1) < 2 or len(x2) < 2:
            return False

        def prepare(x_values, y_values):
            order = np.argsort(np.asarray(x_values, dtype=float))
            x_values = np.asarray(x_values, dtype=float)[order]
            y_values = np.asarray(y_values, dtype=float)[order]
            valid = np.isfinite(x_values) & np.isfinite(y_values)
            x_values = x_values[valid]
            y_values = y_values[valid]
            if len(x_values) < 2:
                return None, None
            unique, indices = np.unique(x_values, return_index=True)
            return unique - unique[0], y_values[indices]

        x1, y1 = prepare(x1, y1)
        x2, y2 = prepare(x2, y2)
        if x1 is None or x2 is None:
            return False

        overlap_end = min(x1[-1], x2[-1])
        if overlap_end <= 0:
            return False

        dt_values = np.concatenate((np.diff(x1), np.diff(x2)))
        dt_values = dt_values[np.isfinite(dt_values) & (dt_values > 0)]
        step = float(np.median(dt_values)) if len(dt_values) else 0.01
        step = max(step, 0.001)
        sample_count = min(4000, max(100, int(overlap_end / step) + 1))
        common_t = np.linspace(0.0, overlap_end, sample_count)
        reference = np.interp(common_t, x1, y1)
        comparison = np.interp(common_t, x2, y2)

        def trend(values, derivative=False):
            values = np.asarray(values, dtype=float)
            if len(values) >= 3:
                time_axis = np.linspace(0.0, overlap_end, len(values))
                coefficients = np.polyfit(time_axis, values, 1)
                values = values - np.polyval(coefficients, time_axis)
            if derivative:
                values = np.gradient(values, common_t)
            smooth_window = max(3, min(101, (len(values) // 100) * 2 + 1))
            if smooth_window > 3:
                kernel = np.ones(smooth_window, dtype=float) / smooth_window
                values = np.convolve(values, kernel, mode="same")
            scale = float(np.nanstd(values))
            return values / scale if scale > 1e-12 else values

        if method == "derivative":
            reference = trend(reference, derivative=True)
            comparison = trend(comparison, derivative=True)
        elif method == "trend":
            reference = trend(reference)
            comparison = trend(comparison)

        max_shift = (
            float(max_shift)
            if max_shift is not None and np.isfinite(max_shift)
            else float(min(x1[-1], x2[-1]) * 0.5)
        )
        max_shift = min(max_shift, float(min(x1[-1], x2[-1]) * 0.9))
        candidate_count = min(801, max(101, int((2.0 * max_shift) / step) + 1))
        candidate_offsets = np.linspace(-max_shift, max_shift, candidate_count)
        errors = np.full(candidate_offsets.shape, np.inf, dtype=float)
        for index, offset in enumerate(candidate_offsets):
            valid = (common_t - offset >= x2[0]) & (common_t - offset <= x2[-1])
            if valid.sum() < 10:
                continue
            shifted = np.interp(common_t[valid] - offset, common_t, comparison)
            errors[index] = float(np.mean((reference[valid] - shifted) ** 2))

        if not np.isfinite(errors).any():
            return False

        best_index = int(np.nanargmin(errors))
        self._speed_comparison_offset = float(candidate_offsets[best_index])
        self._speed_alignment_method = method
        self._update_speed_comparison()
        self.speed_comparison_plot.show()
        self.speed_comparison_plot.raise_()
        self.statusBar().showMessage(
            f"Speed profiles aligned. Test 2 shifted by {self._speed_comparison_offset:+.3f} s."
        )
        return True

    def _update_speed_comparison(self, auto_range: bool = True):
        if self._comparison_plot is None:
            return

        offsets = self._speed_comparison_offsets
        offset_summary = " | ".join(
            f"{label}: {value:+.3f} s"
            + (f" ({self._speed_alignments[label].quality:.0%})"
               if label in self._speed_alignments and self._speed_alignments[label].quality
               else "")
            for label, value in offsets.items()
        )
        self._comparison_plot.setTitle(
            f"Signal Comparison | {offset_summary}",
            color=self.current_fg(),
            size="11pt",
        )
        for item in list(self._comparison_plot.listDataItems()):
            self._comparison_plot.removeItem(item)
        if self._comparison_plot.legend is not None:
            self._comparison_plot.legend.clear()
        self._comparison_curves = []

        selections = [
            (selection, label, self.current_palette()[index % len(self.current_palette())])
            for index, (label, selection) in enumerate(self._speed_comparison_selection)
        ]

        reference_curve = None
        for selection, label, color in selections:
            if not selection:
                continue
            file_id, name, unit = selection
            data = self._fetch_raw_signal("file", file_id, name)
            if data is None:
                continue
            raw_x, raw_y, _ = data
            if len(raw_x) == 0 or len(raw_y) == 0:
                continue
            order = np.argsort(raw_x)
            x = raw_x[order].astype(float)
            y = raw_y[order].astype(float)
            x = x - x[0]
            base_x = x
            x = base_x + offsets.get(label, 0.0)
            x, y = self._downsample(x, y, self.max_plot_points)
            base_x = x - offsets.get(label, 0.0)

            if self._comparison_difference_mode and reference_curve is None:
                reference_curve = {"x": base_x, "y": y.copy(), "label": label, "unit": unit}
                continue

            if self._comparison_difference_mode and reference_curve is not None:
                ref_y = reference_curve["y"]
                if len(ref_y) != len(y):
                    common_len = min(len(ref_y), len(y))
                    ref_y = ref_y[:common_len]
                    y = y[:common_len]
                if np.allclose(ref_y, 0.0, atol=1e-9):
                    delta = np.zeros_like(y, dtype=float)
                else:
                    delta = 100.0 * (y - ref_y) / np.abs(ref_y)
                legend_label = f"{label} vs {reference_curve['label']} (Δ%)"
                y_plot = delta
                plot_title = "Difference vs reference"
            else:
                legend_label = f"{label}: {extract_signal_only_name(name)}"
                y_plot = y
                plot_title = "Signal Comparison"

            curve = self._comparison_plot.plot(
                x,
                y_plot,
                pen=pg.mkPen(color, width=1.8),
                name=legend_label,
                antialias=False,
            )
            self._comparison_curves.append({
                "type": "file",
                "file_id": file_id,
                "name": name,
                "unit": unit,
                "x": x,
                "base_x": base_x,
                "y": y_plot,
                "x_offset": offsets.get(label, 0.0),
                "curve": curve,
                "label": label,
            })

        if self._comparison_difference_mode and reference_curve is not None:
            self._comparison_plot.setTitle(
                f"Difference vs {reference_curve['label']} | {offset_summary}",
                color=self.current_fg(),
                size="11pt",
            )
            self._comparison_plot.setLabel("left", "Difference (%)")
        else:
            self._comparison_plot.setLabel("left", "Speed", units="km/h" if any(
                selection and selection[2].lower() in {"km/h", "kph"}
                for selection, _label, _color in selections
            ) else None)

        self.speed_comparison_plot.set_trace_data([
            {
                "label": curve["label"],
                "x": curve["x"],
                "base_x": curve["base_x"],
                "y": curve["y"],
                "offset": curve["x_offset"],
                "curve": curve["curve"],
            }
            for curve in self._comparison_curves
        ])
        if self._comparison_curves and auto_range:
            # Fit once, then freeze the view. Leaving auto-range enabled makes
            # every horizontal drag change the zoom and destroys alignment.
            view_box = self._comparison_plot.getViewBox()
            view_box.autoRange()
            view_box.disableAutoRange()
        if self._comparison_curves:
            self.speed_comparison_plot.show()
            self.speed_comparison_plot.raise_()

    def on_file_selected(self):
        self.filter_signals(self.search_box.text())

    def populate_signal_tree(self):
        tree = self.signal_tree
        tree.setUpdatesEnabled(False)
        tree.blockSignals(True)
        sorting_enabled = tree.isSortingEnabled()
        tree.setSortingEnabled(False)
        try:
            tree.clear()
            for name, info in self.derived_signals.items():
                leaf = extract_signal_only_name(name)
                unit = info.get("unit", "")
                item = QTreeWidgetItem([leaf, unit, "Derived"])
                item.setData(0, Qt.UserRole, {"type": "derived", "name": name, "file_id": "derived"})
                item.setData(0, Qt.UserRole + 1, f"{name} {leaf} {unit} derived")
                item.setToolTip(0, f"Signal: {leaf}\nFull Name: {name}\nSource: Derived (⚡)\nUnit: {unit or 'None'}")
                item.setForeground(0, QColor("#FF9F0A"))
                tree.addTopLevelItem(item)

            for fid, rec in self.loaded_files.items():
                for name in rec.signal_names:
                    leaf = extract_signal_only_name(name)
                    unit = rec.signal_units.get(name, "")
                    description = rec.signal_descriptions.get(name, "")
                    display_name = rec.signal_display_names.get(name, leaf)
                    source_type = getattr(rec, "source_type", Path(rec.path).suffix.upper().lstrip("."))
                    source_tail = str(fid)
                    if len(source_tail) > 22:
                        source_tail = f"...{source_tail[-19:]}"
                    source_label = f"{source_type} · {source_tail}"
                    item = QTreeWidgetItem([leaf, unit, source_label])
                    item.setData(0, Qt.UserRole, {"type": "file", "name": name, "file_id": fid})
                    item.setData(
                        0,
                        Qt.UserRole + 1,
                        f"{name} {leaf} {display_name} {unit} {description} "
                        f"{source_type} {fid}",
                    )
                    standard = getattr(rec, "chn_standard", "")
                    description_text = f"\nDescription: {description}" if description else ""
                    standard_text = f"\nStandard: {standard}" if standard else ""
                    display_text = f"\nDisplay name: {display_name}" if display_name != leaf else ""
                    full_tooltip = f"Signal: {leaf}\nFull Channel: {name}{display_text}\nFile: {fid}\nUnit: {unit or 'None'}{description_text}{standard_text}"
                    item.setToolTip(0, full_tooltip)
                    item.setToolTip(2, f"Full source: {fid}")
                    item.setForeground(2, QColor(rec.color_tag))
                    tree.addTopLevelItem(item)
        finally:
            tree.setSortingEnabled(sorting_enabled)
            tree.blockSignals(False)
            tree.setUpdatesEnabled(True)
        self.filter_signals(self.search_box.text())

    def filter_signals(self, text: str):
        query = text.lower().strip()
        tokens = query.split()
        selected_item = self.file_list.currentItem()
        selected_file_id = selected_item.data(Qt.UserRole) if selected_item else None
        visible_count = 0
        total_in_scope = 0

        for i in range(self.signal_tree.topLevelItemCount()):
            item = self.signal_tree.topLevelItem(i)
            payload = item.data(0, Qt.UserRole) or {}
            sig_file_id = payload.get("file_id")

            if selected_file_id is not None and sig_file_id != selected_file_id:
                item.setHidden(True)
                continue

            total_in_scope += 1
            search_text = str(item.data(0, Qt.UserRole + 1) or "").lower()
            if not search_text:
                search_text = (
                    f"{payload.get('name', item.text(0))} {item.text(0)} "
                    f"{item.text(1)} {item.text(2)} {payload.get('file_id', '')}"
                ).lower()
            matches = not tokens or all(token in search_text for token in tokens)
            item.setHidden(not matches)
            if matches:
                visible_count += 1

        if query:
            self.lbl_search_results.setText(
                f"{visible_count:,} of {total_in_scope:,} signals match — clear search to restore all"
            )
        else:
            self.lbl_search_results.setText(f"{visible_count:,} signals available")

    def handle_signals_dropped(self, sig_list: List[Dict[str, Any]]):
        for sig_info in sig_list:
            self._add_signal_to_plotted(sig_info)
        self.plot_current_curves(rebuild_layout=False)

    def plot_selected_signals(self):
        selected_items = self.signal_tree.selectedItems()
        if not selected_items:
            QMessageBox.information(self, "No Signals Selected", "Please select one or more signals from the list to plot.")
            return

        if not self.chk_keep.isChecked():
            self._plotted_curves.clear()

        for item in selected_items:
            sig_info = item.data(0, Qt.UserRole)
            if sig_info:
                self._add_signal_to_plotted(sig_info)

        self.plot_current_curves(rebuild_layout=not self.chk_keep.isChecked())

    def _add_signal_to_plotted(self, sig_info: Dict[str, Any]):
        file_id = sig_info.get("file_id")
        name = sig_info.get("name")
        for c in self._plotted_curves:
            if c["file_id"] == file_id and c["name"] == name:
                return
        self._plotted_curves.append(sig_info)

    def on_align_zero_toggled(self, checked: bool):
        self.align_to_zero = checked
        self.plot_current_curves(rebuild_layout=True)

    def clear_plots(self):
        self._plotted_curves.clear()
        self._reset_plot_layout()
        if self._comparison_plot is not None:
            for item in list(self._comparison_plot.listDataItems()):
                self._comparison_plot.removeItem(item)
            if self._comparison_plot.legend is not None:
                self._comparison_plot.legend.clear()
            self._comparison_curves.clear()
            self._speed_comparison_offset = 0.0
            self._speed_comparison_offsets = {"Test 1": 0.0, "Test 2": 0.0}
            self._speed_alignments = {}
            self._speed_comparison_selection = []
            self.speed_comparison_plot.hide()
        self.statusBar().showMessage("Plot cleared.")

    def _reset_plot_layout(self):
        comparison_cursor_items = []
        if self._comparison_plot is not None:
            for cursor_map in (
                self._cursor_lines,
                self._cursor_delta_lines,
                self._cursor_horizontal_lines,
                self._cursor_delta_horizontal_lines,
                self._cursor_value_labels,
                self._cursor_delta_labels,
            ):
                item = cursor_map.get(self._comparison_plot)
                if item is not None:
                    comparison_cursor_items.append(item)
            for item in comparison_cursor_items:
                try:
                    self._comparison_plot.removeItem(item)
                except Exception:
                    pass
        for vb, _axis in self._overlay_extra_viewboxes.values():
            try:
                vb.setXLink(None)
                scene = vb.scene()
                if scene is not None:
                    scene.removeItem(vb)
            except Exception:
                pass

        self.plot_canvas.clear()
        self._cursor_lines.clear()
        self._cursor_value_labels.clear()
        self._cursor_delta_lines.clear()
        self._cursor_delta_labels.clear()
        self._cursor_horizontal_lines.clear()
        self._cursor_delta_horizontal_lines.clear()
        self._plot_signal_map.clear()
        self._plot_legend_map.clear()
        self._legend_signal_map.clear()
        self._overlay_plot = None
        self._overlay_legend = None
        self._overlay_unit_main = None
        self._overlay_extra_viewboxes.clear()
        self._separate_plots.clear()
        self.cursor1_x = None
        self.cursor2_x = None
        self.cursor1_y = None
        self.cursor2_y = None
        if self._comparison_plot is not None:
            self._install_cursor(self._comparison_plot)

    def set_cursor_visible(self, visible: bool):
        self.cursor_visible = visible
        if not visible:
            self.cursor_locked = False
        for line in self._cursor_lines.values():
            # InfiniteLine starts at x=0.  Do not briefly expose that default
            # before the first real mouse position has been received.
            line.setVisible(visible and self.cursor1_x is not None)
        for line in self._cursor_delta_lines.values():
            line.setVisible(
                visible and self.delta_cursor_visible and self.cursor2_x is not None
            )
        for label in self._cursor_value_labels.values():
            label.setVisible(False)
        if not visible:
            self.statusBar().showMessage("Cursor hidden")

    def set_horizontal_cursor_visible(self, visible: bool):
        for line in self._cursor_horizontal_lines.values():
            line.setVisible(visible and self.cursor1_y is not None)
        for line in self._cursor_delta_horizontal_lines.values():
            line.setVisible(
                visible and self.delta_cursor_visible and self.cursor2_y is not None
            )
        if not visible:
            self.statusBar().showMessage("Horizontal cursor hidden")

    def set_delta_cursor_visible(self, visible: bool):
        self.delta_cursor_visible = visible
        for line in self._cursor_delta_lines.values():
            line.setVisible(
                visible and self.cursor_visible and self.cursor2_x is not None
            )
        for line in self._cursor_delta_horizontal_lines.values():
            line.setVisible(
                visible
                and self.chk_horizontal_cursor.isChecked()
                and self.cursor2_y is not None
            )
        for label in self._cursor_delta_labels.values():
            label.setVisible(False)
        if not visible:
            self.statusBar().showMessage("Secondary cursor hidden")

    def _install_cursor(self, plot: pg.PlotItem):
        line = pg.InfiniteLine(angle=90, movable=False, pen=pg.mkPen("#FFD60A", width=1.5))
        plot.addItem(line, ignoreBounds=True)
        line.setVisible(self.cursor_visible and self.cursor1_x is not None)
        self._cursor_lines[plot] = line

        delta_line = pg.InfiniteLine(
            angle=90,
            movable=False,
            pen=pg.mkPen("#FF9F0A", width=1.5, style=Qt.DashLine),
        )
        plot.addItem(delta_line, ignoreBounds=True)
        delta_line.setVisible(
            self.cursor_visible and self.delta_cursor_visible and self.cursor2_x is not None
        )
        self._cursor_delta_lines[plot] = delta_line

        horizontal_line = pg.InfiniteLine(
            angle=0,
            movable=False,
            pen=pg.mkPen("#64D2FF", width=1.3),
        )
        plot.addItem(horizontal_line, ignoreBounds=True)
        horizontal_line.setVisible(
            self.chk_horizontal_cursor.isChecked() and self.cursor1_y is not None
        )
        self._cursor_horizontal_lines[plot] = horizontal_line

        delta_horizontal_line = pg.InfiniteLine(
            angle=0,
            movable=False,
            pen=pg.mkPen("#30D158", width=1.3, style=Qt.DashLine),
        )
        plot.addItem(delta_horizontal_line, ignoreBounds=True)
        delta_horizontal_line.setVisible(
            self.delta_cursor_visible
            and self.chk_horizontal_cursor.isChecked()
            and self.cursor2_y is not None
        )
        self._cursor_delta_horizontal_lines[plot] = delta_horizontal_line

        label = pg.TextItem(anchor=(0, 1), color=self.current_fg(), fill=pg.mkBrush(
            "#000000CC" if self.dark_mode else "#FFFFFFDD"
        ))
        label.setVisible(False)
        plot.addItem(label, ignoreBounds=True)
        self._cursor_value_labels[plot] = label

        delta_label = pg.TextItem(anchor=(0, 1), color=self.current_fg(), fill=pg.mkBrush(
            "#000000CC" if self.dark_mode else "#FFFFFFDD"
        ))
        delta_label.setVisible(False)
        plot.addItem(delta_label, ignoreBounds=True)
        self._cursor_delta_labels[plot] = delta_label

        plot.getViewBox().setMouseEnabled(x=True, y=not self.x_zoom_only)

    def toggle_cursor_lock(self, event):
        if not self.cursor_visible or event.button() != Qt.LeftButton:
            return
        target_plot = next(
            (plot for plot in self._cursor_lines
             if plot.vb.sceneBoundingRect().contains(event.scenePos())),
            None,
        )
        if target_plot is None:
            return

        view_point = target_plot.vb.mapSceneToView(event.scenePos())
        x_range, y_range = target_plot.vb.viewRange()
        cursor_x = float(np.clip(view_point.x(), x_range[0], x_range[1]))
        cursor_y = float(np.clip(view_point.y(), y_range[0], y_range[1]))
        if self.delta_cursor_visible:
            if self.cursor1_x is None or self.cursor2_x is not None:
                self.cursor1_x = cursor_x
                self.cursor1_y = cursor_y
                self.cursor2_x = None
                self.cursor2_y = None
            else:
                self.cursor2_x = cursor_x
                self.cursor2_y = cursor_y
            self.cursor_locked = True
        else:
            self.cursor1_x = cursor_x
            self.cursor1_y = cursor_y
            self.cursor_locked = not self.cursor_locked
        self._set_cursor_position(self.cursor1_x, self.cursor2_x, self.cursor1_y, self.cursor2_y)
        if self.cursor_locked:
            self.update_cursor(event.scenePos(), force=True)
            hint = "[Click to set the second x point]" if self.delta_cursor_visible and self.cursor2_x is None else "[Two x points aligned]" if self.delta_cursor_visible else "[Cursor locked — click again to release]"
            self.statusBar().showMessage(self.statusBar().currentMessage() + "  " + hint)
        else:
            self.statusBar().showMessage("Cursor tracking")

    def update_cursor(self, position, force: bool = False):
        if not self.cursor_visible or (self.cursor_locked and not force):
            return

        target_plot = None
        for plot in self._cursor_lines:
            # Restrict tracking to the actual data ViewBox.  PlotItem's scene
            # rectangle also includes the left axis; mapping that area into
            # data coordinates caused the cursor to jump near the left edge.
            if plot.vb.sceneBoundingRect().contains(position):
                target_plot = plot
                break
        if target_plot is None:
            return

        view_point = target_plot.vb.mapSceneToView(position)
        x_range, y_range = target_plot.vb.viewRange()
        cursor_x = float(np.clip(view_point.x(), x_range[0], x_range[1]))
        cursor_y = float(np.clip(view_point.y(), y_range[0], y_range[1]))
        if not self.cursor_locked or self.cursor1_x is None:
            self.cursor1_x = cursor_x
            self.cursor1_y = cursor_y
        self._set_cursor_position(self.cursor1_x, self.cursor2_x, self.cursor1_y, self.cursor2_y)
        status_parts = [f"x1 = {self.cursor1_x:.3f} s"]
        label_lines = [f"x1 = {self.cursor1_x:.3f} s"]

        if target_plot is self._comparison_plot:
            curves_here = self._comparison_curves
        else:
            curves_here = self._plotted_curves if target_plot is self._overlay_plot else \
                [self._plot_signal_map.get(target_plot)] if self._plot_signal_map.get(target_plot) else []

        for sig in curves_here:
            data_res = self._fetch_raw_signal(sig.get("type", "file"), sig.get("file_id"), sig.get("name"))
            if data_res is None:
                continue
            x, y, unit = data_res
            if target_plot is self._comparison_plot and len(x) > 0:
                x = x - x[0] + float(sig.get("x_offset", 0.0))
            elif self.align_to_zero and len(x) > 0:
                x = x - x[0]
            if len(x) == 0 or cursor_x < x[0] or cursor_x > x[-1]:
                continue
            order = np.argsort(x)
            x = x[order]
            y = y[order]
            value = float(np.interp(self.cursor1_x, x, y))
            short = self._short_label(sig.get("file_id"), sig.get("name"), "")
            line_text = f"{short}: {value:.6g} {unit}".rstrip()
            status_parts.append(line_text)
            label_lines.append(line_text)
            if self.cursor2_x is not None and x[0] <= self.cursor2_x <= x[-1]:
                value2 = float(np.interp(self.cursor2_x, x, y))
                delta_y = value2 - value
                delta_text = f"{short}: ΔY={delta_y:.6g} {unit}".rstrip()
                status_parts.append(delta_text)
                label_lines.append(delta_text)

        if self.cursor2_x is not None:
            delta_x = self.cursor2_x - self.cursor1_x
            delta_y = self.cursor2_y - self.cursor1_y if self.cursor2_y is not None and self.cursor1_y is not None else None
            delta_summary = f", ΔY = {delta_y:.6g}" if delta_y is not None else ""
            status_parts.insert(1, f"x2 = {self.cursor2_x:.3f} s, ΔX = {delta_x:.3f} s{delta_summary}")
            label_lines.insert(1, f"x2 = {self.cursor2_x:.3f} s, ΔX = {delta_x:.3f} s{delta_summary}")

        self.statusBar().showMessage(" | ".join(status_parts) + ("  [locked]" if self.cursor_locked else ""))

        for plot, lbl in self._cursor_value_labels.items():
            lbl.setVisible(plot is target_plot)
        active_label = self._cursor_value_labels.get(target_plot)
        if active_label is not None:
            active_label.setText("\n".join(label_lines))
            y_min, y_max = target_plot.vb.viewRange()[1]
            active_label.setPos(cursor_x, y_min + (y_max - y_min) * 0.9)

    def _set_cursor_position(
        self,
        first: Optional[float],
        second: Optional[float],
        first_y: Optional[float] = None,
        second_y: Optional[float] = None,
    ):
        """Synchronize cursor lines so moving across stacked plots does not jump them."""
        if first is None and first_y is None:
            return
        for plot, line in self._cursor_lines.items():
            if first is not None:
                line.setValue(float(first))
                line.setVisible(self.cursor_visible)
            horizontal_line = self._cursor_horizontal_lines.get(plot)
            if horizontal_line is not None and first_y is not None:
                horizontal_line.setValue(float(first_y))
                horizontal_line.setVisible(self.chk_horizontal_cursor.isChecked())
            delta_line = self._cursor_delta_lines.get(plot)
            if delta_line is not None:
                if second is None or not self.delta_cursor_visible or not self.cursor_visible:
                    delta_line.setVisible(False)
                else:
                    delta_line.setValue(float(second))
                    delta_line.setVisible(True)
            delta_horizontal_line = self._cursor_delta_horizontal_lines.get(plot)
            if delta_horizontal_line is not None:
                if (
                    second_y is None
                    or not self.delta_cursor_visible
                    or not self.chk_horizontal_cursor.isChecked()
                ):
                    delta_horizontal_line.setVisible(False)
                else:
                    delta_horizontal_line.setValue(float(second_y))
                    delta_horizontal_line.setVisible(True)

    def set_x_zoom_only(self, enabled: bool):
        self.x_zoom_only = enabled
        for plot in self._cursor_lines:
            plot.getViewBox().setMouseEnabled(x=True, y=not enabled)
        for vb, _axis in self._overlay_extra_viewboxes.values():
            vb.setMouseEnabled(x=True, y=not enabled)
        if self._comparison_plot is not None:
            self._comparison_plot.getViewBox().setMouseEnabled(x=True, y=not enabled)

    @staticmethod
    def _downsample(x: np.ndarray, y: np.ndarray, max_points: int) -> Tuple[np.ndarray, np.ndarray]:
        if len(x) <= max_points:
            return x, y
        step = int(np.ceil(len(x) / max_points))
        return x[::step], y[::step]

    def _short_label(self, file_id: str, name: str, unit: str) -> str:
        core = simplify_signal_name(name)
        if unit:
            core = f"{core} [{unit}]"
        if len(self.loaded_files) > 1 and file_id != "derived":
            core = f"{file_id}: {core}"
        elif file_id == "derived":
            core = f"⚡ {core}"
        return core

    def _sync_overlay_views(self):
        if self._overlay_plot is None:
            return
        rect = self._overlay_plot.vb.sceneBoundingRect()
        for vb, _ in self._overlay_extra_viewboxes.values():
            vb.setGeometry(rect)
            vb.linkedViewChanged(self._overlay_plot.vb, vb.XAxis)

    def _style_plot(self, plot: pg.PlotItem, title: Optional[str] = None):
        fg = self.current_fg()
        for axis_name in ("left", "bottom"):
            axis = plot.getAxis(axis_name)
            axis.setPen(pg.mkPen(self.current_border(), width=1))
            axis.setTextPen(pg.mkPen(fg))
        if title:
            plot.setTitle(title, color=fg, size="11pt")
        plot.showGrid(x=True, y=True, alpha=0.15)
        plot.setClipToView(True)
        plot.setDownsampling(mode="peak", auto=True)
        plot.getViewBox().setMouseMode(pg.ViewBox.PanMode)
        self._customize_plot_context_menu(plot)

    def _customize_plot_context_menu(self, plot: pg.PlotItem):
        vb = plot.getViewBox()
        menu = vb.menu
        if menu is None:
            return
        if menu.actions():
            menu.addSeparator()
        axis_menu = menu.addMenu("Y-axis")
        main_axis_action = axis_menu.addAction("Reset all to primary Y-axis")
        main_axis_action.triggered.connect(lambda: self._set_selected_y_axis(None))
        secondary_menu = axis_menu.addMenu("Use secondary Y-axis")
        for sig in self._plotted_curves:
            label = self._short_label(sig.get("file_id"), sig.get("name"), "")
            action = secondary_menu.addAction(label)
            action.setCheckable(True)
            action.setChecked(sig.get("y_axis") == "secondary")
            action.triggered.connect(lambda checked, selected=sig: self._set_selected_y_axis(selected))
        transform_menu = menu.addMenu("Transform")
        integral_action = transform_menu.addAction("Integral (∫ y dt)")
        integral_action.triggered.connect(lambda: self._apply_integral_transform(plot))
        integral_abs_action = transform_menu.addAction("Integral Abs (|∫| |y| dt)")
        integral_abs_action.triggered.connect(lambda: self._apply_integral_abs_transform(plot))

    def _set_selected_y_axis(self, selected_sig: Optional[Dict[str, Any]]):
        if selected_sig is None:
            for sig in self._plotted_curves:
                sig.pop("y_axis", None)
        else:
            selected_sig["y_axis"] = "secondary"
        self.plot_current_curves(rebuild_layout=True)
        axis_name = "secondary" if selected_sig is not None else "primary"
        signal_name = selected_sig.get("name") if selected_sig is not None else "all signals"
        self.statusBar().showMessage(f"{signal_name} moved to {axis_name} Y-axis")

    def _apply_integral_transform(self, plot: pg.PlotItem):
        try:
            curves_to_integrate = []
            for curve in plot.listDataItems():
                if hasattr(curve, 'xData') and hasattr(curve, 'yData'):
                    x, y = curve.xData, curve.yData
                    if x is not None and y is not None and len(y) > 0:
                        name = curve.name() if hasattr(curve, 'name') else "Integral"
                        curves_to_integrate.append({'x': x, 'y': y, 'original_name': name})

            if not curves_to_integrate:
                self.statusBar().showMessage("No curves found to integrate")
                return

            legend = self._plot_legend_map.get(plot)
            palette = self.current_palette()
            for idx, curve_info in enumerate(curves_to_integrate):
                x, y = curve_info['x'], curve_info['y']
                original_name = curve_info['original_name']
                if len(x) > 1:
                    integral_y = np.zeros_like(y)
                    for i in range(1, len(x)):
                        dt = x[i] - x[i-1]
                        integral_y[i] = integral_y[i-1] + (y[i-1] + y[i]) / 2.0 * dt
                else:
                    integral_y = y.copy()

                color = palette[idx % len(palette)]
                integral_name = f"∫ {original_name}"
                integral_curve = plot.plot(x, integral_y, pen=pg.mkPen(color, width=1.5, style=Qt.DashLine),
                                          name=integral_name, antialias=False)
                if legend is not None:
                    self._register_legend_signal(legend, integral_curve, integral_name, {
                        'name': integral_name, 'file_id': 'derived', 'type': 'derived'
                    })
            self.statusBar().showMessage(f"Integrated {len(curves_to_integrate)} signal(s)")
        except Exception as e:
            QMessageBox.critical(self, "Integral Error", f"Failed to calculate integral:\n{e}")

    def _apply_integral_abs_transform(self, plot: pg.PlotItem):
        try:
            curves_to_integrate = []
            for curve in plot.listDataItems():
                if hasattr(curve, 'xData') and hasattr(curve, 'yData'):
                    x, y = curve.xData, curve.yData
                    if x is not None and y is not None and len(y) > 0:
                        name = curve.name() if hasattr(curve, 'name') else "Integral Abs"
                        curves_to_integrate.append({'x': x, 'y': y, 'original_name': name})

            if not curves_to_integrate:
                self.statusBar().showMessage("No curves found to integrate")
                return

            legend = self._plot_legend_map.get(plot)
            palette = self.current_palette()
            for idx, curve_info in enumerate(curves_to_integrate):
                x, y = curve_info['x'], curve_info['y']
                original_name = curve_info['original_name']
                if len(x) > 1:
                    integral_y = np.zeros_like(y)
                    for i in range(1, len(x)):
                        dt = x[i] - x[i-1]
                        integral_y[i] = integral_y[i-1] + np.abs((y[i-1] + y[i]) / 2.0 * dt)
                else:
                    integral_y = np.abs(y.copy())

                color = palette[idx % len(palette)]
                integral_abs_name = f"|∫| {original_name}"
                integral_curve = plot.plot(x, integral_y, pen=pg.mkPen(color, width=1.5, style=Qt.DotLine),
                                          name=integral_abs_name, antialias=False)
                if legend is not None:
                    self._register_legend_signal(legend, integral_curve, integral_abs_name, {
                        'name': integral_abs_name, 'file_id': 'derived', 'type': 'derived'
                    })
            self.statusBar().showMessage(f"Integrated (abs) {len(curves_to_integrate)} signal(s)")
        except Exception as e:
            QMessageBox.critical(self, "Integral Abs Error", f"Failed to calculate integral (abs):\n{e}")

    def _register_legend_signal(self, legend: pg.LegendItem, curve, label: str, sig: Dict[str, Any]):
        legend_label = legend.getLabel(curve)
        if legend_label is None:
            legend.addItem(curve, label)
            legend_label = legend.getLabel(curve)
        if legend_label is not None:
            legend_label.setToolTip("Right-click to remove this signal")
            legend_label.setAcceptedMouseButtons(Qt.LeftButton | Qt.RightButton)
            legend_label.installEventFilter(self)
            self._legend_signal_map[legend_label] = sig

    def eventFilter(self, watched, event):
        if event.type() == QEvent.GraphicsSceneMousePress and event.button() == Qt.RightButton:
            sig = self._legend_signal_map.get(watched)
            if sig is not None:
                self._plotted_curves = [
                    item for item in self._plotted_curves if item is not sig
                ]
                self.plot_current_curves(rebuild_layout=True)
                self.statusBar().showMessage(f"Removed signal: {sig.get('name')}")
                return True
        return super().eventFilter(watched, event)

    def plot_current_curves(self, rebuild_layout: bool = True):
        if not self._plotted_curves:
            self._reset_plot_layout()
            return

        overlay = self.combo_plot_mode.currentIndex() == 0
        if rebuild_layout or overlay != self._last_overlay_mode:
            self._reset_plot_layout()
        self._last_overlay_mode = overlay

        palette = self.current_palette()
        plotted_count = 0

        try:
            if overlay:
                if self._overlay_plot is None:
                    plot = self.plot_canvas.addPlot()
                    self._style_plot(plot, title="Multi-Signal Overlay Comparison")
                    plot.setLabel("bottom", "Time", units="s")
                    self._overlay_legend = plot.addLegend(offset=(10, 10), labelTextColor=self.current_fg())
                    self._plot_legend_map[plot] = self._overlay_legend
                    self._overlay_plot = plot
                    plot.vb.sigResized.connect(self._sync_overlay_views)
                    self._install_cursor(plot)

                plot = self._overlay_plot

                signal_units = []
                for sig in self._plotted_curves:
                    data_res = self._fetch_raw_signal(sig.get("type", "file"), sig.get("file_id"), sig.get("name"))
                    if data_res is not None:
                        signal_units.append((sig, data_res[2] or "value"))
                primary_sig = next(
                    (sig for sig, _unit in signal_units if sig.get("y_axis") != "secondary"),
                    signal_units[0][0] if signal_units else None,
                )
                if primary_sig is not None:
                    self._overlay_unit_main = next(
                        unit for sig, unit in signal_units if sig is primary_sig
                    )
                    plot.setLabel("left", self._overlay_unit_main)

                for idx, sig in enumerate(self._plotted_curves):
                    file_id = sig.get("file_id")
                    name = sig.get("name")
                    stype = sig.get("type", "file")

                    data_res = self._fetch_raw_signal(stype, file_id, name)
                    if data_res is None:
                        continue

                    raw_x, raw_y, unit = data_res
                    if self.align_to_zero and len(raw_x) > 0:
                        raw_x = raw_x - raw_x[0]

                    x, y = self._downsample(raw_x, raw_y, self.max_plot_points)
                    color = palette[idx % len(palette)]
                    label = self._short_label(file_id, name, unit)

                    unit_key = unit or "value"
                    if sig.get("y_axis") != "secondary" and unit_key == self._overlay_unit_main:
                        curve = plot.plot(x, y, pen=pg.mkPen(color, width=1.5), name=label, antialias=False)
                        self._register_legend_signal(self._overlay_legend, curve, label, sig)
                    else:
                        axis_key = unit_key if sig.get("y_axis") != "secondary" else f"secondary:{unit_key}"
                        axis_label = unit_key if sig.get("y_axis") != "secondary" else f"Secondary {unit_key}"
                        if axis_key not in self._overlay_extra_viewboxes:
                            vb = pg.ViewBox()
                            axis = pg.AxisItem("right")
                            plot.scene().addItem(vb)
                            plot.layout.addItem(axis, 2, 3 + len(self._overlay_extra_viewboxes))
                            axis.linkToView(vb)
                            vb.setXLink(plot.vb)
                            vb.setMouseEnabled(x=True, y=not self.x_zoom_only)
                            axis.setLabel(axis_label, color=color)
                            axis.setPen(pg.mkPen(color))
                            axis.setTextPen(pg.mkPen(color))
                            self._overlay_extra_viewboxes[axis_key] = (vb, axis)

                        vb, _ = self._overlay_extra_viewboxes[axis_key]
                        curve = pg.PlotDataItem(x, y, pen=pg.mkPen(color, width=1.5), antialias=False)
                        vb.addItem(curve)
                        if self._overlay_legend:
                            self._register_legend_signal(self._overlay_legend, curve, label, sig)

                    plotted_count += 1
                self._sync_overlay_views()

            else:
                self._reset_plot_layout()
                first_plot = None
                for idx, sig in enumerate(self._plotted_curves):
                    file_id = sig.get("file_id")
                    name = sig.get("name")
                    stype = sig.get("type", "file")

                    data_res = self._fetch_raw_signal(stype, file_id, name)
                    if data_res is None:
                        continue

                    raw_x, raw_y, unit = data_res
                    if self.align_to_zero and len(raw_x) > 0:
                        raw_x = raw_x - raw_x[0]

                    x, y = self._downsample(raw_x, raw_y, self.max_plot_points)
                    color = palette[idx % len(palette)]
                    title = self._short_label(file_id, name, unit)

                    p = self.plot_canvas.addPlot(row=idx, col=0)
                    self._style_plot(p, title=title)
                    self._install_cursor(p)
                    self._plot_signal_map[p] = sig
                    legend = p.addLegend(offset=(10, 10), labelTextColor=self.current_fg())
                    self._plot_legend_map[p] = legend
                    p.setLabel("left", unit or "value")
                    p.setLabel("bottom", "Time", units="s")
                    curve = p.plot(x, y, pen=pg.mkPen(color, width=1.5), name=title, antialias=False)
                    self._register_legend_signal(legend, curve, title, sig)

                    if first_plot is None:
                        first_plot = p
                    else:
                        p.setXLink(first_plot)

                    self._separate_plots.append(p)
                    plotted_count += 1

            self.statusBar().showMessage(f"Displaying {plotted_count} signal curve(s).")

        except Exception as e:
            QMessageBox.critical(self, "Plot Error", f"Failed to render plots:\n{e}\n\n{traceback.format_exc()}")

    def _fetch_raw_signal(self, stype: str, file_id: str, name: str) -> Optional[Tuple[np.ndarray, np.ndarray, str]]:
        if stype == "derived" or file_id == "derived":
            if name in self.derived_signals:
                info = self.derived_signals[name]
                return info["x"], info["y"], info.get("unit", "")
            return None

        if file_id in self.loaded_files:
            rec = self.loaded_files[file_id]
            data = rec.get_signal_data(name)
            if data is not None:
                unit = rec.signal_units.get(name, "")
                return data[0], data[1], unit
        return None

    def export_plotted_csv(self):
        if not self._plotted_curves:
            QMessageBox.warning(self, "Nothing to Export", "No signals are currently plotted.")
            return

        path, _ = QFileDialog.getSaveFileName(self, "Export Plotted Signals CSV", "plotted_signals.csv", "CSV Files (*.csv)")
        if not path:
            return

        try:
            self.statusBar().showMessage("Resampling and exporting data...")
            QApplication.processEvents()

            raster = self.export_raster
            cache = {}
            t_min, t_max = None, None

            for sig in self._plotted_curves:
                data_res = self._fetch_raw_signal(sig.get("type", "file"), sig.get("file_id"), sig.get("name"))
                if data_res is None:
                    continue
                x, y, unit = data_res
                if self.align_to_zero and len(x) > 0:
                    x = x - x[0]
                key = f"[{sig.get('file_id')}] {sig.get('name')} [{unit}]" if unit else f"[{sig.get('file_id')}] {sig.get('name')}"
                cache[key] = (x, y)
                t_min = x[0] if t_min is None else min(t_min, x[0])
                t_max = x[-1] if t_max is None else max(t_max, x[-1])

            if not cache or t_min is None or t_max is None:
                QMessageBox.warning(self, "Export Failed", "No valid signal data found to export.")
                return

            n_steps = max(2, int(np.ceil((t_max - t_min) / raster)) + 1)
            common_t = t_min + np.arange(n_steps) * raster
            common_t = common_t[common_t <= t_max]

            df_dict = {"time_s": common_t}
            for col_name, (x, y) in cache.items():
                df_dict[col_name] = np.interp(common_t, x, y)

            df = pd.DataFrame(df_dict)
            df.to_csv(path, index=False)
            self.statusBar().showMessage(f"Exported {len(df)} rows to {Path(path).name}")
            QMessageBox.information(
                self, "Export Successful", f"Successfully exported {len(df)} rows to:\n{path}"
            )
        except Exception as e:
            QMessageBox.critical(
                self, "Export Error", f"Failed to export CSV:\n{e}\n\n{traceback.format_exc()}"
            )
            self.statusBar().showMessage("Export failed.")

    def closeEvent(self, event):
        try:
            self._reset_plot_layout()
        except Exception:
            pass
        for rec in self.loaded_files.values():
            rec.close()
        event.accept()


class EngineMountNvhDialog(QDialog):
    """CHN-aware vibration analysis view for loaded CSV/XLSX/CHN tests."""

    def __init__(self, parent, records, dark_mode=True):
        super().__init__(parent)
        self.records = records
        self.dark_mode = dark_mode
        self._rows = []
        self._build_ui()
        self._populate_channels()
        self._refresh_analysis()

    def _build_ui(self):
        self.setWindowTitle("Engine Mount and NVH Analysis")
        self.resize(1450, 900)
        layout = QVBoxLayout(self)
        toolbar = QHBoxLayout()
        toolbar.addWidget(QLabel("Test:"))
        self.test_combo = QComboBox()
        self.test_combo.addItems([record.file_id for record in self.records])
        self.test_combo.currentIndexChanged.connect(self._populate_channels)
        toolbar.addWidget(self.test_combo)
        toolbar.addWidget(QLabel("Channel:"))
        self.channel_combo = QComboBox()
        self.channel_combo.currentIndexChanged.connect(self._refresh_analysis)
        toolbar.addWidget(self.channel_combo)
        self.code_search = QLineEdit()
        self.code_search.setPlaceholderText("Search channel code...")
        self.code_search.returnPressed.connect(self._find_channel_code)
        toolbar.addWidget(self.code_search)
        self.view_combo = QComboBox()
        self.view_combo.addItems(["Time signal", "FFT spectrum", "PSD spectrum"])
        self.view_combo.currentIndexChanged.connect(self._refresh_plot)
        toolbar.addWidget(self.view_combo)
        self.band_combo = QComboBox()
        self.band_combo.addItems(["All frequencies", "0-50 Hz", "50-200 Hz", "200-500 Hz", "500-1000 Hz", "1000-2000 Hz"])
        self.band_combo.currentIndexChanged.connect(self._refresh_analysis)
        toolbar.addWidget(self.band_combo)
        self.compare_combo = QComboBox()
        self.compare_combo.addItem("No comparison")
        self.compare_combo.addItems([record.file_id for record in self.records])
        self.compare_combo.currentIndexChanged.connect(self._refresh_analysis)
        toolbar.addWidget(QLabel("Compare:"))
        toolbar.addWidget(self.compare_combo)
        export_button = QPushButton("Export report")
        export_button.clicked.connect(self._export_report)
        toolbar.addWidget(export_button)
        inventory_button = QPushButton("Export CHN map")
        inventory_button.clicked.connect(self._export_channel_inventory)
        toolbar.addWidget(inventory_button)
        toolbar.addStretch()
        layout.addLayout(toolbar)

        self.summary = QLabel()
        self.summary.setWordWrap(True)
        layout.addWidget(self.summary)
        self.tabs = QTabWidget()
        layout.addWidget(self.tabs)

        dashboard = QWidget()
        dashboard_layout = QVBoxLayout(dashboard)
        self.stats_table = QTableWidget(0, 9)
        self.stats_table.setHorizontalHeaderLabels(["Test", "Location", "Direction", "Unit", "Peak", "Min", "Mean", "RMS", "Crest factor"])
        self.stats_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.stats_table.setSortingEnabled(True)
        self.stats_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        dashboard_layout.addWidget(self.stats_table)
        self.quality_text = QTextEdit()
        self.quality_text.setReadOnly(True)
        dashboard_layout.addWidget(self.quality_text)
        self.tabs.addTab(dashboard, "Dashboard")

        plot_page = QWidget()
        plot_layout = QVBoxLayout(plot_page)
        self.plot = pg.PlotWidget()
        plot_layout.addWidget(self.plot)
        self.tabs.addTab(plot_page, "Signal analysis")

    def _record(self, file_id):
        return next((record for record in self.records if record.file_id == file_id), None)

    def _find_channel_code(self):
        query = _compact_chn_key(self.code_search.text())
        if not query:
            return
        for index in range(self.channel_combo.count()):
            name = self.channel_combo.itemData(index)
            record = self._record(self.test_combo.currentText())
            metadata = record.signal_metadata.get(name, {}) if record else {}
            aliases = [name, metadata.get("channel_code", ""), metadata.get("laboratory_channel_code", ""), metadata.get("customer_channel_code", "")]
            if any(query == _compact_chn_key(alias) or query in _compact_chn_key(alias) for alias in aliases):
                self.channel_combo.setCurrentIndex(index)
                return
        self.quality_text.setPlainText(f"Channel code not found: {self.code_search.text()}")

    def _metadata(self, record, name):
        metadata = record.signal_metadata.get(name, {})
        code = name
        description = record.signal_descriptions.get(name, "")
        display = metadata.get("channel_name") or record.signal_display_names.get(name, name)
        text = f"{display} {description} {code}".lower()
        direction = metadata.get("direction", "") or next((axis for axis in ("X", "Y", "Z") if re.search(rf"(?:^|[^a-z]){axis.lower()}(?:$|[^a-z])", text)), "")
        if not direction:
            direction = next((axis for axis in ("X", "Y", "Z") if axis.lower() in code.lower()), "")
        location = metadata.get("location", "") or (description.split(":", 1)[0].strip() if ":" in description else display)
        return display, location, direction, metadata.get("unit", "") or record.signal_units.get(name, "")

    def _populate_channels(self):
        record = self._record(self.test_combo.currentText())
        self.channel_combo.blockSignals(True)
        self.channel_combo.clear()
        if record:
            for name in record.signal_names:
                display, location, direction, unit = self._metadata(record, name)
                code = record.signal_metadata.get(name, {}).get("channel_code", name)
                self.channel_combo.addItem(
                    f"{display} | {code} | {direction or '?'} | {location or 'Unknown location'}",
                    name,
                )
        self.channel_combo.blockSignals(False)
        self._refresh_analysis()

    def _signal(self, record, name):
        data = record.get_signal_data(name)
        if data is None:
            return None
        timestamps, values = data
        valid = np.isfinite(timestamps) & np.isfinite(values)
        if valid.sum() < 2:
            return None
        return timestamps[valid].astype(float), values[valid].astype(float)

    def _stats(self, x, y):
        centered = y - np.mean(y)
        rms = float(np.sqrt(np.mean(y ** 2)))
        peak = float(np.max(np.abs(y)))
        return {
            "peak": peak, "minimum": float(np.min(y)), "mean": float(np.mean(y)),
            "rms": rms, "std": float(np.std(y)),
            "crest": peak / rms if rms else np.nan,
            "peak_to_peak": float(np.ptp(y)), "samples": len(y),
            "duration": float(x[-1] - x[0]), "offset": float(np.mean(centered)),
        }

    def _spectrum(self, x, y):
        dt = np.diff(x)
        dt = dt[np.isfinite(dt) & (dt > 0)]
        if len(dt) == 0:
            return np.array([]), np.array([]), np.array([])
        fs = 1.0 / float(np.median(dt))
        values = y - np.mean(y)
        frequencies = np.fft.rfftfreq(len(values), 1.0 / fs)
        fft_values = np.abs(np.fft.rfft(values)) / max(1, len(values)) * 2.0
        window = np.hanning(len(values))
        psd = np.abs(np.fft.rfft(values * window)) ** 2 / max(1.0, fs * np.sum(window ** 2))
        return frequencies, fft_values, psd

    def _refresh_analysis(self):
        record = self._record(self.test_combo.currentText())
        name = self.channel_combo.currentData()
        if record is None or not name:
            return
        signal = self._signal(record, name)
        if signal is None:
            return
        x, y = signal
        stats = self._stats(x, y)
        display, location, direction, unit = self._metadata(record, name)
        comparison = ""
        compare_id = self.compare_combo.currentText()
        if compare_id != "No comparison":
            other = self._record(compare_id)
            other_signal = self._signal(other, name) if other else None
            if other_signal:
                other_stats = self._stats(*other_signal)
                change = 100.0 * (stats["rms"] - other_stats["rms"]) / abs(other_stats["rms"]) if other_stats["rms"] else np.nan
                comparison = f" | {compare_id} RMS {other_stats['rms']:.6g}, change {change:+.2f}%"
        self.summary.setText(
            f"{display} | {location} | Direction {direction or 'unknown'} | "
            f"Samples {stats['samples']} | Duration {stats['duration']:.3f} s | "
            f"RMS {stats['rms']:.6g} {unit}{comparison}"
        )
        self.stats_table.setRowCount(1)
        values = [record.file_id, location, direction, unit, stats["peak"], stats["minimum"], stats["mean"], stats["rms"], stats["crest"]]
        for column, value in enumerate(values):
            self.stats_table.setItem(0, column, QTableWidgetItem(f"{value:.6g}" if isinstance(value, float) else str(value)))
        self._quality_report(record, name, x, y)
        self._refresh_plot()

    def _quality_report(self, record, name, x, y):
        issues = []
        if len(x) > 2:
            intervals = np.diff(x)
            median = np.median(intervals)
            if median <= 0 or np.max(np.abs(intervals - median)) > max(1e-9, median * 0.05):
                issues.append("Sampling interval is irregular")
        if len(y) and (not np.isfinite(y).all()):
            issues.append("Corrupted or non-finite samples detected")
        if np.ptp(y) == 0:
            issues.append("Signal is constant")
        metadata = record.signal_metadata.get(name, {})
        status = "VALID" if not issues else "; ".join(issues)
        lines = [
            f"Channel health: {status}",
            f"Sensor name: {metadata.get('channel_name', name)}",
            f"Source column: {name}",
            f"Channel code: {metadata.get('channel_code', name)}",
            f"Laboratory code: {metadata.get('laboratory_channel_code', '')}",
            f"Customer code: {metadata.get('customer_channel_code', '')}",
            f"Location: {metadata.get('location', '')}",
            f"Direction: {metadata.get('direction', '')}",
            f"Unit: {metadata.get('unit', record.signal_units.get(name, ''))}",
            f"Sensor type: {metadata.get('transducer_type', metadata.get('sensor_type', ''))}",
            f"Transducer ID: {metadata.get('transducer_id', '')}",
            f"Calibration date: {metadata.get('calibration_date', '')}",
            f"Calibration due date: {metadata.get('calibration_due_date', '')}",
            f"Sampling rate: {metadata.get('sampling_rate', 1.0 / np.median(np.diff(x)) if len(x) > 1 else '')} Hz",
            f"Sample count: {len(y)}",
            f"Maximum: {np.max(y):.6g}",
            f"Minimum: {np.min(y):.6g}",
        ]
        self.quality_text.setPlainText("\n".join(lines))

    def _refresh_plot(self):
        record = self._record(self.test_combo.currentText())
        name = self.channel_combo.currentData()
        if record is None or not name:
            return
        signal = self._signal(record, name)
        if signal is None:
            return
        x, y = signal
        frequencies, fft_values, psd = self._spectrum(x, y)
        self.plot.clear()
        mode = self.view_combo.currentText()
        color = self.parent().current_palette()[0] if hasattr(self.parent(), "current_palette") else "#0A84FF"
        if mode == "Time signal":
            self.plot.plot(x, y, pen=pg.mkPen(color, width=1.4))
            self.plot.setLabel("bottom", "Time", units="s")
            self.plot.setLabel("left", "Amplitude")
        elif mode == "FFT spectrum":
            self.plot.plot(frequencies, fft_values, pen=pg.mkPen(color, width=1.4))
            self.plot.setLabel("bottom", "Frequency", units="Hz")
            self.plot.setLabel("left", "Amplitude")
        else:
            self.plot.plot(frequencies, psd, pen=pg.mkPen(color, width=1.4))
            self.plot.setLabel("bottom", "Frequency", units="Hz")
            self.plot.setLabel("left", "PSD")
        self.plot.showGrid(x=True, y=True, alpha=0.2)

    def _export_report(self):
        record = self._record(self.test_combo.currentText())
        name = self.channel_combo.currentData()
        signal = self._signal(record, name) if record and name else None
        if signal is None:
            return
        x, y = signal
        stats = self._stats(x, y)
        display, location, direction, unit = self._metadata(record, name)
        destination, _ = QFileDialog.getSaveFileName(self, "Export NVH report", "engine_mount_report.xlsx", "Excel (*.xlsx);;CSV (*.csv)")
        if not destination:
            return
        frame = pd.DataFrame([{
            "Test": record.file_id, "Channel": display, "Location": location, "Direction": direction,
            "Channel code": record.signal_metadata.get(name, {}).get("channel_code", name),
            "Laboratory channel code": record.signal_metadata.get(name, {}).get("laboratory_channel_code", ""),
            "Customer channel code": record.signal_metadata.get(name, {}).get("customer_channel_code", ""),
            "Unit": unit, **stats,
        }])
        try:
            if destination.lower().endswith(".csv"):
                frame.to_csv(destination, index=False)
            else:
                frame.to_excel(destination, index=False)
            self.quality_text.append(f"\nReport exported: {destination}")
        except Exception as exc:
            QMessageBox.critical(self, "Export failed", str(exc))

    def _export_channel_inventory(self):
        destination, _ = QFileDialog.getSaveFileName(
            self, "Export CHN channel map", "chn_channel_inventory.xlsx",
            "Excel (*.xlsx);;CSV (*.csv)"
        )
        if not destination:
            return
        rows = []
        for record in self.records:
            for name in record.signal_names:
                display, location, direction, unit = self._metadata(record, name)
                metadata = record.signal_metadata.get(name, {})
                data = record.get_signal_data(name)
                rows.append({
                    "Test": record.file_id,
                    "Signal name": display,
                    "Source column": name,
                    "Channel code": metadata.get("channel_code", name),
                    "Laboratory channel code": metadata.get("laboratory_channel_code", ""),
                    "Customer channel code": metadata.get("customer_channel_code", ""),
                    "Location": location,
                    "Direction": direction,
                    "Unit": unit,
                    "Samples": len(data[1]) if data else 0,
                })
        frame = pd.DataFrame(rows)
        try:
            if destination.lower().endswith(".csv"):
                frame.to_csv(destination, index=False)
            else:
                frame.to_excel(destination, index=False)
            self.quality_text.append(f"\nCHN map exported: {destination}")
        except Exception as exc:
            QMessageBox.critical(self, "Export failed", str(exc))


def main():
    app = QApplication(sys.argv)
    window = MF4VehicleExplorer()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
